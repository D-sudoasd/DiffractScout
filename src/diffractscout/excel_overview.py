"""Build a compact, navigable overview sheet for exported workbooks."""

from __future__ import annotations

import math
from typing import Any
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.cell import quote_sheetname
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet


_NAVY = "16324F"
_TEAL = "167D8D"
_PALE_BLUE = "F2F6FA"
_TEXT = "203040"
_WHITE = "FFFFFF"
_FONT_NAME = "Microsoft YaHei"
_TITLE_FONT = Font(name=_FONT_NAME, size=18, bold=True, color=_WHITE)
_SECTION_FONT = Font(name=_FONT_NAME, size=11, bold=True, color=_WHITE)
_HEADER_FONT = Font(name=_FONT_NAME, bold=True, color=_WHITE)
_BODY_FONT = Font(name=_FONT_NAME, size=10, color=_TEXT)
_LINK_FONT = Font(name=_FONT_NAME, color=_TEAL, underline="single")
_TITLE_ALIGNMENT = Alignment(vertical="center")
_LEFT_ALIGNMENT = Alignment(horizontal="left", vertical="center")
_CENTER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
_CENTER_WRAP_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
_TOP_WRAP_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
_DEFAULT_ALIGNMENT = Alignment()
_NAVY_FILL = PatternFill("solid", fgColor=_NAVY)
_TEAL_FILL = PatternFill("solid", fgColor=_TEAL)
_PALE_BLUE_FILL = PatternFill("solid", fgColor=_PALE_BLUE)
_WHITE_FILL = PatternFill("solid", fgColor=_WHITE)
_BOTTOM_BORDER = Border(bottom=Side(style="hair", color="D5DEE7"))


def _safe_text(value: Any) -> str:
    """Keep untrusted workbook text literal instead of letting Excel parse formulas."""

    text = str(value)
    if text and text[0] in {"=", "+", "-", "@", "\t", "\r", "\n"}:
        return "'" + text
    return text


def _summary_values(summary: list[dict[str, Any]]) -> dict[str, Any]:
    values: dict[str, Any] = {}
    for row in summary:
        key = row.get("key")
        if isinstance(key, str):
            values[key] = row.get("value")
    return values


def _summary_text(value: Any, *, missing: str) -> str:
    if value is None or value == "":
        return missing
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (list, tuple)):
        if len(value) != 2 or any(item is None for item in value):
            return missing
        left, right = (_summary_scalar(item) for item in value)
        return f"{left}–{right}°"
    return _summary_scalar(value)


def _summary_scalar(value: Any) -> str:
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float) and math.isfinite(value):
        return format(value, ".5g")
    return _safe_text(value)


def _significant_number_format(value: int | float, digits: int = 5) -> str:
    """Return an Excel format with at most ``digits`` significant figures."""

    numeric = float(value)
    if not math.isfinite(numeric) or numeric == 0:
        return "0"
    exponent = math.floor(math.log10(abs(numeric)))
    decimal_places = max(0, digits - 1 - exponent)
    if exponent >= digits or decimal_places > 12:
        return "0.0000E+00"
    return "0" if decimal_places == 0 else "0." + ("#" * decimal_places)


def _wrapped_line_count(value: Any, width: float) -> int:
    """Estimate wrapped lines using Excel column units and Unicode character width."""

    text = str(value or "").replace("\r\n", "\n").replace("\r", "\n")
    capacity = max(1.0, float(width) - 1.0)
    lines = 1
    used = 0.0
    for character in text:
        if character == "\n":
            lines += 1
            used = 0.0
            continue
        if unicodedata.combining(character):
            advance = 0.0
        else:
            advance = 2.0 if unicodedata.east_asian_width(character) in {"W", "F"} else 1.0
        if used > 0 and used + advance > capacity:
            lines += 1
            used = 0.0
        used += advance
    return lines


def _estimated_row_height(values_and_widths: tuple[tuple[Any, float], ...]) -> float:
    lines = max((_wrapped_line_count(value, width) for value, width in values_and_widths), default=1)
    return min(409.5, max(30.0, lines * 15.0 + 6.0))


def _write_numeric_or_text(cell: Any, value: Any, *, missing: str) -> None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        numeric = float(value)
        if math.isfinite(numeric):
            cell.value = value
            cell.number_format = _significant_number_format(numeric)
            cell.alignment = _TITLE_ALIGNMENT
            return
    cell.value = _summary_scalar(value) if value is not None else missing


def _display_radiation_source(value: Any, *, lab_views: bool) -> str:
    if value is None or value == "":
        return "无数据" if lab_views else "No data"
    source = str(value)
    if source.startswith("source_preset:"):
        source = source.split(":", 1)[1]
    sources = {
        "Cu Ka": "Cu Kα",
        "Co Ka": "Co Kα",
        "Fe Ka": "Fe Kα",
        "Mo Ka": "Mo Kα",
        "Ag Ka": "Ag Kα",
    }
    if source in sources:
        return sources[source]
    if source == "wavelength_A":
        return "波长输入" if lab_views else "Wavelength input"
    if source == "energy_keV":
        return "能量输入" if lab_views else "Energy input"
    if source == "custom_source_wavelength":
        return "自定义波长" if lab_views else "Custom wavelength"
    return _safe_text(source)


def _display_profile_model(value: Any, *, lab_views: bool, missing: str) -> str:
    if value is None or value == "":
        return missing
    model = str(value).casefold().replace("-", "_")
    names = {
        "pseudo_voigt": ("伪Voigt", "pseudo-Voigt"),
        "gaussian": ("高斯", "Gaussian"),
        "lorentzian": ("洛伦兹", "Lorentzian"),
    }
    if model in names:
        return names[model][0 if lab_views else 1]
    return _safe_text(value)


def _display_input_mode(value: Any, *, lab_views: bool, missing: str) -> str:
    if value is None or value == "":
        return missing
    modes = {
        "source": ("预设光源", "Source preset"),
        "wavelength": ("波长", "Wavelength"),
        "energy": ("能量", "Photon energy"),
    }
    mode = str(value).casefold()
    if mode in modes:
        return modes[mode][0 if lab_views else 1]
    return _safe_text(value)


def add_overview_sheet(
    workbook: Workbook,
    *,
    summary: list[dict[str, Any]],
    phases: list[dict[str, Any]],
    diagnostics: list[dict[str, Any]],
    lab_views: bool,
) -> Worksheet:
    """Add a concise overview using only the supplied export rows and summary.

    The caller should invoke this after creating the workbook's data sheets so
    the navigation section can link only to sheets that actually exist.
    """

    title = "结果概览" if lab_views else "Overview"
    if title in workbook.sheetnames:
        raise ValueError(f"Workbook already contains an overview sheet named {title!r}.")

    labels = {
        "title": "本次衍射分析结果概览" if lab_views else "Diffraction analysis overview",
        "summary": "导出摘要" if lab_views else "Export summary",
        "run": "运行条件" if lab_views else "Run conditions",
        "phases": "各物相摘要" if lab_views else "Phase summary",
        "navigation": "工作表导航" if lab_views else "Worksheet navigation",
        "metric": "指标" if lab_views else "Metric",
        "count": "数量" if lab_views else "Count",
        "phase_count": "物相数" if lab_views else "Phases",
        "reflection_count": "反射数" if lab_views else "Reflections",
        "error_count": "错误数" if lab_views else "Errors",
        "warning_count": "警告数" if lab_views else "Warnings",
        "diagnostic_count": "诊断数" if lab_views else "Diagnostics",
        "candidate_count": "候选数" if lab_views else "Candidates",
        "parameter": "参数" if lab_views else "Parameter",
        "value": "值" if lab_views else "Value",
        "wavelength": "波长 (Å)" if lab_views else "Wavelength (Å)",
        "energy": "能量 (keV)" if lab_views else "Energy (keV)",
        "radiation": "预设光源" if lab_views else "Radiation source",
        "range": "有效 2θ 范围" if lab_views else "Effective 2θ range",
        "analysis_range": "计算 2θ 范围" if lab_views else "Analysis 2θ range",
        "profile_range": "谱线采样范围" if lab_views else "Profile sampled range",
        "profile": "线型模型" if lab_views else "Profile model",
        "input_mode": "输入方式" if lab_views else "Input mode",
        "phase_name": "物相名称" if lab_views else "Phase name",
        "formula": "化学式" if lab_views else "Formula",
        "space_group": "空间群" if lab_views else "Space group",
        "peak_count": "反射数" if lab_views else "Reflections",
        "elastic": "弹性数据状态" if lab_views else "Elastic data status",
        "no_phases": "无分析物相" if lab_views else "No analyzed phases",
        "no_data": "无数据" if lab_views else "No data",
        "sheet": "工作表" if lab_views else "Worksheet",
        "contents": "内容" if lab_views else "Contents",
    }

    values = _summary_values(summary)
    error_count = sum(str(row.get("level", "")).casefold() == "error" for row in diagnostics)
    warning_count = sum(
        str(row.get("level", "")).casefold() == "warning" for row in diagnostics
    )
    reflection_count = sum(
        count
        for phase in phases
        if isinstance((count := phase.get("reflection_count")), int)
        and not isinstance(count, bool)
        and count >= 0
    )
    if not any(
        isinstance(phase.get("reflection_count"), int)
        and not isinstance(phase.get("reflection_count"), bool)
        for phase in phases
    ):
        reported_reflections = values.get("reflection_count")
        if isinstance(reported_reflections, int) and not isinstance(reported_reflections, bool):
            reflection_count = reported_reflections

    worksheet = workbook.create_sheet(title=title)
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_view.zoomScale = 90

    worksheet.merge_cells("A1:E1")
    heading = worksheet["A1"]
    heading.value = _safe_text(labels["title"])
    heading.font = _TITLE_FONT
    heading.fill = _NAVY_FILL
    heading.alignment = _TITLE_ALIGNMENT
    worksheet.row_dimensions[1].height = 36

    row = 3

    def section(label: str) -> int:
        nonlocal row
        section_row = row
        worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        cell = worksheet.cell(row=row, column=1, value=_safe_text(label))
        cell.font = _SECTION_FONT
        cell.fill = _TEAL_FILL
        cell.alignment = _TITLE_ALIGNMENT
        worksheet.row_dimensions[row].height = 23
        row += 1
        return section_row

    section(labels["summary"])
    for column, value in enumerate(
        (labels["metric"], labels["count"], labels["metric"], labels["count"]),
        start=1,
    ):
        cell = worksheet.cell(row=row, column=column, value=_safe_text(value))
        cell.font = _HEADER_FONT
        cell.fill = _NAVY_FILL
        cell.alignment = _CENTER_ALIGNMENT
    row += 1

    candidate_count = values.get("candidate_count")
    metric_rows = (
        (labels["phase_count"], len(phases), labels["reflection_count"], reflection_count),
        (labels["error_count"], error_count, labels["warning_count"], warning_count),
        (
            labels["diagnostic_count"],
            len(diagnostics),
            labels["candidate_count"],
            candidate_count
            if isinstance(candidate_count, int) and not isinstance(candidate_count, bool)
            else labels["no_data"],
        ),
    )
    for index, metric_row in enumerate(metric_rows):
        for column, value in enumerate(metric_row, start=1):
            cell = worksheet.cell(row=row, column=column)
            cell.value = _safe_text(value) if isinstance(value, str) else value
            cell.alignment = _LEFT_ALIGNMENT if column % 2 else _CENTER_ALIGNMENT
            cell.fill = _PALE_BLUE_FILL if index % 2 == 0 else _WHITE_FILL
        row += 1

    row += 1
    section(labels["run"])
    worksheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
    for column, value in ((1, labels["parameter"]), (2, labels["value"])):
        cell = worksheet.cell(row=row, column=column, value=_safe_text(value))
        cell.font = _HEADER_FONT
        cell.fill = _NAVY_FILL
        cell.alignment = _CENTER_ALIGNMENT
    row += 1
    missing = labels["no_data"]
    wavelength = values.get("effective_wavelength_A")
    energy = values.get("effective_energy_keV")
    radiation = values.get("effective_radiation_source")
    requested_source = values.get("source_preset_applied")
    if radiation is None:
        radiation = requested_source
    run_rows = (
        (labels["wavelength"], wavelength, True),
        (labels["energy"], energy, True),
        (labels["radiation"], _display_radiation_source(radiation, lab_views=lab_views), False),
        (labels["range"], _summary_text(values.get("effective_two_theta_range_deg"), missing=missing), False),
        (labels["analysis_range"], _summary_text(values.get("analysis_two_theta_range_deg"), missing=missing), False),
        (labels["profile_range"], _summary_text(values.get("profile_sampled_two_theta_range_deg"), missing=missing), False),
        (labels["profile"], _display_profile_model(values.get("profile_model"), lab_views=lab_views, missing=missing), False),
        (
            labels["input_mode"],
            _display_input_mode(values.get("input_mode"), lab_views=lab_views, missing=missing),
            False,
        ),
    )
    run_value_width = sum((24, 23, 16, 25))
    for index, (parameter, value, numeric) in enumerate(run_rows):
        worksheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        label_cell = worksheet.cell(row=row, column=1, value=_safe_text(parameter))
        value_cell = worksheet.cell(row=row, column=2)
        if numeric:
            _write_numeric_or_text(value_cell, value, missing=missing)
        else:
            value_cell.value = _safe_text(value)
        label_cell.alignment = _TOP_WRAP_ALIGNMENT
        value_cell.alignment = _TOP_WRAP_ALIGNMENT
        fill = _PALE_BLUE_FILL if index % 2 == 0 else _WHITE_FILL
        label_cell.fill = fill
        value_cell.fill = fill
        worksheet.row_dimensions[row].height = _estimated_row_height(
            ((label_cell.value, 25), (value_cell.value, run_value_width))
        )
        row += 1

    row += 1
    section(labels["phases"])
    phase_headers = (
        labels["phase_name"],
        labels["formula"],
        labels["space_group"],
        labels["peak_count"],
        labels["elastic"],
    )
    for column, value in enumerate(phase_headers, start=1):
        cell = worksheet.cell(row=row, column=column, value=_safe_text(value))
        cell.font = _HEADER_FONT
        cell.fill = _NAVY_FILL
        cell.alignment = _CENTER_WRAP_ALIGNMENT
    worksheet.row_dimensions[row].height = 30
    row += 1
    if phases:
        for index, phase in enumerate(phases):
            symbol = phase.get("space_group_symbol")
            number = phase.get("space_group_number")
            if symbol not in (None, "") and number not in (None, ""):
                space_group = f"{_safe_text(symbol)} ({_safe_text(number)})"
            elif symbol not in (None, ""):
                space_group = _safe_text(symbol)
            elif number not in (None, ""):
                space_group = _safe_text(number)
            else:
                space_group = missing
            phase_values = (
                phase.get("phase_name"),
                phase.get("formula"),
                space_group,
                phase.get("reflection_count"),
                phase.get("elastic_status"),
            )
            for column, value in enumerate(phase_values, start=1):
                cell = worksheet.cell(row=row, column=column)
                if isinstance(value, str):
                    cell.value = _safe_text(value)
                elif value is None:
                    cell.value = missing
                else:
                    cell.value = value
                cell.alignment = _TOP_WRAP_ALIGNMENT
                if index % 2 == 1:
                    cell.fill = _PALE_BLUE_FILL
            worksheet.row_dimensions[row].height = _estimated_row_height(
                tuple(
                    (cell.value, width)
                    for cell, width in zip(
                        worksheet[row][:5], (25, 24, 23, 16, 25), strict=True
                    )
                )
            )
            row += 1
    else:
        worksheet.cell(row=row, column=1, value=_safe_text(labels["no_phases"]))
        row += 1

    row += 1
    section(labels["navigation"])
    worksheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
    for column, value in ((1, labels["sheet"]), (2, labels["contents"])):
        cell = worksheet.cell(row=row, column=column, value=_safe_text(value))
        cell.font = _HEADER_FONT
        cell.fill = _NAVY_FILL
        cell.alignment = _CENTER_ALIGNMENT
    row += 1
    descriptions = {
        "Summary": "导出设置与来源摘要" if lab_views else "Export settings and provenance summary",
        "Peaks": "完整反射数据表" if lab_views else "Full reflection table",
        "Phases": "物相结构与分析参数" if lab_views else "Phase structures and analysis parameters",
        "Elasticity": "弹性张量与来源信息" if lab_views else "Elastic tensors and provenance",
        "Candidates": "候选材料" if lab_views else "Candidate materials",
        "Downloads": "文件下载状态" if lab_views else "Download status",
        "Diagnostics": "错误与警告记录" if lab_views else "Errors and warnings",
        "Patterns": "衍射谱线采样数据" if lab_views else "Diffraction profile samples",
        "推荐峰表": "便于筛选的中文反射表" if lab_views else "Chinese recommended-reflection view",
        "使用说明": "字段与使用说明" if lab_views else "Field and usage guide",
    }
    for index, sheet_title in enumerate(workbook.sheetnames):
        if sheet_title == title:
            continue
        name_cell = worksheet.cell(row=row, column=1, value=_safe_text(sheet_title))
        name_cell.hyperlink = Hyperlink(
            ref=name_cell.coordinate,
            location=f"{quote_sheetname(sheet_title)}!A1",
        )
        name_cell.font = _LINK_FONT
        name_cell.alignment = _TOP_WRAP_ALIGNMENT
        description = descriptions.get(
            sheet_title,
            "各反射与该物相相关数据" if lab_views else "Reflection data for this phase",
        )
        worksheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        content_cell = worksheet.cell(row=row, column=2, value=_safe_text(description))
        content_cell.alignment = _TOP_WRAP_ALIGNMENT
        worksheet.row_dimensions[row].height = _estimated_row_height(
            ((name_cell.value, 25), (content_cell.value, sum((24, 23, 16, 25))))
        )
        if index % 2 == 1:
            for cell in worksheet[row][:2]:
                cell.fill = _PALE_BLUE_FILL
        row += 1

    for cells in worksheet.iter_rows(min_row=1, max_row=row - 1, min_col=1, max_col=5):
        for cell in cells:
            if cell.row > 1 and cell.value is not None:
                cell.border = _BOTTOM_BORDER
                if cell.alignment == _DEFAULT_ALIGNMENT:
                    cell.alignment = _TOP_WRAP_ALIGNMENT
            if cell.font.name != _FONT_NAME:
                cell.font = _BODY_FONT

    for column, width in {"A": 25, "B": 24, "C": 23, "D": 16, "E": 25}.items():
        worksheet.column_dimensions[column].width = width
    worksheet.freeze_panes = "A4"
    worksheet.print_area = f"A1:E{row - 1}"
    worksheet.print_title_rows = "1:1"
    worksheet.page_setup.orientation = worksheet.ORIENTATION_LANDSCAPE
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_margins.left = 0.3
    worksheet.page_margins.right = 0.3
    worksheet.page_margins.top = 0.5
    worksheet.page_margins.bottom = 0.5
    return worksheet
