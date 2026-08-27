"""Traceable CSV, Excel, JSON, and bundle-manifest exports."""

from __future__ import annotations

import csv
import json
import math
from dataclasses import asdict, fields, is_dataclass
import os
from pathlib import Path
import tempfile
from typing import Any

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .diffraction import SCIENTIFIC_BOUNDARY
from .export_views import (
    ANALYSIS_PEAK_COLUMNS,
    BEGINNER_PEAK_HEADERS_ZH,
    beginner_peak_rows_zh,
    safe_excel_sheet_title,
    user_guide_rows,
)
from .models import (
    AnalysisSettings,
    CandidateRecord,
    DiagnosticRecord,
    DiscoveryResult,
    DownloadArtifact,
    PhaseAnalysis,
)
from .utils import package_versions, runtime_environment, sha256_file, to_jsonable, utc_now_iso, write_json

CANDIDATE_HEADERS = [item.name for item in fields(CandidateRecord)]
DOWNLOAD_HEADERS = [
    "material_id",
    "formula",
    "cif_path",
    "cif_sha256",
    "elasticity_path",
    "elasticity_sha256",
    "status",
    "error",
    "elasticity_status",
    "elasticity_error",
    "source_provider",
    "source_url",
    "provider_metadata",
]
PHASE_HEADERS = [
    "phase_name",
    "cif_name",
    "cif_path",
    "cif_sha256",
    "data_block",
    "formula",
    "space_group_symbol",
    "space_group_number",
    "a_A",
    "b_A",
    "c_A",
    "alpha_deg",
    "beta_deg",
    "gamma_deg",
    "site_count_asymmetric",
    "site_count_unit_cell",
    "partial_occupancy",
    "reflection_count",
    "wavelength_A",
    "energy_keV",
    "wavelength_source",
    "formula_weight_g_mol",
    "density_g_cm3",
    "cell_volume_A3",
    "elastic_status",
    "elastic_source_provider",
    "elastic_source_record_id",
    "elastic_coordinate_frame",
    "warnings",
    "source_metadata",
]
# Analysis-first order: identity → hkl → geometry → intensity → ranks → SF extras → elastic → meta.
# All historical column names retained (additive reordering only).
PEAK_HEADERS = [
    # Identity (multi-phase filter)
    "phase_name",
    "cif_name",
    "cif_sha256",
    "formula",
    "space_group",
    # Miller
    "h",
    "k",
    "i",
    "l",
    "hkl",
    "family_label",
    "multiplicity",
    # Geometry
    "d_spacing_A",
    "theta_deg",
    "two_theta_deg",
    "two_theta_cu_ka_deg",
    "q_invA",
    "g_invA",
    # Display + LP channels
    "normalized_intensity",
    "rank_by_intensity",
    "intensity_with_lp",
    "intensity_no_lp",
    "lp_factor",
    # Volume-normalized J (+ legacy R_hkl aliases)
    "volume_normalized_intensity_with_lp",
    "volume_normalized_intensity_no_lp",
    "material_scattering_factor_R_hkl",
    "material_scattering_factor_R_hkl_no_lp",
    "inverse_R_hkl",
    "inverse_R_hkl_no_lp",
    "phase_relative_R_hkl_pct",
    "phase_relative_R_hkl_no_lp_pct",
    "rank_by_R_hkl",
    "rank_by_R_hkl_no_lp",
    # Structure-factor detail
    "structure_factor_sq",
    "mean_structure_factor_sq_per_multiplicity",
    "mean_structure_factor_abs_per_multiplicity",
    "sin_theta",
    "cos_theta",
    "sin_theta_over_lambda",
    "sin2_theta_over_lambda2",
    "is_multi_family_peak",
    "coincident_hkl_family_count",
    # Elastic + run meta
    "young_modulus_hkl_normal_GPa",
    "elastic_status",
    "elastic_note",
    "wavelength_A",
    "energy_keV",
    "formula_weight_g_mol",
    "density_g_cm3",
    "cell_volume_A3",
    "r_hkl_model_note",
    "scientific_boundary",
]
ELASTICITY_HEADERS = [
    "phase_name",
    "cif_name",
    "status",
    "unit",
    "source_provider",
    "source_record_id",
    "source_url",
    "methodology_url",
    "nature_of_data",
    "coordinate_frame",
    "raw_payload_path",
    "warnings",
    *[f"C{i}{j}_GPa" for i in range(1, 7) for j in range(1, 7)],
]
PATTERN_HEADERS = [
    "phase_name",
    "cif_name",
    "two_theta_deg",
    "d_A",
    "q_invA",
    "g_invA",
    "x_axis_mode",
    "x",
    "relative_intensity",
    "wavelength_A",
]
DIAGNOSTIC_HEADERS = ["stage", "item", "level", "message"]
SUMMARY_HEADERS = ["key", "value"]
EXCEL_DATA_ROW_LIMIT = 900_000


def _portable_path(value: Path, bundle_root: Path | None = None) -> str:
    """Return a stable POSIX path without leaking a transaction directory."""

    if not value.is_absolute():
        return value.as_posix()
    if bundle_root is not None:
        root = Path(os.path.abspath(os.fspath(bundle_root)))
        absolute = Path(os.path.abspath(os.fspath(value)))
        try:
            return absolute.relative_to(root).as_posix()
        except ValueError:
            pass
    return value.name


def _safe_spreadsheet_text(value: str) -> str:
    """Prevent externally sourced text from being interpreted as a spreadsheet formula."""

    if value and value[0] in {"=", "+", "-", "@", "\t", "\r", "\n"}:
        return "'" + value
    return value


def _cell_value(value: Any, bundle_root: Path | None = None) -> Any:
    if value is None:
        return ""
    if isinstance(value, Path):
        return _safe_spreadsheet_text(_portable_path(value, bundle_root))
    if isinstance(value, (np.floating, float)):
        resolved = float(value)
        return resolved if math.isfinite(resolved) else ""
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, bool):
        return value
    if isinstance(value, (dict, list, tuple, set, frozenset)):
        text = json.dumps(
            _portable_json_value(value, bundle_root),
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        )
        return _safe_spreadsheet_text(text)
    if isinstance(value, str):
        return _safe_spreadsheet_text(value)
    return value


def _portable_json_value(value: Any, bundle_root: Path | None = None) -> Any:
    """Normalize nested export values deterministically and JSON-safely."""

    if is_dataclass(value):
        return {
            field.name: _portable_json_value(getattr(value, field.name), bundle_root)
            for field in fields(value)
            if field.name
            not in {"small_structure", "structure_factor_structure", "space_group_object"}
        }
    if isinstance(value, Path):
        return _portable_path(value, bundle_root)
    if isinstance(value, dict):
        normalized: list[tuple[str, Any]] = []
        for key, item in value.items():
            portable_key = _portable_json_value(key, bundle_root)
            if not isinstance(portable_key, str):
                portable_key = json.dumps(
                    portable_key,
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
            normalized.append(
                (portable_key, _portable_json_value(item, bundle_root))
            )
        return {key: item for key, item in sorted(normalized, key=lambda pair: pair[0])}
    if isinstance(value, (list, tuple)):
        return [_portable_json_value(item, bundle_root) for item in value]
    if isinstance(value, (set, frozenset)):
        normalized = [_portable_json_value(item, bundle_root) for item in value]
        return sorted(
            normalized,
            key=lambda item: json.dumps(
                item,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ),
        )
    return to_jsonable(value)


def _write_text_atomic(path: Path, text: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=path.parent,
            prefix=f".{path.name}.",
            suffix=".tmp",
            delete=False,
        ) as handle:
            temporary = Path(handle.name)
            handle.write(text)
            handle.flush()
            os.fsync(handle.fileno())
        temporary.replace(path)
    finally:
        _unlink_temporary(temporary)
    return path


def _unlink_temporary(path: Path | None) -> None:
    """Best-effort cleanup that never masks an export result or its error."""

    if path is None:
        return
    try:
        path.unlink(missing_ok=True)
    except OSError:
        # A published result (or the primary write/replace exception) is more
        # important than a cleanup failure caused by another process or the
        # filesystem.  The leftover same-directory temp is recoverable.
        return


def _write_csv(
    path: Path,
    rows: list[dict[str, Any]],
    fieldnames: list[str],
    *,
    bundle_root: Path | None = None,
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            newline="",
            encoding="utf-8-sig",
            dir=path.parent,
            prefix=f".{path.name}.",
            suffix=".tmp",
            delete=False,
        ) as handle:
            temporary = Path(handle.name)
            writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                writer.writerow(
                    {
                        name: _cell_value(row.get(name), bundle_root)
                        for name in fieldnames
                    }
                )
            handle.flush()
            os.fsync(handle.fileno())
        temporary.replace(path)
    finally:
        _unlink_temporary(temporary)
    return path


def candidate_rows(discovery: DiscoveryResult | None) -> list[dict[str, Any]]:
    if discovery is None:
        return []
    return [asdict(candidate) for candidate in discovery.candidates]


def diagnostic_rows(diagnostics: list[DiagnosticRecord]) -> list[dict[str, Any]]:
    return [asdict(item) for item in diagnostics]


def download_rows(downloads: list[DownloadArtifact]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for item in downloads:
        rows.append(
            {
                "material_id": item.candidate.material_id,
                "formula": item.candidate.formula,
                "cif_path": item.cif_path,
                "cif_sha256": sha256_file(item.cif_path)
                if item.cif_path and item.cif_path.is_file()
                else "",
                "elasticity_path": item.elasticity_path,
                "elasticity_sha256": sha256_file(item.elasticity_path)
                if item.elasticity_path and item.elasticity_path.is_file()
                else "",
                "status": item.status,
                "error": item.error,
                "elasticity_status": item.elasticity_status,
                "elasticity_error": item.elasticity_error,
                "source_provider": item.candidate.source_provider,
                "source_url": item.candidate.source_url,
                "provider_metadata": item.provider_metadata,
            }
        )
    return rows


def phase_rows(analyses: list[PhaseAnalysis]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for analysis in analyses:
        structure = analysis.structure
        cell = structure.cell_parameters
        tensor = analysis.elastic_tensor
        meta = analysis.metadata
        rows.append(
            {
                "phase_name": analysis.phase_name,
                "cif_name": structure.cif_path.name,
                "cif_path": structure.cif_path,
                "cif_sha256": structure.cif_sha256,
                "data_block": structure.data_block,
                "formula": structure.formula,
                "space_group_symbol": structure.space_group_symbol,
                "space_group_number": structure.space_group_number,
                "a_A": cell[0],
                "b_A": cell[1],
                "c_A": cell[2],
                "alpha_deg": cell[3],
                "beta_deg": cell[4],
                "gamma_deg": cell[5],
                "site_count_asymmetric": structure.site_count_asymmetric,
                "site_count_unit_cell": structure.site_count_unit_cell,
                "partial_occupancy": structure.has_partial_occupancy,
                "reflection_count": len(analysis.reflections),
                "wavelength_A": analysis.wavelength_A,
                "energy_keV": analysis.energy_keV,
                "wavelength_source": analysis.wavelength_source,
                "formula_weight_g_mol": meta.get("formula_weight_g_mol"),
                "density_g_cm3": meta.get("density_g_cm3"),
                "cell_volume_A3": meta.get("cell_volume_A3"),
                "elastic_status": (
                    tensor.status
                    if tensor
                    else "not_requested"
                    if meta.get("elasticity_requested") is False
                    else "not_available"
                ),
                "elastic_source_provider": tensor.source_provider if tensor else "",
                "elastic_source_record_id": tensor.source_record_id if tensor else "",
                "elastic_coordinate_frame": tensor.coordinate_frame if tensor else "",
                "warnings": analysis.warnings,
                "source_metadata": structure.source_metadata,
            }
        )
    return rows


def peak_rows(analyses: list[PhaseAnalysis]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for analysis in analyses:
        structure = analysis.structure
        meta = analysis.metadata
        for reflection in analysis.reflections:
            if reflection.i is not None:
                hkl_label = f"({reflection.h} {reflection.k} {reflection.i} {reflection.l})"
            else:
                hkl_label = f"({reflection.h} {reflection.k} {reflection.l})"
            rows.append(
                {
                    "phase_name": analysis.phase_name,
                    "cif_name": structure.cif_path.name,
                    "cif_sha256": structure.cif_sha256,
                    "formula": structure.formula,
                    "space_group": structure.space_group_symbol,
                    "h": reflection.h,
                    "k": reflection.k,
                    "i": reflection.i,
                    "l": reflection.l,
                    "hkl": hkl_label,
                    "family_label": reflection.family_label,
                    "multiplicity": reflection.multiplicity,
                    "d_spacing_A": reflection.d_spacing_A,
                    "theta_deg": reflection.theta_deg,
                    "two_theta_deg": reflection.two_theta_deg,
                    "two_theta_cu_ka_deg": reflection.two_theta_cu_ka_deg,
                    "q_invA": reflection.q_invA,
                    "g_invA": reflection.g_invA,
                    "sin_theta": reflection.sin_theta,
                    "cos_theta": reflection.cos_theta,
                    "sin_theta_over_lambda": reflection.sin_theta_over_lambda,
                    "sin2_theta_over_lambda2": reflection.sin2_theta_over_lambda2,
                    "structure_factor_sq": reflection.structure_factor_sq,
                    "mean_structure_factor_sq_per_multiplicity": (
                        reflection.mean_structure_factor_sq_per_multiplicity
                    ),
                    "mean_structure_factor_abs_per_multiplicity": (
                        reflection.mean_structure_factor_abs_per_multiplicity
                    ),
                    "intensity_no_lp": reflection.intensity_no_lp,
                    "lp_factor": reflection.lp_factor,
                    "intensity_with_lp": reflection.intensity_with_lp,
                    "normalized_intensity": reflection.normalized_intensity,
                    "volume_normalized_intensity_with_lp": reflection.material_scattering_factor_R_hkl,
                    "volume_normalized_intensity_no_lp": reflection.material_scattering_factor_R_hkl_no_lp,
                    "material_scattering_factor_R_hkl": reflection.material_scattering_factor_R_hkl,
                    "material_scattering_factor_R_hkl_no_lp": reflection.material_scattering_factor_R_hkl_no_lp,
                    "inverse_R_hkl": reflection.inverse_R_hkl,
                    "inverse_R_hkl_no_lp": reflection.inverse_R_hkl_no_lp,
                    "phase_relative_R_hkl_pct": reflection.phase_relative_R_hkl_pct,
                    "phase_relative_R_hkl_no_lp_pct": reflection.phase_relative_R_hkl_no_lp_pct,
                    "rank_by_intensity": reflection.rank_by_intensity,
                    "rank_by_R_hkl": reflection.rank_by_R_hkl,
                    "rank_by_R_hkl_no_lp": reflection.rank_by_R_hkl_no_lp,
                    "is_multi_family_peak": reflection.is_multi_family_peak,
                    "coincident_hkl_family_count": reflection.coincident_hkl_family_count,
                    "young_modulus_hkl_normal_GPa": reflection.young_modulus_hkl_normal_GPa,
                    "elastic_status": reflection.elastic_status,
                    "elastic_note": reflection.elastic_note,
                    "wavelength_A": analysis.wavelength_A,
                    "energy_keV": analysis.energy_keV,
                    "formula_weight_g_mol": meta.get("formula_weight_g_mol"),
                    "density_g_cm3": meta.get("density_g_cm3"),
                    "cell_volume_A3": meta.get("cell_volume_A3"),
                    "r_hkl_model_note": reflection.r_hkl_model_note,
                    "scientific_boundary": SCIENTIFIC_BOUNDARY,
                }
            )
    return rows


def elasticity_rows(analyses: list[PhaseAnalysis]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for analysis in analyses:
        tensor = analysis.elastic_tensor
        if tensor is None:
            continue
        row: dict[str, Any] = {
            "phase_name": analysis.phase_name,
            "cif_name": analysis.structure.cif_path.name,
            "status": tensor.status,
            "unit": tensor.unit,
            "source_provider": tensor.source_provider,
            "source_record_id": tensor.source_record_id,
            "source_url": tensor.source_url,
            "methodology_url": tensor.methodology_url,
            "nature_of_data": tensor.nature_of_data,
            "coordinate_frame": tensor.coordinate_frame,
            "raw_payload_path": tensor.raw_payload_path,
            "warnings": tensor.warnings,
        }
        if tensor.stiffness_GPa.shape == (6, 6):
            for i in range(6):
                for j in range(6):
                    row[f"C{i + 1}{j + 1}_GPa"] = float(tensor.stiffness_GPa[i, j])
        rows.append(row)
    return rows


def _pattern_axis_coordinates(
    two_theta_deg: float,
    wavelength_A: float,
    x_axis_mode: str,
) -> tuple[float | None, float | None, float | None, float | None]:
    """Return d_A, q_invA, g_invA, and selected x for a profile sample."""

    theta_rad = math.radians(float(two_theta_deg) / 2.0)
    sin_theta = math.sin(theta_rad)
    wavelength_valid = math.isfinite(wavelength_A) and wavelength_A > 0
    if not math.isfinite(sin_theta) or sin_theta < 0 or not wavelength_valid:
        d_A = None
        q_invA = None
        g_invA = None
    elif sin_theta == 0.0:
        # d is singular at 2theta=0, but reciprocal coordinates are finite
        # physical zeros and must remain numeric in CSV/Excel profiles.
        d_A = None
        q_invA = 0.0
        g_invA = 0.0
    else:
        d_A = float(wavelength_A) / (2.0 * sin_theta)
        if not math.isfinite(d_A) or d_A <= 0:
            d_A = None
        if d_A is None:
            q_invA = None
            g_invA = None
        else:
            # Direct sin(theta) expressions avoid deriving a finite reciprocal
            # coordinate through the intentionally blank d=inf endpoint.
            q_invA = float(4.0 * math.pi * sin_theta / float(wavelength_A))
            g_invA = float(2.0 * sin_theta / float(wavelength_A))
            if not math.isfinite(q_invA):
                q_invA = None
            if not math.isfinite(g_invA):
                g_invA = None
    mode = str(x_axis_mode or "two_theta")
    if mode == "d_spacing":
        x_value = d_A
    elif mode == "q":
        x_value = q_invA
    elif mode == "g":
        x_value = g_invA
    else:
        x_value = float(two_theta_deg)
    return d_A, q_invA, g_invA, x_value


def pattern_rows(analyses: list[PhaseAnalysis]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for analysis in analyses:
        x_axis_mode = str(analysis.metadata.get("pattern_axis") or "two_theta")
        wavelength = float(analysis.wavelength_A)
        for angle, intensity in zip(
            analysis.two_theta_grid,
            analysis.intensity_profile,
            strict=True,
        ):
            two_theta = float(angle)
            d_A, q_invA, g_invA, x_value = _pattern_axis_coordinates(
                two_theta, wavelength, x_axis_mode
            )
            rows.append(
                {
                    "phase_name": analysis.phase_name,
                    "cif_name": analysis.structure.cif_path.name,
                    "two_theta_deg": two_theta,
                    "d_A": d_A,
                    "q_invA": q_invA,
                    "g_invA": g_invA,
                    "x_axis_mode": x_axis_mode,
                    "x": x_value,
                    "relative_intensity": float(intensity),
                    "wavelength_A": analysis.wavelength_A,
                }
            )
    return rows


def _add_sheet(
    workbook: Workbook,
    title: str,
    rows: list[dict[str, Any]],
    headers: list[str],
    *,
    bundle_root: Path | None = None,
) -> None:
    sheet = workbook.create_sheet(title=title)
    sheet.append(headers)
    if len(rows) > EXCEL_DATA_ROW_LIMIT:
        note: dict[str, Any] = {headers[0]: "omitted_from_workbook"}
        if len(headers) > 1:
            note[headers[1]] = (
                f"{len(rows):,} data rows exceed the conservative Excel limit of "
                f"{EXCEL_DATA_ROW_LIMIT:,}; use the corresponding CSV file."
            )
        rows = [note]
    if rows:
        for row in rows:
            sheet.append(
                [_cell_value(row.get(header), bundle_root) for header in headers]
            )
    else:
        sheet.append(["no rows", *([""] * (len(headers) - 1))])
    # Frozen header + autofilter: required for analysis-ready long tables.
    sheet.freeze_panes = "A2"
    last_col = get_column_letter(max(len(headers), 1))
    last_row = max(sheet.max_row, 1)
    sheet.auto_filter.ref = f"A1:{last_col}{last_row}"
    header_fill = PatternFill("solid", fgColor="16324F")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column_index, header in enumerate(headers, start=1):
        sampled = [str(header)] + [
            str(_cell_value(row.get(header), bundle_root)) for row in rows[:200]
        ]
        # Slightly wider for analysis headers with units in the name.
        width = min(max(max(len(value) for value in sampled) + 2, 10), 48)
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _add_guide_sheet(workbook: Workbook, title: str, rows: list[list[str]]) -> None:
    sheet = workbook.create_sheet(title=title)
    if not rows:
        sheet.append(["no rows", ""])
    else:
        for row in rows:
            sheet.append([_cell_value(cell) for cell in row])
    sheet.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="16324F")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column_index in range(1, 3):
        sampled = [
            str(_cell_value(row[column_index - 1]) if column_index - 1 < len(row) else "")
            for row in rows[:200]
        ] or [""]
        width = min(max(max(len(value) for value in sampled) + 2, 12), 72)
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_excel_workbook(
    path: Path,
    *,
    summary: list[dict[str, Any]],
    phases: list[dict[str, Any]],
    peaks: list[dict[str, Any]],
    elasticity: list[dict[str, Any]],
    candidates: list[dict[str, Any]],
    downloads: list[dict[str, Any]],
    diagnostics: list[dict[str, Any]],
    patterns: list[dict[str, Any]],
    analyses: list[PhaseAnalysis] | None = None,
    export_lab_views: bool = False,
    include_patterns: bool = True,
    bundle_root: Path | None = None,
) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    # Lab-first when enabled: guide → Chinese long table → full Peaks, then metadata sheets.
    if export_lab_views:
        _add_guide_sheet(workbook, "使用说明", user_guide_rows())
        zh_headers = list(BEGINNER_PEAK_HEADERS_ZH.keys())
        _add_sheet(workbook, "推荐峰表", beginner_peak_rows_zh(peaks), zh_headers)
    _add_sheet(workbook, "Summary", summary, SUMMARY_HEADERS, bundle_root=bundle_root)
    _add_sheet(workbook, "Peaks", peaks, PEAK_HEADERS, bundle_root=bundle_root)
    _add_sheet(workbook, "Phases", phases, PHASE_HEADERS, bundle_root=bundle_root)
    _add_sheet(
        workbook, "Elasticity", elasticity, ELASTICITY_HEADERS, bundle_root=bundle_root
    )
    _add_sheet(
        workbook, "Candidates", candidates, CANDIDATE_HEADERS, bundle_root=bundle_root
    )
    _add_sheet(
        workbook, "Downloads", downloads, DOWNLOAD_HEADERS, bundle_root=bundle_root
    )
    _add_sheet(
        workbook, "Diagnostics", diagnostics, DIAGNOSTIC_HEADERS, bundle_root=bundle_root
    )
    if include_patterns:
        _add_sheet(workbook, "Patterns", patterns, PATTERN_HEADERS, bundle_root=bundle_root)
    if export_lab_views:
        analysis_list = analyses or []
        if 0 < len(analysis_list) <= 20:
            used_titles: set[str] = set(workbook.sheetnames)
            for analysis in analysis_list:
                if not analysis.reflections:
                    continue
                phase_peaks = [row for row in peaks if row.get("phase_name") == analysis.phase_name]
                if not phase_peaks:
                    continue
                title = safe_excel_sheet_title(f"峰_{analysis.phase_name}", used=used_titles)
                _add_sheet(
                    workbook,
                    title,
                    phase_peaks,
                    PEAK_HEADERS,
                    bundle_root=bundle_root,
                )
        # Open on the Chinese analysis long table when present.
        if "推荐峰表" in workbook.sheetnames:
            workbook.active = workbook["推荐峰表"]
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{path.name}.",
            suffix=".tmp",
            dir=path.parent,
            delete=False,
        ) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        temporary.replace(path)
    finally:
        _unlink_temporary(temporary)
    return path


def _summary_rows(
    analyses: list[PhaseAnalysis],
    discovery: DiscoveryResult | None,
    settings: AnalysisSettings,
    diagnostics: list[DiagnosticRecord],
    *,
    include_excel: bool = True,
) -> list[dict[str, Any]]:
    first_analysis = analyses[0] if analyses else None
    first_metadata = first_analysis.metadata if first_analysis is not None else {}
    requested_range = [settings.two_theta_min_deg, settings.two_theta_max_deg]
    grid_range = first_metadata.get("two_theta_range_deg", requested_range)
    if not isinstance(grid_range, (list, tuple)) or len(grid_range) != 2:
        grid_range = requested_range
    sampled_range = first_metadata.get("profile_sampled_two_theta_range_deg")
    if sampled_range is not None and (
        not isinstance(sampled_range, (list, tuple)) or len(sampled_range) != 2
    ):
        sampled_range = None
    effective_range = first_metadata.get("effective_two_theta_range_deg")
    if first_analysis is None and effective_range is None:
        # No phase was analyzed, so there is no wavelength-specific filter
        # intersection to report; retain the historical summary fallback.
        effective_range = requested_range
    if effective_range is not None and (
        not isinstance(effective_range, (list, tuple)) or len(effective_range) != 2
    ):
        effective_range = None
    effective_window_empty = bool(first_metadata.get("effective_window_empty", False))
    geometric_d_min = first_metadata.get(
        "geometric_d_min_A", first_metadata.get("d_min_A")
    )
    effective_export_lab_views = bool(include_excel and settings.export_lab_views)
    effective_include_figures = bool(settings.include_figures and analyses)
    effective_source = (
        first_analysis.wavelength_source if first_analysis is not None else None
    )
    effective_wavelength = (
        first_analysis.wavelength_A if first_analysis is not None else None
    )
    effective_energy = first_analysis.energy_keV if first_analysis is not None else None
    return [
        {"key": "generated_at_utc", "value": utc_now_iso()},
        {"key": "phase_count", "value": len(analyses)},
        {"key": "reflection_count", "value": sum(len(item.reflections) for item in analyses)},
        {"key": "candidate_count", "value": len(discovery.candidates) if discovery else 0},
        {"key": "diagnostic_count", "value": len(diagnostics)},
        {"key": "error_count", "value": sum(item.level == "error" for item in diagnostics)},
        {"key": "input_mode", "value": settings.input_mode},
        {"key": "source_preset", "value": settings.source_preset},
        # Preserve the established Summary meaning for downstream readers.
        {"key": "two_theta_range_deg", "value": requested_range},
        {"key": "requested_two_theta_range_deg", "value": requested_range},
        {"key": "analysis_two_theta_range_deg", "value": list(grid_range)},
        {
            "key": "profile_sampled_two_theta_range_deg",
            "value": list(sampled_range) if sampled_range is not None else None,
        },
        {"key": "effective_two_theta_range_deg", "value": (
            list(effective_range) if effective_range is not None else None
        )},
        {"key": "effective_window_empty", "value": effective_window_empty},
        {"key": "effective_wavelength_A", "value": effective_wavelength},
        {"key": "effective_energy_keV", "value": effective_energy},
        {"key": "effective_radiation_source", "value": effective_source},
        {
            "key": "source_preset_applied",
            "value": (
                settings.source_preset
                if first_analysis is not None and settings.input_mode == "source"
                else None
            ),
        },
        {"key": "profile_model", "value": settings.profile_model},
        {"key": "geometric_d_min_A", "value": geometric_d_min},
        {"key": "d_min_A", "value": settings.d_min_A},
        {"key": "d_max_A", "value": settings.d_max_A},
        {
            "key": "filter_d_min_A",
            "value": first_metadata.get("filter_d_min_A", settings.d_min_A),
        },
        {
            "key": "filter_d_max_A",
            "value": first_metadata.get("filter_d_max_A", settings.d_max_A),
        },
        {"key": "pattern_axis", "value": settings.pattern_axis},
        {"key": "include_excel", "value": bool(include_excel)},
        {"key": "include_patterns", "value": bool(settings.include_patterns)},
        {"key": "include_figures", "value": bool(settings.include_figures)},
        {"key": "effective_include_figures", "value": effective_include_figures},
        {"key": "export_lab_views", "value": bool(settings.export_lab_views)},
        {"key": "effective_export_lab_views", "value": effective_export_lab_views},
        {
            "key": "peak_longtable_note",
            "value": (
                "Peaks sheet: analysis-first column order; freeze header + autofilter; "
                "filter phase_name for multi-phase. Lab views add 推荐峰表 + 使用说明."
            ),
        },
        {
            "key": "analysis_peak_columns",
            "value": list(ANALYSIS_PEAK_COLUMNS),
        },
        {"key": "scientific_boundary", "value": SCIENTIFIC_BOUNDARY},
    ]


def _bundle_readme(
    analyses: list[PhaseAnalysis],
    discovery: DiscoveryResult | None,
    diagnostics: list[DiagnosticRecord],
    settings: AnalysisSettings | None = None,
    *,
    include_excel: bool = True,
) -> str:
    settings = settings or AnalysisSettings()
    effective_export_lab_views = bool(include_excel and settings.export_lab_views)
    effective_include_figures = bool(settings.include_figures and analyses)
    contents = [
        "# DiffractScout result bundle",
        "",
        "This directory is a self-contained, provenance-oriented output from DiffractScout.",
        "",
        "## Contents",
        "",
        "- `phase_summary.csv`: structure identity, unit cell, validation warnings, and source metadata.",
        "- `peak_reference.csv`: indexed theoretical powder reflections and optional hkl-normal Young's modulus.",
    ]
    if settings.include_patterns:
        contents.append(
            f"- `pattern_profiles.csv`: normalized {settings.profile_model} display profiles; not an instrument model."
        )
    else:
        contents.append(
            f"- Selected profile model: `{settings.profile_model}` (continuous pattern export is disabled)."
        )
    contents.extend(
        [
            "- `elasticity.csv`: paired 6x6 tensors and explicit provenance.",
            "- `candidate_index.csv`: source-database search results when discovery was used.",
            "- `download_index.csv`: source download and elasticity-query outcomes with hashes.",
            "- `diagnostics.csv`: structured warnings and failures, including phases that could not be analyzed.",
        ]
    )
    if include_excel:
        contents.append("- `results.xlsx`: human-readable workbook containing the same tables.")
    contents.extend(
        [
            "- `provenance.json`: settings, definitions, source metadata, and scientific boundaries.",
            "- `manifest.json`: SHA-256 inventory used by `diffractscout verify`.",
            "",
            "## Scientific boundary",
            "",
            SCIENTIFIC_BOUNDARY,
            "",
            "Materials Project Cij is labeled as DFT-derived. Raw/POSCAR and IEEE tensor records are retained with `frame_transform_required` until an explicit verified transform to the emitted CIF Cartesian frame is available. Missing tensors remain missing.",
            "",
            f"Analyzed phases: {len(analyses)}.",
            f"Diagnostics: {len(diagnostics)} ({sum(item.level == 'error' for item in diagnostics)} errors).",
        ]
    )
    if include_excel:
        if effective_export_lab_views:
            contents.append(
                "Excel lab views are enabled: `results.xlsx` includes `使用说明`, `推荐峰表`, "
                "and eligible per-phase convenience sheets."
            )
    if effective_include_figures:
        contents.append("2theta figure files are emitted under `figures/`.")
    if discovery is not None:
        contents.append(
            f"Discovered candidates: {len(discovery.candidates)} across {len(discovery.subsystems)} queried subsystems."
        )
    return "\n".join(contents) + "\n"


def export_result_bundle(
    output_dir: str | Path,
    *,
    analyses: list[PhaseAnalysis],
    settings: AnalysisSettings,
    discovery: DiscoveryResult | None = None,
    downloads: list[DownloadArtifact] | None = None,
    diagnostics: list[DiagnosticRecord] | None = None,
    include_excel: bool = True,
) -> Path:
    output = Path(output_dir).expanduser().resolve()
    output.mkdir(parents=True, exist_ok=True)
    if downloads is None:
        downloads = []
    if diagnostics is None:
        diagnostics = []

    candidates = candidate_rows(discovery)
    download_table = download_rows(downloads)
    phases = phase_rows(analyses)
    peaks = peak_rows(analyses)
    elasticity = elasticity_rows(analyses)
    patterns = pattern_rows(analyses) if settings.include_patterns else []
    if include_excel:
        workbook_tables = {
            "Phases": phases,
            "Peaks": peaks,
            "Elasticity": elasticity,
            "Candidates": candidates,
            "Downloads": download_table,
        }
        if settings.include_patterns:
            workbook_tables["Patterns"] = patterns
        for title, rows in workbook_tables.items():
            if len(rows) > EXCEL_DATA_ROW_LIMIT:
                diagnostics.append(
                    DiagnosticRecord(
                        "export",
                        title,
                        "warning",
                        f"{len(rows):,} rows exceed the conservative Excel data-row limit "
                        f"of {EXCEL_DATA_ROW_LIMIT:,}; the workbook contains an omission note and "
                        "the complete data remain in CSV.",
                    )
                )
    diagnostic_table = diagnostic_rows(diagnostics)
    summary = _summary_rows(
        analyses,
        discovery,
        settings,
        diagnostics,
        include_excel=include_excel,
    )

    _write_csv(output / "phase_summary.csv", phases, PHASE_HEADERS, bundle_root=output)
    _write_csv(output / "peak_reference.csv", peaks, PEAK_HEADERS, bundle_root=output)
    if settings.include_patterns:
        _write_csv(
            output / "pattern_profiles.csv",
            patterns,
            PATTERN_HEADERS,
            bundle_root=output,
        )
    _write_csv(output / "elasticity.csv", elasticity, ELASTICITY_HEADERS, bundle_root=output)
    _write_csv(output / "candidate_index.csv", candidates, CANDIDATE_HEADERS, bundle_root=output)
    _write_csv(
        output / "download_index.csv",
        download_table,
        DOWNLOAD_HEADERS,
        bundle_root=output,
    )
    _write_csv(
        output / "diagnostics.csv",
        diagnostic_table,
        DIAGNOSTIC_HEADERS,
        bundle_root=output,
    )
    _write_text_atomic(
        output / "README.md",
        _bundle_readme(
            analyses,
            discovery,
            diagnostics,
            settings,
            include_excel=include_excel,
        ),
    )

    effective_include_figures = bool(settings.include_figures and analyses)
    effective_export_lab_views = bool(include_excel and settings.export_lab_views)
    provenance = {
        "schema": "diffractscout_provenance_v1",
        "generated_at_utc": utc_now_iso(),
        "analysis_settings": asdict(settings),
        "output_flags": {
            "include_excel": bool(include_excel),
            "include_patterns": bool(settings.include_patterns),
            "include_figures": effective_include_figures,
            "export_lab_views": effective_export_lab_views,
            "effective_export_lab_views": effective_export_lab_views,
        },
        "discovery": (
            _portable_json_value(discovery, output) if discovery is not None else None
        ),
        "downloads": [_portable_json_value(row, output) for row in download_table],
        "diagnostics": _portable_json_value(diagnostics, output),
        "phase_metadata": [
            {
                "phase_name": analysis.phase_name,
                "structure": {
                    "cif_name": analysis.structure.cif_path.name,
                    "cif_sha256": analysis.structure.cif_sha256,
                    "formula": analysis.structure.formula,
                    "space_group": analysis.structure.space_group_symbol,
                    "source_metadata": _portable_json_value(
                        analysis.structure.source_metadata, output
                    ),
                },
                "analysis_metadata": _portable_json_value(analysis.metadata, output),
                "warnings": analysis.warnings,
            }
            for analysis in analyses
        ],
        "definitions": {
            "q": "2*pi/d = 4*pi*sin(theta)/lambda",
            "intensity_no_lp": "multiplicity * |F_xray|^2",
            "intensity_with_lp": "intensity_no_lp * Lorentz-polarization factor",
            "volume_normalized_intensity_with_lp": "intensity_with_lp / unit_cell_volume^2",
            "volume_normalized_intensity_no_lp": "intensity_no_lp / unit_cell_volume^2",
            "legacy_R_hkl_aliases": (
                "material_scattering_factor_R_hkl and material_scattering_factor_R_hkl_no_lp "
                "are compatibility aliases for the two volume-normalized theoretical intensity channels; "
                "they are not crystallographic residual factors or standardized QPA coefficients"
            ),
            "formula_weight_g_mol": (
                "Expanded unit-cell mass in g/mol, calculated from all occupied sites in the "
                "crystallographic unit cell (including Z); it is not the empirical formula mass."
            ),
            "elastic_modulus": "E(n) = 1 / (q(n)^T S q(n)) under engineering-shear Voigt convention",
            "lab_views_schema": (
                "When effective_export_lab_views is true, the emitted results.xlsx opens on 推荐峰表 after 使用说明: "
                "analysis-first Chinese long table (BEGINNER_PEAK_HEADERS_ZH) with freeze+autofilter, "
                "optional per-phase peak sheets (≤20 phases). Peaks/CSV remain the English machine schema "
                "with the same peak_rows values; R_hkl names are volume-normalized J aliases, not residuals."
            ),
            "analysis_peak_columns": list(ANALYSIS_PEAK_COLUMNS),
            "pattern_axis_columns": (
                "When pattern_profiles.csv is emitted (include_patterns is true), it includes "
                "two_theta_deg, d_A, q_invA, g_invA, "
                "x_axis_mode (settings.pattern_axis), x (selected axis value), and relative_intensity"
            ),
            "pattern_axis_sampling": (
                "Pattern profiles are sampled on a uniform 2theta grid and transformed to the "
                "selected d, q, or g coordinate; they are not uniformly resampled in q or d, "
                "and no Jacobian is applied."
            ),
            "two_theta_ranges": (
                "In phase metadata, two_theta_range_deg is the configured analysis bound; "
                "profile_sampled_two_theta_range_deg records the first and last emitted profile "
                "coordinates. In the Excel Summary, legacy two_theta_range_deg and "
                "requested_two_theta_range_deg are the user request, while "
                "analysis_two_theta_range_deg records the configured per-phase bounds. "
                "effective_two_theta_range_deg is the inclusive intersection with the d-spacing "
                "filter, or null when that intersection is empty. geometric_d_min_A is the Bragg "
                "lower d bound from the configured analysis upper angle and is distinct from "
                "filter d_min_A."
            ),
            "conditional_outputs": (
                "results.xlsx is emitted when include_excel is true; pattern_profiles.csv is "
                "emitted when include_patterns is true; figures/ is emitted when include_figures "
                "is true and at least one phase is analyzed."
            ),
        },
        "scientific_boundary": SCIENTIFIC_BOUNDARY,
        "excel_data_row_limit": EXCEL_DATA_ROW_LIMIT,
        "software_versions": package_versions(),
        "runtime_environment": runtime_environment(),
    }
    write_json(output / "provenance.json", provenance)

    if include_excel:
        write_excel_workbook(
            output / "results.xlsx",
            summary=summary,
            phases=phases,
            peaks=peaks,
            elasticity=elasticity,
            candidates=candidates,
            downloads=download_table,
            diagnostics=diagnostic_table,
            patterns=patterns,
            analyses=analyses,
            export_lab_views=bool(settings.export_lab_views),
            include_patterns=bool(settings.include_patterns),
            bundle_root=output,
        )

    if settings.include_figures and analyses:
        from .plotting import export_phase_figures

        figures_dir = output / "figures"
        figure_preset = settings.figure_preset or "publication"
        multi = len(analyses) > 1
        for index, analysis in enumerate(analyses, start=1):
            export_phase_figures(
                analysis,
                figures_dir,
                preset=figure_preset,
                formats=("svg", "png"),
                index=index if multi else None,
            )

    files: list[dict[str, Any]] = []
    for path in sorted(output.rglob("*")):
        if path.is_symlink():
            raise ValueError(f"Result bundles cannot contain symbolic links: {path}")
        if not path.is_file() or path == output / "manifest.json":
            continue
        relative = path.relative_to(output).as_posix()
        suffix = path.suffix.lower()
        if relative.startswith("inputs/"):
            role = "source_input"
        elif relative.startswith("figures/") or suffix in {".svg", ".png", ".pdf", ".eps", ".tif", ".tiff"}:
            role = "figure"
        elif suffix in {".csv", ".xlsx"}:
            role = "tabular_result"
        elif suffix == ".json":
            role = "provenance"
        else:
            role = "documentation"
        files.append(
            {
                "path": relative,
                "sha256": sha256_file(path),
                "size_bytes": path.stat().st_size,
                "role": role,
            }
        )
    manifest = {
        "schema": "diffractscout_bundle_manifest_v1",
        "generated_at_utc": utc_now_iso(),
        "summary": {
            "phase_count": len(analyses),
            "reflection_count": len(peaks),
            "candidate_count": len(candidates),
            "download_count": len(downloads),
            "diagnostic_count": len(diagnostics),
            "error_count": sum(item.level == "error" for item in diagnostics),
        },
        "files": files,
    }
    return write_json(output / "manifest.json", manifest)
