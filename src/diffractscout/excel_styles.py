"""Reusable, presentation-only styles for exported Excel workbooks."""

from __future__ import annotations

from math import ceil
import re
import unicodedata
from typing import Mapping, Sequence

from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .export_views import BEGINNER_PEAK_HEADERS_ZH


_FONT_NAME = "Microsoft YaHei"
_BODY_FONT = Font(name=_FONT_NAME, size=10, color="243447")
_HEADER_FONT = Font(name=_FONT_NAME, size=10, bold=True, color="FFFFFF")
_SECTION_FONT = Font(name=_FONT_NAME, size=11, bold=True, color="16324F")
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
_BODY_ALIGNMENT = Alignment(vertical="center")
_WRAP_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
_HEADER_BOTTOM = Border(bottom=Side(style="medium", color="FFFFFF"))

_GROUP_COLORS = {
    "identity": "24476B",
    "indices": "52628A",
    "geometry": "28677A",
    "intensity": "806017",
    "elasticity": "32664F",
    "quality": "714A79",
    "metadata": "596779",
}
_GROUP_FILLS = {
    name: PatternFill("solid", fgColor=color) for name, color in _GROUP_COLORS.items()
}
_BAND_FILL = PatternFill("solid", fgColor="F1F5F9")
_SECTION_FILL = PatternFill("solid", fgColor="E4ECF4")
_OVERVIEW_HEADER_FILL = PatternFill("solid", fgColor="24476B")
_ZH_TO_CANONICAL = {label: canonical for label, canonical in BEGINNER_PEAK_HEADERS_ZH.items()}

_IDENTITY_HEADERS = {
    "phase_name",
    "cif_name",
    "formula",
    "space_group",
    "space_group_symbol",
    "material_id",
    "物相名称",
    "CIF文件",
    "化学式",
    "空间群",
}
_INDEX_HEADERS = {"h", "k", "i", "l", "hkl", "family_label", "multiplicity"}
_INDEX_LABELS = {"晶面指标", "晶面族", "多重度"}
_GEOMETRY_HEADERS = {
    "d_spacing_A",
    "d_A",
    "theta_deg",
    "two_theta_deg",
    "two_theta_cu_ka_deg",
    "q_invA",
    "g_invA",
    "a_A",
    "b_A",
    "c_A",
    "alpha_deg",
    "beta_deg",
    "gamma_deg",
    "cell_volume_A3",
    "wavelength_A",
    "energy_keV",
    "d間距_Å",
    "θ_deg",
    "2θ_当前_deg",
    "2θ_CuKa_deg",
    "q_1/Å",
    "g_1/Å",
    "波长_Å",
    "能量_keV",
    "晶胞体积_Å3",
}
_INTENSITY_HEADERS = {
    "normalized_intensity",
    "relative_intensity",
    "intensity_with_lp",
    "intensity_no_lp",
    "lp_factor",
    "volume_normalized_intensity_with_lp",
    "volume_normalized_intensity_no_lp",
    "material_scattering_factor_R_hkl",
    "material_scattering_factor_R_hkl_no_lp",
    "inverse_R_hkl",
    "inverse_R_hkl_no_lp",
    "phase_relative_R_hkl_pct",
    "phase_relative_R_hkl_no_lp_pct",
    "rank_by_intensity",
    "rank_by_R_hkl",
    "rank_by_R_hkl_no_lp",
    "structure_factor_sq",
    "mean_structure_factor_sq_per_multiplicity",
    "mean_structure_factor_abs_per_multiplicity",
    "relative_intensity",
    "相对强度_相内max100",
    "强度排序_相内",
    "强度_含LP",
    "强度_无LP",
    "LP因子",
    "体积归一强度J_含LP_R_hkl别名",
    "体积归一强度J_无LP",
    "1/J_含LP",
    "1/J_无LP",
    "相内相对J_含LP_%",
    "相内相对J_无LP_%",
    "J_含LP排序",
    "J_无LP排序",
}
_ELASTICITY_HEADERS = {
    "young_modulus_hkl_normal_GPa",
    "elastic_status",
    "elastic_note",
    "杨氏模量_hkl法向_GPa",
    "弹性状态",
    "弹性备注",
}
_QUALITY_HEADERS = {
    "status",
    "level",
    "error",
    "elasticity_status",
    "elasticity_error",
    "message",
    "warnings",
    "is_multi_family_peak",
    "coincident_hkl_family_count",
    "是否多族共2θ",
    "共位hkl族数",
}

# These columns remain in the workbook and retain their full values. They start
# collapsed so the main scientific columns fit on screen; each can be expanded.
_DETAIL_HEADERS = {
    "cif_sha256",
    "elasticity_sha256",
    "data_block",
    "source_url",
    "methodology_url",
    "raw_payload_path",
    "source_metadata",
    "provider_metadata",
    "scientific_boundary",
    "r_hkl_model_note",
}
_PEAK_DETAIL_HEADERS = {
    "structure_factor_sq",
    "mean_structure_factor_sq_per_multiplicity",
    "mean_structure_factor_abs_per_multiplicity",
    "sin_theta",
    "cos_theta",
    "sin_theta_over_lambda",
    "sin2_theta_over_lambda2",
    "elastic_note",
    "formula_weight_g_mol",
    "density_g_cm3",
    "cell_volume_A3",
    "wavelength_A",
    "energy_keV",
    "wavelength_source",
}
_LONG_TEXT_HEADERS = {
    "cif_path",
    "elasticity_path",
    "source_url",
    "methodology_url",
    "raw_payload_path",
    "source_metadata",
    "provider_metadata",
    "warnings",
    "error",
    "elasticity_error",
    "elastic_note",
    "scientific_boundary",
    "r_hkl_model_note",
    "弹性备注",
    "J_R_hkl通道说明",
}
_INTEGER_HEADERS = {
    "h",
    "k",
    "i",
    "l",
    "multiplicity",
    "rank_by_intensity",
    "rank_by_R_hkl",
    "rank_by_R_hkl_no_lp",
    "space_group_number",
    "site_count_asymmetric",
    "site_count_unit_cell",
    "reflection_count",
    "coincident_hkl_family_count",
    "frequency",
}
_PERCENT_SCALE_HEADERS = {
    "phase_relative_R_hkl_pct",
    "phase_relative_R_hkl_no_lp_pct",
}
_NORMALIZED_INTENSITY_HEADERS = {"normalized_intensity", "relative_intensity"}
_RELATIVE_BAR_HEADERS = _PERCENT_SCALE_HEADERS | _NORMALIZED_INTENSITY_HEADERS
_DEGREE_HEADERS = {
    "theta_deg",
    "two_theta_deg",
    "two_theta_cu_ka_deg",
    "alpha_deg",
    "beta_deg",
    "gamma_deg",
    "θ_deg",
    "2θ_当前_deg",
    "2θ_CuKa_deg",
}
_FOUR_DECIMAL_HEADERS = {
    "d_spacing_A",
    "d_A",
    "q_invA",
    "g_invA",
    "wavelength_A",
    "energy_keV",
    "a_A",
    "b_A",
    "c_A",
    "young_modulus_hkl_normal_GPa",
    "cell_volume_A3",
    "formula_weight_g_mol",
    "density_g_cm3",
    "q_1/Å",
    "g_1/Å",
    "x",
}
_SCIENTIFIC_HEADERS = {
    "intensity_with_lp",
    "intensity_no_lp",
    "volume_normalized_intensity_with_lp",
    "volume_normalized_intensity_no_lp",
    "material_scattering_factor_R_hkl",
    "material_scattering_factor_R_hkl_no_lp",
    "inverse_R_hkl",
    "inverse_R_hkl_no_lp",
    "structure_factor_sq",
    "mean_structure_factor_sq_per_multiplicity",
    "mean_structure_factor_abs_per_multiplicity",
}


def cjk_display_width(value: object) -> int:
    """Estimate Excel character width, counting full-width CJK characters as two."""

    width = 0
    for character in str(value):
        if unicodedata.category(character) in {"Mn", "Mc", "Me"}:
            continue
        width += 2 if unicodedata.east_asian_width(character) in {"F", "W"} else 1
    return width


def _canonical_header(header: str) -> str:
    return _ZH_TO_CANONICAL.get(str(header), str(header))


def number_format_for_header(header: str) -> str | None:
    """Choose a display format by field meaning without changing stored values."""

    name = _canonical_header(header)
    if name in _INTEGER_HEADERS or name.endswith("_count") or name.endswith("_number"):
        return "0"
    if name in _PERCENT_SCALE_HEADERS:
        # Values already use a 0–100 scale; Excel's percent format would multiply by 100.
        return '0.00"%"'
    if name in _NORMALIZED_INTENSITY_HEADERS:
        return "0.00"
    if name in _DEGREE_HEADERS or name.endswith("_deg"):
        return "0.00000"
    if name in _SCIENTIFIC_HEADERS or name in {"lp_factor", "sin_theta", "cos_theta"}:
        return "0.000E+00"
    if name.startswith("C") and re.fullmatch(r"C[1-6][1-6]_GPa", name):
        return "0.000"
    if name in _FOUR_DECIMAL_HEADERS or name.endswith("_A") or name.endswith("_invA"):
        return "0.0000"
    return None


def number_format_for_value(header: str, value: object) -> str | None:
    """Use scientific notation for nonzero values that fixed decimals would erase."""

    number_format = number_format_for_header(header)
    if number_format is None or isinstance(value, bool) or not isinstance(value, (int, float)):
        return number_format
    if value == 0 or "E+00" in number_format or number_format == "0":
        return number_format
    decimal_match = re.search(r"\.(0+)", number_format)
    if decimal_match is None:
        return number_format
    half_last_place = 0.5 * (10.0 ** -len(decimal_match.group(1)))
    if abs(float(value)) >= half_last_place:
        return number_format
    percent_suffix = '"%"' if number_format.endswith('"%"') else ""
    return f"0.000E+00{percent_suffix}"


def _header_group(header: str) -> str:
    name = _canonical_header(header)
    low = name.lower()
    if name in _IDENTITY_HEADERS:
        return "identity"
    if name in _INDEX_HEADERS or name in _INDEX_LABELS:
        return "indices"
    if name in _GEOMETRY_HEADERS or name.endswith(("_deg", "_invA", "_A")):
        return "geometry"
    if name in _INTENSITY_HEADERS or "intensity" in low or "r_hkl" in low or "_j_" in low:
        return "intensity"
    if name in _ELASTICITY_HEADERS or low.startswith("cij") or re.fullmatch(r"C[1-6][1-6]_GPa", name):
        return "elasticity"
    if name in _QUALITY_HEADERS or low.endswith(("_status", "_error", "_warning", "_warnings")):
        return "quality"
    if name in _DETAIL_HEADERS or any(
        token in low for token in ("sha256", "metadata", "_path", "_url", "boundary", "note")
    ):
        return "metadata"
    return "quality"


def _column_width(header: str, samples: Sequence[object]) -> float:
    name = _canonical_header(header)
    if name in _LONG_TEXT_HEADERS or any(
        token in name.lower() for token in ("_url", "_path", "metadata", "boundary", "note")
    ):
        return 38.0
    header_width = cjk_display_width(header) + 2
    sample_width = max((cjk_display_width(sample) for sample in samples), default=0) + 2
    width = max(header_width, sample_width)
    if name == "phase_name":
        return min(max(width, 20), 32)
    if name == "cif_name":
        return min(max(width, 20), 32)
    if name in {"cif_sha256", "elasticity_sha256"}:
        return 18.0
    if name in _INDEX_HEADERS:
        return min(max(width, 10), 20)
    if name in _GEOMETRY_HEADERS or name.endswith(("_deg", "_invA", "_A")):
        return min(max(width, 11), 18)
    if name in _INTENSITY_HEADERS or "intensity" in name.lower() or "r_hkl" in name.lower():
        return min(max(width, 13), 24)
    return min(max(width, 11), 28)


def _is_wrapped_text_column(header: str, sheet_title: str) -> bool:
    name = _canonical_header(header)
    if sheet_title == "Patterns" and name in {"phase_name", "cif_name"}:
        return False
    if sheet_title == "Diagnostics" and name in {"item", "message"}:
        return True
    return name in {"phase_name", "cif_name"} or name in _LONG_TEXT_HEADERS or any(
        token in name.lower() for token in ("_url", "_path", "metadata", "boundary", "note")
    )


def _identity_row_height(
    values: Sequence[object],
    headers: Sequence[str],
    widths: Mapping[str, float],
) -> float:
    line_counts = [1]
    for header, value in zip(headers, values, strict=False):
        if _canonical_header(str(header)) not in {"phase_name", "cif_name"}:
            continue
        width = max(widths.get(str(header), 1.0) - 2.0, 1.0)
        line_counts.append(ceil(cjk_display_width(value) / width))
    return min(120.0, max(19.0, 4.0 + max(line_counts) * 15.0))


def _wrapped_line_count(value: object, width: float) -> int:
    text = str(value or "")
    segments = text.splitlines() or [""]
    capacity = max(width - 2.0, 1.0)
    return sum(max(1, ceil(cjk_display_width(segment) / capacity)) for segment in segments)


def _detail_column(header: str, sheet_title: str) -> bool:
    name = _canonical_header(header)
    if name in _DETAIL_HEADERS:
        return True
    if name in {"cif_path", "elasticity_path"}:
        return sheet_title != "Downloads"
    is_peak_sheet = sheet_title in {"Peaks", "推荐峰表"} or sheet_title.startswith("峰_")
    return is_peak_sheet and name in _PEAK_DETAIL_HEADERS


def _contiguous_ranges(indices: Sequence[int]) -> list[tuple[int, int]]:
    if not indices:
        return []
    ordered = sorted(set(indices))
    ranges: list[tuple[int, int]] = []
    start = previous = ordered[0]
    for index in ordered[1:]:
        if index != previous + 1:
            ranges.append((start, previous))
            start = index
        previous = index
    ranges.append((start, previous))
    return ranges


def _phase_identity(header: str) -> bool:
    return _canonical_header(header) == "phase_name"


def _set_print_layout(
    sheet: Worksheet,
    *,
    end_column: int,
    repeat_header: bool,
    print_area_end_row: int | None = None,
) -> None:
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.45
    sheet.page_margins.bottom = 0.45
    sheet.print_area = f"A1:{get_column_letter(max(end_column, 1))}{print_area_end_row or sheet.max_row}"
    if repeat_header:
        sheet.print_title_rows = "1:1"


def _apply_phase_data_bars(
    sheet: Worksheet,
    headers: Sequence[str],
    last_row: int,
) -> None:
    if last_row < 2:
        return
    for column_index, header in enumerate(headers, start=1):
        canonical = _canonical_header(header)
        if canonical not in _RELATIVE_BAR_HEADERS:
            continue
        column = get_column_letter(column_index)
        color = "64A884" if canonical in _NORMALIZED_INTENSITY_HEADERS else "5A9BD5"
        rule = DataBarRule(
            start_type="num",
            start_value=0,
            end_type="num",
            end_value=100,
            color=color,
            showValue=True,
        )
        sheet.conditional_formatting.add(f"{column}2:{column}{last_row}", rule)


def style_data_sheet(
    sheet: Worksheet,
    headers: Sequence[str],
) -> None:
    """Style an existing header-first data sheet without changing its values."""

    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90 if len(headers) > 12 else 100
    has_phase_identity = any(_phase_identity(str(header)) for header in headers)
    sheet.freeze_panes = "C2" if has_phase_identity and len(headers) >= 3 else "A2"
    sheet.sheet_properties.outlinePr.summaryRight = True
    sheet.sheet_view.showOutlineSymbols = True
    last_row = max(sheet.max_row, 1)
    last_column = max(len(headers), 1)
    sheet.auto_filter.ref = f"A1:{get_column_letter(last_column)}{last_row}"
    sheet.sheet_format.defaultRowHeight = 19

    header_lines = 1
    identity_widths: dict[str, float] = {}
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column_index)
        cell.font = _HEADER_FONT
        cell.fill = _GROUP_FILLS[_header_group(str(header))]
        cell.alignment = _HEADER_ALIGNMENT
        cell.border = _HEADER_BOTTOM
        samples = [
            sheet.cell(row=row_index, column=column_index).value
            for row_index in range(2, min(last_row, 81) + 1)
        ]
        width = _column_width(str(header), samples)
        canonical_header = _canonical_header(str(header))
        if sheet.title == "Diagnostics" and canonical_header == "message":
            width = 80.0
        elif sheet.title == "Diagnostics" and canonical_header == "item":
            width = 32.0
        letter = get_column_letter(column_index)
        dimension = sheet.column_dimensions[letter]
        dimension.width = width
        if _canonical_header(str(header)) in {"phase_name", "cif_name"}:
            identity_widths[str(header)] = width
        if _detail_column(str(header), sheet.title):
            dimension.hidden = True
            dimension.outlineLevel = 1
        wrap_capacity = max(width - 2.0, 1.0)
        header_lines = max(header_lines, ceil(cjk_display_width(header) / wrap_capacity))
    sheet.row_dimensions[1].height = max(36, min(72, 10 + header_lines * 15))

    hidden_columns = [
        index
        for index, header in enumerate(headers, start=1)
        if _detail_column(str(header), sheet.title)
    ]
    for _, end in _contiguous_ranges(hidden_columns):
        sheet.column_dimensions[get_column_letter(end)].collapsed = True
    identity_columns = [
        (index, str(header))
        for index, header in enumerate(headers)
        if _canonical_header(str(header)) in {"phase_name", "cif_name"}
    ]
    identity_height_cache: dict[tuple[object, ...], float] = {}
    diagnostic_text_columns = [
        (index, str(header))
        for index, header in enumerate(headers)
        if sheet.title == "Diagnostics"
        and _canonical_header(str(header)) in {"item", "message"}
    ]
    diagnostic_widths = {
        header: 80.0 if _canonical_header(header) == "message" else 32.0
        for _, header in diagnostic_text_columns
    }

    for row in sheet.iter_rows(min_row=2, max_row=last_row):
        for column_index, cell in enumerate(row, start=1):
            header = str(headers[column_index - 1]) if column_index <= len(headers) else ""
            cell.font = _BODY_FONT
            wrapped_text = _is_wrapped_text_column(header, sheet.title)
            cell.alignment = _WRAP_ALIGNMENT if wrapped_text else _BODY_ALIGNMENT
            if cell.row % 2 == 0:
                cell.fill = _BAND_FILL
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                number_format = number_format_for_value(header, cell.value)
                if number_format:
                    cell.number_format = number_format
        if identity_columns and sheet.title != "Patterns":
            identity_values = tuple(row[index].value for index, _ in identity_columns)
            height = identity_height_cache.get(identity_values)
            if height is None:
                height = _identity_row_height(
                    identity_values,
                    [header for _, header in identity_columns],
                    identity_widths,
                )
                identity_height_cache[identity_values] = height
            if height > 19.0:
                sheet.row_dimensions[row[0].row].height = height
        if diagnostic_text_columns:
            line_count = max(
                _wrapped_line_count(row[index].value, diagnostic_widths[header])
                for index, header in diagnostic_text_columns
            )
            height = min(409.5, max(19.0, 4.0 + line_count * 15.0))
            if height > 19.0:
                sheet.row_dimensions[row[0].row].height = height

    if last_row >= 2 and last_column:
        _apply_phase_data_bars(sheet, headers, last_row)

    # The overview is the preferred print view; long peak sheets print only
    # through normalized intensity while retaining every later column in Excel.
    normalized_column = next(
        (
            index
            for index, header in enumerate(headers, start=1)
            if _canonical_header(str(header)) in _NORMALIZED_INTENSITY_HEADERS
        ),
        None,
    )
    print_column = normalized_column or last_column
    _set_print_layout(
        sheet,
        end_column=print_column,
        repeat_header=True,
        print_area_end_row=last_row,
    )


def style_guide_sheet(sheet: Worksheet, rows: Sequence[Sequence[object]]) -> None:
    """Style the two-column usage guide, including its section rows."""

    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 95
    sheet.freeze_panes = "A2"
    sheet.sheet_format.defaultRowHeight = 21
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 90
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
        cell.fill = _OVERVIEW_HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT
        cell.border = _HEADER_BOTTOM
    sheet.row_dimensions[1].height = 34
    for row_index, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        source_index = row_index - 1
        section = (
            source_index < len(rows)
            and bool(rows[source_index])
            and str(rows[source_index][0]).startswith("——")
        )
        for cell in row:
            cell.font = _SECTION_FONT if section else _BODY_FONT
            cell.alignment = _WRAP_ALIGNMENT
            if section:
                cell.fill = _SECTION_FILL
        if section:
            sheet.row_dimensions[row_index].height = 25
    _set_print_layout(sheet, end_column=2, repeat_header=True)
