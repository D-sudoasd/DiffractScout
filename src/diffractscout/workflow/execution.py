"""Execution callbacks shared by the existing numerical pipeline."""
from __future__ import annotations

import json
import os
import shutil
from pathlib import Path
from typing import Any

from ..models import DiagnosticRecord
from ..pipeline import _analyze_paths
from ..utils import sha256_file, to_jsonable, write_json
from .checkpoints import fingerprint, load_analysis, publish_checkpoint, save_analysis, verify_checkpoint


class WorkflowCancelled(BaseException):
    """Bypass provider/per-phase Exception handlers at a safe boundary."""


class WorkflowIntegrityError(BaseException):
    """An invalid saved checkpoint must not be converted to an empty query."""


class Execution:
    def __init__(self, store, request, *, api_key: str | None = None):
        self.store, self.root, self.request = store, store.root, request
        self.max_attempts = request.execution["max_attempts"]
        self.secrets = [key for key in (api_key, os.environ.get("MP_API_KEY")) if key]
        self.extra_diagnostics: list[DiagnosticRecord] = []
        self.provider = None
        self.analysis_refs: list[dict[str, str]] = []

    def scrub(self, value: Any) -> Any:
        if isinstance(value, str):
            for key in self.secrets:
                value = value.replace(key, "[redacted]")
            return value
        if isinstance(value, dict):
            return {key: self.scrub(item) for key, item in value.items()}
        if isinstance(value, (list, tuple)):
            return [self.scrub(item) for item in value]
        return value

    def diagnostic(self, code: str, error: Any, *, stage: str = "execution", item: str = "",
                   retryable: bool = False, missing_fields=None) -> dict[str, Any]:
        return {"code": code, "message": self.scrub(str(error)), "stage": stage, "item": item,
                "retryable": retryable, "missing_fields": missing_fields or []}

    def check_cancel(self) -> None:
        if (self.root / "cancel.json").exists():
            raise WorkflowCancelled()

    def analyze_paths(self, paths_and_tensors, settings):
        items = list(paths_and_tensors)
        analyses, diagnostics = [], []
        self.store.stage("analysis", total=len(items))
        self.analysis_refs = []
        for index, (path, tensor) in enumerate(items):
            self.check_cancel()
            # The sidecar path contains a random staging directory. Its content
            # and scientific metadata, not that temporary path, identify reuse.
            tensor_data = to_jsonable(tensor)
            if tensor_data:
                tensor_data.pop("raw_payload_path", None)
            key = fingerprint({"cif": sha256_file(path), "name": path.name,
                               "settings": settings, "tensor": tensor_data})
            checkpoint = self.root / "analyses" / key
            if not checkpoint.exists() and self.request.parent_run_id:
                from .store import job_path, read_json
                parent = job_path(self.root.parent.parent, self.request.parent_run_id)
                previous = parent / "analyses" / key
                if previous.exists() and read_json(parent / "environment.json") == read_json(self.root / "environment.json"):
                    verify_checkpoint(previous)
                    publish_checkpoint(checkpoint, lambda directory: shutil.copytree(previous, directory, dirs_exist_ok=True))
                    self.store.event("parent_checkpoint_reused", item=path.name, parent_run_id=self.request.parent_run_id)
            if checkpoint.exists():
                analysis = load_analysis(checkpoint, cif_path=path)
                analysis.elastic_tensor = tensor
                self.store.count("reused")
                item_diagnostics = [DiagnosticRecord("analysis", path.name, "warning", value)
                                    for value in analysis.warnings]
                results = [analysis]
            else:
                self.store.count("analyses")
                results, item_diagnostics = _analyze_paths([(path, tensor)], settings)
                if results:
                    publish_checkpoint(checkpoint, lambda directory: save_analysis(directory, results[0]))
            if results:
                self.analysis_refs.append({"checkpoint": str(checkpoint.relative_to(self.root)), "cif_name": path.name})
                analyses.extend(results)
            diagnostics.extend(item_diagnostics)
            self.store.update(completed_items=index+1, total_items=len(items))
            self.store.event("phase_completed", item=path.name, success=bool(results))
        return analyses, diagnostics

    def before_export(self, diagnostics) -> None:
        self.check_cancel()
        diagnostics.extend(self.extra_diagnostics)
        if self.provider:
            diagnostics.extend(DiagnosticRecord(row["stage"], row["item"], "error", row["message"])
                               for row in self.provider.failures)
        diagnostics[:] = [DiagnosticRecord(row.stage, self.scrub(row.item), row.level, self.scrub(row.message))
                          for row in diagnostics]
        self.store.stage("export")

    def before_verify(self, staging: Path) -> None:
        self.check_cancel()
        self.store.stage("verification")
        if self.request.exports["template"] == "peak_table":
            self._compact_workbook(staging)
        self.decorate_bundle(staging, analysis=self.request.workflow != "discover")

    def decorate_bundle(self, staging: Path, *, analysis: bool) -> None:
        """Add reproducibility records before the original atomic publication."""
        from .preparation import manifest
        from .store import read_json
        request = read_json(self.root / "request.json")
        write_json(staging / "workflow_request.json", request)
        write_json(staging / "workflow_environment.json", read_json(self.root / "environment.json"))
        write_json(staging / "workflow_summary.json", {
            "run_id": self.store.state["run_id"], "workflow": self.request.workflow,
            "parameter_sources": read_json(self.root / "plan.json").get("parameter_sources", {}),
            "counters": self.store.state["counters"], "scientific_scope": "theoretical_reference",
        })
        sources = self.root / "sources"
        if sources.exists():
            shutil.copytree(sources, staging / "workflow_sources")
        replay = dict(request)
        replay.update(output_dir=None, submission_id=None, parent_run_id=None)
        if analysis:
            replay.update(workflow="analyze", inputs=[f"inputs/{path.name}" for path in sorted((staging / "inputs").glob("*.cif"))],
                          composition="", preparation={}, adaptation={}, elastic_overrides={})
            for key, value in self.request.elastic_overrides.items():
                payload = dict(value)
                if payload.get("raw_payload_path"):
                    payload["raw_payload_path"] = "workflow_sources/" + Path(payload["raw_payload_path"]).relative_to(sources).as_posix()
                replay["elastic_overrides"][Path(key).name] = payload
        elif sources.exists():
            mapping = read_json(sources / "source-map.json")
            replay["inputs"] = ["workflow_sources/"+name for name in mapping["inputs"]]
            replay["preparation"] = dict(replay["preparation"])
            if mapping["templates"]:
                replay["preparation"]["templates"] = {phase: "workflow_sources/"+name for phase, name in mapping["templates"].items()}
            if mapping["parameter_file"]:
                replay["preparation"]["parameter_file"] = "workflow_sources/"+mapping["parameter_file"]
        replay["execution"] = {**replay["execution"], "overwrite": False}
        write_json(staging / "replay_request.json", replay)
        manifest(staging)

    def _compact_workbook(self, staging: Path) -> None:
        if not self.request.exports["include_excel"]:
            return
        import csv
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Peak reference"
        sheet.append(["Phase", "hkl", "2theta (deg)", "d (Angstrom)"])
        with (staging / "peak_reference.csv").open(encoding="utf-8-sig", newline="") as stream:
            for row in csv.DictReader(stream):
                phase = row["phase_name"]
                if phase.startswith(("=", "+", "-", "@")):
                    phase = "'" + phase
                sheet.append([phase, row.get("family_label") or f"({row['h']} {row['k']} {row['l']})",
                              float(row["two_theta_deg"]), float(row["d_spacing_A"])])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E8E8E8")
        for column, width in (("A", 28), ("B", 18), ("C", 20), ("D", 20)):
            sheet.column_dimensions[column].width = width
        for row in sheet.iter_rows(min_row=2, min_col=3, max_col=4):
            for cell in row:
                cell.number_format = "0.000000"
        path = staging / "peak_table.xlsx"
        workbook.save(path)
        manifest = json.loads((staging / "manifest.json").read_text(encoding="utf-8"))
        manifest["files"].append({"path": path.name, "sha256": sha256_file(path),
                                  "size_bytes": path.stat().st_size, "role": "tabular_result"})
        write_json(staging / "manifest.json", manifest)
