"""Focused checks for the workbook overview and its internal navigation."""

from __future__ import annotations

from openpyxl import Workbook, load_workbook

from diffractscout.excel_overview import add_overview_sheet


def _workbook_with_data_sheets() -> Workbook:
    workbook = Workbook()
    workbook.active.title = "Summary"
    workbook.create_sheet("User's Results")
    workbook.create_sheet("Diagnostics")
    return workbook


def test_overview_summarizes_supplied_rows_and_all_phases(tmp_path) -> None:
    workbook = _workbook_with_data_sheets()
    phases = [
        {
            "phase_name": f"phase-{index}",
            "formula": "Al2O3",
            "space_group_symbol": "R -3 c",
            "space_group_number": 167,
            "reflection_count": index,
            "elastic_status": "not_requested",
        }
        for index in range(1, 27)
    ]
    summary = [
        {"key": "effective_wavelength_A", "value": 1.540612345678},
        {"key": "effective_energy_keV", "value": 8.047000000123},
        {"key": "effective_radiation_source", "value": "source_preset:Cu Ka"},
        {"key": "effective_two_theta_range_deg", "value": [10.0, 90.0]},
        {"key": "analysis_two_theta_range_deg", "value": [5.0, 120.0]},
        {"key": "profile_sampled_two_theta_range_deg", "value": [5.0, 120.0]},
        {"key": "profile_model", "value": "pseudo_voigt"},
        {"key": "input_mode", "value": "source"},
        {"key": "candidate_count", "value": 4},
    ]
    diagnostics = [
        {"level": "error", "message": "failed"},
        {"level": "warning", "message": "check"},
        {"level": "info", "message": "ready"},
    ]

    overview = add_overview_sheet(
        workbook,
        summary=summary,
        phases=phases,
        diagnostics=diagnostics,
        lab_views=True,
    )

    assert overview.title == "结果概览"
    assert overview["B5"].value == 26
    assert overview["D5"].value == sum(range(1, 27))
    assert overview["B6"].value == 1
    assert overview["D6"].value == 1
    assert overview["B11"].value == 1.540612345678
    assert overview["B11"].number_format == "0.####"
    assert overview["B12"].value == 8.047000000123
    assert overview["B12"].number_format == "0.####"
    assert overview["B13"].value == "Cu Kα"
    assert overview["B14"].value == "10–90°"
    assert overview["B17"].value == "伪Voigt"
    assert overview["B18"].value == "预设光源"
    phase_header = next(
        row for row in range(1, overview.max_row + 1) if overview.cell(row, 1).value == "物相名称"
    )
    assert overview.cell(phase_header + 1, 1).value == "phase-1"
    assert overview.cell(phase_header + 26, 1).value == "phase-26"
    assert overview.page_setup.orientation == "landscape"
    assert str(overview.page_setup.paperSize) == str(overview.PAPERSIZE_A4)
    assert overview.page_setup.fitToWidth == 1
    assert overview.page_setup.fitToHeight == 0
    assert {cell.font.name for cells in overview.iter_rows() for cell in cells} == {
        "Microsoft YaHei"
    }

    path = tmp_path / "overview.xlsx"
    workbook.save(path)
    reopened = load_workbook(path, data_only=False)
    saved = reopened["结果概览"]
    navigation_header = next(
        row for row in range(1, saved.max_row + 1) if saved.cell(row, 1).value == "工作表"
    )
    user_sheet_row = next(
        row
        for row in range(navigation_header + 1, saved.max_row + 1)
        if saved.cell(row, 1).value == "User's Results"
    )
    user_sheet_link = saved.cell(user_sheet_row, 1).hyperlink
    assert user_sheet_link is not None
    assert user_sheet_link.target is None
    assert user_sheet_link.location == "'User''s Results'!A1"


def test_overview_escapes_untrusted_text_and_leaves_no_formulas() -> None:
    workbook = _workbook_with_data_sheets()
    summary = [
        {"key": "effective_radiation_source", "value": "=HYPERLINK(\"https://bad.invalid\")"},
        {"key": "profile_model", "value": "+CMD|' /C calc'!A0"},
        {"key": "input_mode", "value": "energy"},
    ]
    phases = [
        {
            "phase_name": "=1+1",
            "formula": "@SUM(A1:A2)",
            "space_group_symbol": "P1",
            "space_group_number": 1,
            "reflection_count": 2,
            "elastic_status": "-not checked",
        }
    ]

    overview = add_overview_sheet(
        workbook,
        summary=summary,
        phases=phases,
        diagnostics=[],
        lab_views=False,
    )

    assert overview.title == "Overview"
    string_values = [
        cell.value
        for row in overview.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    ]
    assert any(value.startswith("'=HYPERLINK") for value in string_values)
    assert any(value.startswith("'+CMD") for value in string_values)
    assert overview["B18"].value == "Photon energy"
    phase_header = next(
        row for row in range(1, overview.max_row + 1) if overview.cell(row, 1).value == "Phase name"
    )
    assert overview.cell(phase_header + 1, 1).value == "'=1+1"
    assert overview.cell(phase_header + 1, 2).value == "'@SUM(A1:A2)"
    assert overview.cell(phase_header + 1, 5).value == "'-not checked"
    assert all(
        cell.data_type != "f"
        for row in overview.iter_rows()
        for cell in row
    )


def test_empty_overview_reports_no_results_and_links_only_existing_sheets() -> None:
    workbook = Workbook()
    workbook.active.title = "Peaks"

    overview = add_overview_sheet(
        workbook,
        summary=[],
        phases=[],
        diagnostics=[],
        lab_views=False,
    )

    assert overview["B5"].value == 0
    assert overview["D5"].value == 0
    assert any(
        cell.value == "No analyzed phases"
        for row in overview.iter_rows()
        for cell in row
    )
    links = [
        cell.hyperlink
        for row in overview.iter_rows()
        for cell in row
        if cell.hyperlink is not None
    ]
    assert len(links) == 1
    assert links[0].target is None
    assert links[0].location == "'Peaks'!A1"


def test_long_phase_text_gets_wrapped_row_height_and_scientific_number_format() -> None:
    workbook = _workbook_with_data_sheets()
    phase = {
        "phase_name": "长物相名称" * 12,
        "formula": "Al0.12345678901234567890Mg0.98765432109876543210SiO4",
        "space_group_symbol": "P 63/mmc",
        "space_group_number": 194,
        "reflection_count": 18,
        "elastic_status": "elasticity not requested; no source record was supplied; status remains unavailable",
    }

    overview = add_overview_sheet(
        workbook,
        summary=[
            {"key": "effective_wavelength_A", "value": 1e-15},
            {"key": "effective_energy_keV", "value": 1e14},
            {"key": "effective_radiation_source", "value": "custom_source_wavelength"},
        ],
        phases=[phase],
        diagnostics=[],
        lab_views=True,
    )

    assert overview["B11"].value == 1e-15
    assert overview["B11"].number_format == "0.0000E+00"
    assert overview["B12"].value == 1e14
    assert overview["B12"].number_format == "0.0000E+00"
    assert overview["B13"].value == "自定义波长"
    phase_header = next(
        row
        for row in range(1, overview.max_row + 1)
        if overview.cell(row, 1).value == "物相名称"
    )
    phase_row = phase_header + 1
    assert overview.cell(phase_row, 1).value == phase["phase_name"]
    assert overview.row_dimensions[phase_row].height > 30
    assert overview.row_dimensions[phase_row].height <= 409.5
