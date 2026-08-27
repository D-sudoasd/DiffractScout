import csv
import ctypes
import errno
import json
import os
from pathlib import Path
import shutil
import subprocess
import threading
from concurrent.futures import ThreadPoolExecutor
from types import SimpleNamespace

from openpyxl import load_workbook
import pytest

import diffractscout.exporters as exporters
from diffractscout.elasticity_input import parse_cubic_cij
from diffractscout.models import (
    AnalysisSettings,
    CandidateRecord,
    DiscoverySettings,
    DownloadArtifact,
)
from diffractscout.pipeline import (
    _TargetState,
    _acquire_transaction_lock,
    _attempt_restore_backup,
    _capture_target_state,
    _commit_staging_output,
    _copy_local_input,
    _lock_path_for,
    _lookup_elastic_override,
    _normalize_download_artifacts,
    _release_transaction_lock,
    analyze_cifs,
    collect_cif_paths,
    run_pipeline,
)
from diffractscout.utils import write_json
from diffractscout.validation import verify_bundle


_UNSET = object()


def _run_include_elasticity_contract(
    demo_inputs: Path,
    output: Path,
    *,
    settings: AnalysisSettings | None = None,
    include_elasticity: bool | object = _UNSET,
) -> tuple[object, bool]:
    observed: dict[str, bool] = {}

    class Provider:
        name = "include-elasticity-contract"

        def search_subsystem(
            self, chemsys: str, **_kwargs: object
        ) -> list[CandidateRecord]:
            return [
                CandidateRecord(
                    material_id="synthetic-isotropic-cubic",
                    formula="Al",
                    energy_above_hull_eV_atom=0.0,
                    is_stable=True,
                    queried_chemsys=chemsys,
                    source_provider=self.name,
                )
            ]

        def metadata(self) -> dict[str, object]:
            return {"provider": self.name}

        def download_candidates(
            self,
            candidates: object,
            _output_dir: Path,
            **kwargs: object,
        ) -> list[DownloadArtifact]:
            observed["include_elasticity"] = bool(kwargs["include_elasticity"])
            candidate = list(candidates)[0]
            return [
                DownloadArtifact(
                    candidate=candidate,
                    cif_path=demo_inputs / "synthetic_fcc_al.cif",
                    elasticity_path=demo_inputs / "synthetic_fcc_al_elasticity.json",
                    status="ok",
                )
            ]

    options: dict[str, object] = {
        "discovery_settings": DiscoverySettings(
            mode="single_chemsys", max_total=1
        ),
        "include_excel": False,
        "confirm_above": 1,
    }
    if settings is not None:
        options["analysis_settings"] = settings
    if include_elasticity is not _UNSET:
        options["include_elasticity"] = include_elasticity
    result = run_pipeline("Al", Provider(), output, **options)  # type: ignore[arg-type]
    return result, observed["include_elasticity"]


def test_local_pipeline_is_self_contained_and_verifiable(demo_inputs: Path, tmp_path: Path) -> None:
    output = tmp_path / "bundle"
    result = analyze_cifs([demo_inputs], output)
    assert len(result.analyses) == 1
    assert result.manifest_path == output / "manifest.json"
    assert result.analyses[0].structure.cif_path == output / "inputs" / "synthetic_fcc_al.cif"
    assert result.analyses[0].structure.cif_path.is_file()
    assert result.analyses[0].elastic_tensor is not None
    assert result.analyses[0].elastic_tensor.raw_payload_path == (
        output / "inputs" / "synthetic_fcc_al_elasticity.json"
    )
    assert result.analyses[0].elastic_tensor.raw_payload_path.is_file()
    assert (output / "inputs" / "synthetic_fcc_al.cif").is_file()
    assert (output / "peak_reference.csv").is_file()
    assert (output / "results.xlsx").is_file()
    assert verify_bundle(output)["ok"]

    with (output / "peak_reference.csv").open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert rows
    assert rows[0]["hkl"] == "(1 1 1)"
    assert rows[0]["cif_name"] == "synthetic_fcc_al.cif"
    assert rows[0]["volume_normalized_intensity_with_lp"] == rows[0]["material_scattering_factor_R_hkl"]
    assert rows[0]["volume_normalized_intensity_no_lp"] == rows[0]["material_scattering_factor_R_hkl_no_lp"]

    workbook = load_workbook(output / "results.xlsx", read_only=True)
    assert {"Summary", "Phases", "Peaks", "Elasticity", "Candidates", "Downloads", "Patterns"}.issubset(workbook.sheetnames)


def test_run_pipeline_honors_analysis_settings_elasticity_when_override_omitted(
    demo_inputs: Path, tmp_path: Path
) -> None:
    result, observed = _run_include_elasticity_contract(
        demo_inputs,
        tmp_path / "settings-disabled",
        settings=AnalysisSettings(include_elasticity=False),
    )

    assert observed is False
    assert result.analyses[0].metadata["elasticity_requested"] is False


def test_run_pipeline_explicit_elasticity_override_wins_over_analysis_settings(
    demo_inputs: Path, tmp_path: Path
) -> None:
    result, observed = _run_include_elasticity_contract(
        demo_inputs,
        tmp_path / "override-enabled",
        settings=AnalysisSettings(include_elasticity=False),
        include_elasticity=True,
    )

    assert observed is True
    assert result.analyses[0].metadata["elasticity_requested"] is True


def test_run_pipeline_without_settings_preserves_elasticity_default(
    demo_inputs: Path, tmp_path: Path
) -> None:
    result, observed = _run_include_elasticity_contract(
        demo_inputs,
        tmp_path / "default-enabled",
    )

    assert observed is True
    assert result.analyses[0].metadata["elasticity_requested"] is True


def test_excel_omission_warning_is_retained_in_pipeline_result(
    demo_inputs: Path, tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setattr(exporters, "EXCEL_DATA_ROW_LIMIT", 2)
    result = analyze_cifs([demo_inputs], tmp_path / "limited", include_excel=True)
    warnings = [item for item in result.diagnostics if item.stage == "export"]
    assert warnings
    diagnostics_csv = (result.output_dir / "diagnostics.csv").read_text(encoding="utf-8-sig")
    assert warnings[0].message in diagnostics_csv
    workbook = load_workbook(result.output_dir / "results.xlsx", read_only=True, data_only=False)
    assert warnings[0].message in str(workbook["Diagnostics"]["D2"].value)


def test_output_overwrite_requires_known_manifest(demo_inputs: Path, tmp_path: Path) -> None:
    output = tmp_path / "occupied"
    output.mkdir()
    (output / "user.txt").write_text("keep", encoding="utf-8")
    try:
        analyze_cifs([demo_inputs], output, overwrite=True)
    except FileExistsError:
        pass
    else:
        raise AssertionError("Expected refusal to overwrite an unrelated directory")
    assert (output / "user.txt").read_text(encoding="utf-8") == "keep"


def test_output_path_must_be_a_directory(demo_inputs: Path, tmp_path: Path) -> None:
    output = tmp_path / "result-file"
    output.write_text("keep", encoding="utf-8")
    try:
        analyze_cifs([demo_inputs], output)
    except FileExistsError as exc:
        assert "not a directory" in str(exc)
    else:
        raise AssertionError("Expected refusal when the output path is a file")
    assert output.read_text(encoding="utf-8") == "keep"


def test_no_elasticity_does_not_copy_or_load_sidecar(demo_inputs: Path, tmp_path: Path) -> None:
    output = tmp_path / "no-elasticity"
    result = analyze_cifs(
        [demo_inputs],
        output,
        settings=AnalysisSettings(include_elasticity=False),
        include_excel=False,
    )
    assert len(result.analyses) == 1
    assert result.analyses[0].elastic_tensor is None
    assert result.analyses[0].metadata["elasticity_requested"] is False
    assert {
        reflection.elastic_status for reflection in result.analyses[0].reflections
    } == {"not_requested"}
    assert not list((output / "inputs").glob("*_elasticity.json"))


def test_collect_cif_paths_is_case_insensitive(demo_inputs: Path, tmp_path: Path) -> None:
    scan_root = tmp_path / "case-scan"
    scan_root.mkdir()
    upper = scan_root / "UPPER.CIF"
    shutil.copy2(demo_inputs / "synthetic_fcc_al.cif", upper)
    assert collect_cif_paths([scan_root]) == [upper.resolve()]


def test_same_named_inputs_are_preserved_without_collision(
    demo_inputs: Path, tmp_path: Path
) -> None:
    left = tmp_path / "left"
    right = tmp_path / "right"
    left.mkdir()
    right.mkdir()
    shutil.copy2(demo_inputs / "synthetic_fcc_al.cif", left / "phase.cif")
    shutil.copy2(demo_inputs / "synthetic_fcc_al.cif", right / "phase.cif")

    output = tmp_path / "collision-safe"
    result = analyze_cifs(
        [left, right],
        output,
        settings=AnalysisSettings(include_elasticity=False),
        include_excel=False,
    )
    bundled = sorted(path.name for path in (output / "inputs").glob("*.cif"))
    assert len(result.analyses) == 2
    assert bundled[0] == "phase.cif"
    assert len(bundled) == 2
    assert bundled[1].startswith("phase_")


def test_same_named_local_sidecars_are_rebound_to_collision_safe_cifs(
    demo_inputs: Path, tmp_path: Path
) -> None:
    left = tmp_path / "left-with-sidecar"
    right = tmp_path / "right-with-sidecar"
    left.mkdir()
    right.mkdir()
    for directory in (left, right):
        shutil.copy2(demo_inputs / "synthetic_fcc_al.cif", directory / "phase.cif")
        payload = json.loads(
            (demo_inputs / "synthetic_fcc_al_elasticity.json").read_text(encoding="utf-8")
        )
        payload["cif_filename"] = "phase.cif"
        payload["provenance"]["paired_cif"] = "phase.cif"
        payload["diffractscout"]["paired_cif"] = "phase.cif"
        (directory / "phase_elasticity.json").write_text(
            json.dumps(payload),
            encoding="utf-8",
        )

    result = analyze_cifs(
        [left, right],
        tmp_path / "local-sidecar-collision",
        include_excel=False,
    )
    assert len(result.analyses) == 2
    for analysis in result.analyses:
        assert analysis.elastic_tensor is not None
        sidecar = analysis.elastic_tensor.raw_payload_path
        assert sidecar is not None and sidecar.is_file()
        payload = json.loads(sidecar.read_text(encoding="utf-8"))
        assert payload["cif_filename"] == analysis.structure.cif_path.name


def _write_local_elasticity_index(path: Path, rows: list[dict[str, object]]) -> None:
    fields = [
        "cif_name",
        "status",
        "source_provider",
        "source_record_id",
        "coordinate_frame",
        *[f"C{i}{j}_GPa" for i in range(1, 7) for j in range(1, 7)],
    ]
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def test_local_elasticity_csv_index_is_staged_and_loaded(
    demo_inputs: Path, tmp_path: Path
) -> None:
    source_dir = tmp_path / "single-index-source"
    source_dir.mkdir()
    cif = source_dir / "one.cif"
    shutil.copy2(demo_inputs / "synthetic_fcc_al.cif", cif)
    _write_local_elasticity_index(
        source_dir / "elasticity.csv",
        [
            {
                "cif_name": cif.name,
                "status": "valid",
                "source_provider": "local-index",
                "source_record_id": "one",
                "coordinate_frame": "crystal_cartesian_from_cif_lattice",
                **{
                    f"C{i}{j}_GPa": 150 if i == j else 0
                    for i in range(1, 7)
                    for j in range(1, 7)
                },
            }
        ],
    )

    result = analyze_cifs(
        [cif],
        tmp_path / "single-index-output",
        include_excel=False,
    )
    tensor = result.analyses[0].elastic_tensor
    assert tensor is not None
    assert tensor.stiffness_GPa[0, 0] == pytest.approx(150.0)
    assert tensor.raw_payload_path == result.output_dir / "inputs" / "elasticity.csv"
    assert tensor.raw_payload_path.is_file()


def test_shared_local_elasticity_index_is_reused_for_two_cifs(
    demo_inputs: Path, tmp_path: Path
) -> None:
    source_dir = tmp_path / "shared-index-source"
    source_dir.mkdir()
    cifs = [source_dir / "one.cif", source_dir / "two.cif"]
    for cif in cifs:
        shutil.copy2(demo_inputs / "synthetic_fcc_al.cif", cif)
    _write_local_elasticity_index(
        source_dir / "elasticity_index.csv",
        [
            {
                "cif_name": cif.name,
                "status": "valid",
                "source_provider": "shared-local-index",
                "source_record_id": cif.stem,
                "coordinate_frame": "crystal_cartesian_from_cif_lattice",
                **{
                    f"C{i}{j}_GPa": value if i == j else 0
                    for i in range(1, 7)
                    for j in range(1, 7)
                    for value in (160.0 if cif.stem == "one" else 170.0,)
                },
            }
            for cif in cifs
        ],
    )

    result = analyze_cifs(
        [source_dir],
        tmp_path / "shared-index-output",
        include_excel=False,
    )
    assert len(result.analyses) == 2
    values = {
        analysis.structure.cif_path.name: float(analysis.elastic_tensor.stiffness_GPa[0, 0])
        for analysis in result.analyses
        if analysis.elastic_tensor is not None
    }
    assert values == {"one.cif": 160.0, "two.cif": 170.0}
    raw_paths = {
        analysis.elastic_tensor.raw_payload_path
        for analysis in result.analyses
        if analysis.elastic_tensor is not None
    }
    assert raw_paths == {result.output_dir / "inputs" / "elasticity_index.csv"}


def test_missing_explicit_input_is_not_silently_ignored(tmp_path: Path) -> None:
    missing = tmp_path / "missing.cif"
    try:
        collect_cif_paths([missing])
    except FileNotFoundError as exc:
        assert str(missing) in str(exc)
    else:
        raise AssertionError("Expected an explicit missing input to raise FileNotFoundError")


def test_input_output_overlap_is_rejected_before_writes(demo_inputs: Path) -> None:
    output = demo_inputs / "nested-result"
    try:
        analyze_cifs([demo_inputs], output, include_excel=False)
    except ValueError as exc:
        assert "must be disjoint" in str(exc)
    else:
        raise AssertionError("Expected overlapping input/output paths to be rejected")
    assert not output.exists()


def test_invalid_run_settings_fail_before_output_writes(
    demo_inputs: Path, tmp_path: Path
) -> None:
    output = tmp_path / "invalid-settings"
    with pytest.raises(ValueError, match="2theta range"):
        analyze_cifs(
            [demo_inputs],
            output,
            settings=AnalysisSettings(
                two_theta_min_deg=120.0,
                two_theta_max_deg=5.0,
            ),
            include_excel=False,
        )
    assert not output.exists()


def test_invalid_figure_preset_fails_before_output_writes(
    demo_inputs: Path, tmp_path: Path
) -> None:
    output = tmp_path / "invalid-figure-preset"
    with pytest.raises(ValueError, match="Unknown figure export preset"):
        analyze_cifs(
            [demo_inputs],
            output,
            settings=AnalysisSettings(
                include_figures=True,
                figure_preset="not-a-preset",
            ),
            include_excel=False,
        )
    assert not output.exists()


def test_workbook_contains_structured_diagnostics_sheet(demo_inputs: Path, tmp_path: Path) -> None:
    output = tmp_path / "diagnostic-workbook"
    analyze_cifs([demo_inputs], output)
    workbook = load_workbook(output / "results.xlsx", read_only=True)
    assert "Diagnostics" in workbook.sheetnames
    headers = [cell.value for cell in next(workbook["Diagnostics"].iter_rows(max_row=1))]
    assert headers == ["stage", "item", "level", "message"]


def test_overwrite_requires_an_intact_existing_bundle(demo_inputs: Path, tmp_path: Path) -> None:
    output = tmp_path / "damaged-bundle"
    analyze_cifs([demo_inputs], output, include_excel=False)
    (output / "phase_summary.csv").write_text("damaged", encoding="utf-8")
    try:
        analyze_cifs([demo_inputs], output, include_excel=False, overwrite=True)
    except FileExistsError as exc:
        assert "fails integrity verification" in str(exc)
    else:
        raise AssertionError("Expected refusal to replace a damaged result bundle")
    assert (output / "phase_summary.csv").read_text(encoding="utf-8") == "damaged"


def test_failed_staged_export_preserves_previous_verified_bundle(
    demo_inputs: Path, tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    output = tmp_path / "previous-bundle"
    analyze_cifs([demo_inputs], output, include_excel=False)
    original_manifest = (output / "manifest.json").read_bytes()

    def fail_export(*_args, **_kwargs):
        raise RuntimeError("simulated export failure")

    monkeypatch.setattr(pipeline, "export_result_bundle", fail_export)
    try:
        pipeline.analyze_cifs(
            [demo_inputs], output, include_excel=False, overwrite=True
        )
    except RuntimeError as exc:
        assert "simulated export failure" in str(exc)
    else:
        raise AssertionError("Expected staged export failure")
    assert (output / "manifest.json").read_bytes() == original_manifest
    assert verify_bundle(output)["ok"]


def test_target_created_during_run_is_not_silently_replaced(
    demo_inputs: Path, tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    output = tmp_path / "raced-target"
    original_export = pipeline.export_result_bundle

    def export_then_occupy(*args, **kwargs):
        manifest = original_export(*args, **kwargs)
        output.mkdir()
        (output / "user.txt").write_text("keep", encoding="utf-8")
        return manifest

    monkeypatch.setattr(pipeline, "export_result_bundle", export_then_occupy)
    with pytest.raises(FileExistsError, match="not empty"):
        analyze_cifs([demo_inputs], output, include_excel=False)

    assert (output / "user.txt").read_text(encoding="utf-8") == "keep"
    assert not list(tmp_path.glob(".raced-target.diffractscout-*"))


def test_target_created_after_final_validation_is_not_deleted(
    demo_inputs: Path, tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    output = tmp_path / "post-validation-race"
    original_validate = pipeline._validate_output_target
    call_count = 0

    def validate_then_occupy(path: str | Path, *, overwrite: bool) -> Path:
        nonlocal call_count
        call_count += 1
        result = original_validate(path, overwrite=overwrite)
        if call_count == 2:
            result.mkdir()
            (result / "user.txt").write_text("keep", encoding="utf-8")
        return result

    monkeypatch.setattr(pipeline, "_validate_output_target", validate_then_occupy)
    with pytest.raises(FileExistsError, match="changed during"):
        analyze_cifs([demo_inputs], output, include_excel=False)

    assert (output / "user.txt").read_text(encoding="utf-8") == "keep"
    assert not list(tmp_path.glob(".post-validation-race.diffractscout-*"))


def test_manifest_member_mutation_after_final_validation_is_not_overwritten(
    demo_inputs: Path, tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    output = tmp_path / "member-race"
    analyze_cifs([demo_inputs], output, include_excel=False)
    original_validate = pipeline._validate_output_target
    call_count = 0

    def validate_then_mutate(path: str | Path, *, overwrite: bool) -> Path:
        nonlocal call_count
        call_count += 1
        result = original_validate(path, overwrite=overwrite)
        if call_count == 2:
            (result / "phase_summary.csv").write_text(
                "changed-after-final-validation\n",
                encoding="utf-8",
            )
        return result

    monkeypatch.setattr(pipeline, "_validate_output_target", validate_then_mutate)
    with pytest.raises(FileExistsError, match="changed|unverifiable"):
        analyze_cifs(
            [demo_inputs],
            output,
            include_excel=False,
            overwrite=True,
        )

    assert (output / "phase_summary.csv").read_text(encoding="utf-8") == (
        "changed-after-final-validation\n"
    )
    assert not list(tmp_path.glob(".member-race.diffractscout-*"))


def test_rollback_conflict_preserves_old_bundle_backup(
    demo_inputs: Path, tmp_path: Path, monkeypatch
) -> None:
    output = tmp_path / "rollback-conflict"
    analyze_cifs([demo_inputs], output, include_excel=False)
    expected_state = _capture_target_state(output)
    staging = tmp_path / "staging"
    staging.mkdir()
    original_replace = Path.replace

    def isolate_then_occupy(self: Path, target: str | Path) -> Path:
        result = original_replace(self, target)
        if self == output and Path(target).name.startswith(f".{output.name}.backup-"):
            output.mkdir()
            (output / "user.txt").write_text("external", encoding="utf-8")
        return result

    monkeypatch.setattr(Path, "replace", isolate_then_occupy)
    with pytest.raises(FileExistsError, match="Rollback conflict"):
        _commit_staging_output(
            output,
            staging,
            expected_state=expected_state,
        )

    assert (output / "user.txt").read_text(encoding="utf-8") == "external"
    backups = sorted(tmp_path.glob(".rollback-conflict.backup-*"))
    assert len(backups) == 1
    assert verify_bundle(backups[0])["ok"]


def test_rollback_diagnostic_preserves_primary_exception_without_add_note(
    tmp_path: Path,
) -> None:
    backup = tmp_path / "legacy-backup"
    backup.mkdir()
    (backup / "old.txt").write_text("old", encoding="utf-8")
    target = tmp_path / "legacy-target"
    target.mkdir()
    (target / "external.txt").write_text("external", encoding="utf-8")

    class LegacyException(Exception):
        add_note = None

    primary = LegacyException("primary rollback failure")
    with pytest.warns(RuntimeWarning, match="Rollback conflict"):
        _attempt_restore_backup(backup, target, primary_error=primary)

    assert str(primary) == "primary rollback failure"
    assert (target / "external.txt").read_text(encoding="utf-8") == "external"
    assert (backup / "old.txt").read_text(encoding="utf-8") == "old"


def test_rollback_restore_uses_atomic_no_replace_when_target_reappears(
    tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    backup = tmp_path / "restore-backup"
    backup.mkdir()
    (backup / "old.txt").write_text("old", encoding="utf-8")
    target = tmp_path / "restore-target"

    def race(source: Path, destination: Path) -> None:
        assert source == backup
        destination.mkdir()
        (destination / "external.txt").write_text("keep", encoding="utf-8")
        raise FileExistsError("target appeared during restore")

    monkeypatch.setattr(pipeline, "_rename_directory_noreplace", race)
    warning_sink: list[str] = []
    with pytest.warns(RuntimeWarning, match="Could not restore isolated backup"):
        _attempt_restore_backup(backup, target, warning_sink=warning_sink)

    assert (target / "external.txt").read_text(encoding="utf-8") == "keep"
    assert (backup / "old.txt").read_text(encoding="utf-8") == "old"
    assert any("backup was preserved" in message for message in warning_sink)


def test_target_state_survives_same_filesystem_directory_rename(
    demo_inputs: Path, tmp_path: Path
) -> None:
    output = tmp_path / "rename-stable-state"
    analyze_cifs([demo_inputs], output, include_excel=False)
    before = _capture_target_state(output)

    renamed = tmp_path / "rename-stable-state-moved"
    output.rename(renamed)
    after = _capture_target_state(renamed)

    assert after == before
    if before.identity not in {None, (0, 0)}:
        identical_swap = tmp_path / "rename-stable-state-identical"
        shutil.copytree(renamed, identical_swap)
        swapped = _capture_target_state(identical_swap)
        assert swapped.members == before.members
        assert swapped.identity != before.identity


def test_stale_lock_recovery_and_live_lock_protection(
    tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    target = tmp_path / "locked-target"
    lock = _lock_path_for(target)
    lock.write_text(
        json.dumps(
            {
                "version": 1,
                "host": pipeline._lock_host(),
                "pid": 424242,
                "created_at": 1.0,
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(pipeline, "_process_is_alive", lambda pid: pid == os.getpid())

    acquired = _acquire_transaction_lock(target)
    metadata = json.loads(acquired.read_text(encoding="utf-8"))
    assert metadata["pid"] == os.getpid()
    assert metadata["host"] == pipeline._lock_host()
    _release_transaction_lock(acquired)
    assert not lock.exists()

    lock.write_text(
        json.dumps(
            {
                "version": 1,
                "host": pipeline._lock_host(),
                "pid": os.getpid(),
                "created_at": 1.0,
            }
        ),
        encoding="utf-8",
    )
    with pytest.raises(FileExistsError, match="active"):
        _acquire_transaction_lock(target)
    assert lock.exists()


def test_stale_lock_restore_never_replaces_live_lock_after_vacancy_check(
    tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    target = tmp_path / "stale-restore-race-target"
    lock = _lock_path_for(target)
    stale_payload = {
        "version": 1,
        "host": pipeline._lock_host(),
        "pid": 424242,
        "created_at": 1.0,
    }
    live_payload = {
        "version": 1,
        "host": pipeline._lock_host(),
        "pid": os.getpid(),
        "created_at": 2.0,
    }
    stale_raw = (json.dumps(stale_payload, sort_keys=True) + "\n").encode()
    live_raw = (json.dumps(live_payload, sort_keys=True) + "\n").encode()
    lock.write_bytes(stale_raw)
    monkeypatch.setattr(pipeline, "_process_is_alive", lambda _pid: False)

    original_read_bytes = Path.read_bytes
    changed_read = False

    def report_changed_quarantine(path: Path) -> bytes:
        nonlocal changed_read
        raw = original_read_bytes(path)
        if path.name.startswith(f"{lock.name}.stale-") and not changed_read:
            changed_read = True
            return raw + b"changed-after-isolation"
        return raw

    monkeypatch.setattr(Path, "read_bytes", report_changed_quarantine)

    original_exists = pipeline._path_exists
    vacancy_checked = False

    def observe_vacancy(path: Path) -> bool:
        nonlocal vacancy_checked
        result = original_exists(path)
        if path == lock and not result:
            vacancy_checked = True
        return result

    monkeypatch.setattr(pipeline, "_path_exists", observe_vacancy)

    original_replace = Path.replace
    original_link = pipeline.os.link
    live_created = False

    def create_live_lock() -> None:
        nonlocal live_created
        if not live_created:
            lock.write_bytes(live_raw)
            live_created = True

    def race_on_replace(source: Path, destination: str | Path) -> Path:
        if (
            vacancy_checked
            and source.name.startswith(f"{lock.name}.stale-")
            and Path(destination) == lock
        ):
            create_live_lock()
        return original_replace(source, destination)

    def race_on_link(source: str, destination: str) -> None:
        if Path(destination) == lock:
            create_live_lock()
        return original_link(source, destination)

    monkeypatch.setattr(Path, "replace", race_on_replace)
    monkeypatch.setattr(pipeline.os, "link", race_on_link)
    warning_sink: list[str] = []

    with pytest.warns(RuntimeWarning, match="appeared while restoring"):
        with pytest.raises(FileExistsError, match="changed while being isolated"):
            pipeline._recover_stale_transaction_lock(lock, warning_sink)

    assert lock.read_bytes() == live_raw
    quarantines = sorted(tmp_path.glob(f"{lock.name}.stale-*"))
    assert len(quarantines) == 1
    assert quarantines[0].read_bytes() == stale_raw
    assert any("appeared while restoring" in message for message in warning_sink)


def test_lock_snapshot_identity_is_rename_stable_and_replacement_sensitive(
    tmp_path: Path,
) -> None:
    import diffractscout.pipeline as pipeline

    lock = tmp_path / "identity-lock"
    lock.write_bytes(b"original-lock")
    original_raw, original_identity = pipeline._lock_snapshot(lock)

    renamed = tmp_path / "identity-lock-renamed"
    lock.replace(renamed)
    renamed_raw, renamed_identity = pipeline._lock_snapshot(renamed)
    assert renamed_raw == original_raw
    assert renamed_identity == original_identity

    replacement = tmp_path / "identity-lock-replacement"
    replacement.write_bytes(b"replaced-lock")
    replacement.replace(renamed)
    replaced_raw, replaced_identity = pipeline._lock_snapshot(renamed)
    assert replaced_raw != original_raw
    if original_identity[:2] != (0, 0):
        assert replaced_identity != original_identity


def test_lock_release_preserves_replacement_after_metadata_read(
    tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    target = tmp_path / "release-race-target"
    lock = _lock_path_for(target)
    lock.write_text(
        json.dumps(
            {
                "version": 1,
                "host": pipeline._lock_host(),
                "pid": os.getpid(),
                "created_at": 1.0,
            }
        ),
        encoding="utf-8",
    )
    replacement = tmp_path / "replacement-lock"
    original_read = pipeline._read_lock_metadata

    def read_then_replace(path: Path) -> dict[str, object]:
        metadata = original_read(path)
        replacement.write_text(
            json.dumps(
                {
                    "version": 1,
                    "host": pipeline._lock_host(),
                    "pid": os.getpid(),
                    "created_at": 2.0,
                }
            ),
            encoding="utf-8",
        )
        replacement.replace(path)
        return metadata

    monkeypatch.setattr(pipeline, "_read_lock_metadata", read_then_replace)
    warning_sink: list[str] = []
    with pytest.warns(RuntimeWarning, match="changed during release"):
        _release_transaction_lock(lock, warning_sink=warning_sink)

    assert lock.exists()
    assert json.loads(lock.read_text(encoding="utf-8"))["created_at"] == 2.0
    assert any("changed during release" in message for message in warning_sink)


def test_lock_release_preserves_replacement_after_final_snapshot(
    tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    target = tmp_path / "release-final-race-target"
    lock = _lock_path_for(target)
    lock.write_text(
        json.dumps(
            {
                "version": 1,
                "host": pipeline._lock_host(),
                "pid": os.getpid(),
                "created_at": 1.0,
            }
        ),
        encoding="utf-8",
    )
    replacement = tmp_path / "replacement-final-lock"
    original_snapshot = pipeline._lock_snapshot
    snapshot_calls = 0

    def snapshot_then_replace(path: Path) -> tuple[bytes, tuple[int, int, int]]:
        nonlocal snapshot_calls
        snapshot = original_snapshot(path)
        snapshot_calls += 1
        if snapshot_calls == 2:
            replacement.write_text(
                json.dumps(
                    {
                        "version": 1,
                        "host": pipeline._lock_host(),
                        "pid": os.getpid(),
                        "created_at": 3.0,
                    }
                ),
                encoding="utf-8",
            )
            replacement.replace(path)
        return snapshot

    monkeypatch.setattr(pipeline, "_lock_snapshot", snapshot_then_replace)
    warning_sink: list[str] = []
    with pytest.warns(RuntimeWarning, match="changed while being isolated"):
        _release_transaction_lock(lock, warning_sink=warning_sink)

    assert snapshot_calls >= 3
    assert lock.exists()
    assert json.loads(lock.read_text(encoding="utf-8"))["created_at"] == 3.0
    assert any("changed while being isolated" in message for message in warning_sink)


def test_lock_cleanup_failure_does_not_mask_success_or_primary_error(
    tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    output = tmp_path / "cleanup-success"
    staging = tmp_path / "success-staging"
    staging.mkdir()

    def fail_release(*_args, **_kwargs) -> None:
        raise PermissionError("simulated lock cleanup failure")

    monkeypatch.setattr(pipeline, "_release_transaction_lock", fail_release)
    with pytest.warns(RuntimeWarning, match="lock cleanup"):
        _commit_staging_output(output, staging, expected_state=_TargetState(False))
    assert output.is_dir()
    assert _lock_path_for(output).exists()

    primary_output = tmp_path / "cleanup-primary"
    primary_staging = tmp_path / "primary-staging"
    primary_staging.mkdir()
    def fail_staging_rename(source: Path, target: Path) -> None:
        if source == primary_staging:
            raise RuntimeError("primary publish failure")
        raise AssertionError("unexpected publication source")

    monkeypatch.setattr(pipeline, "_rename_directory_noreplace", fail_staging_rename)
    with pytest.raises(RuntimeError, match="primary publish failure"):
        _commit_staging_output(
            primary_output,
            primary_staging,
            expected_state=_TargetState(False),
        )
    assert not primary_output.exists()


def test_publication_primitive_target_race_preserves_external_target_and_backup(
    demo_inputs: Path, tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    output = tmp_path / "publication-race"
    analyze_cifs([demo_inputs], output, include_excel=False)
    expected_state = _capture_target_state(output)
    staging = tmp_path / "publication-staging"
    staging.mkdir()

    def publish_then_occupy(source: Path, target: Path) -> None:
        assert source == staging
        target.mkdir()
        raise FileExistsError(errno.EEXIST, "target appeared at publication")

    monkeypatch.setattr(pipeline, "_rename_directory_noreplace", publish_then_occupy)
    with pytest.raises(FileExistsError, match="target appeared"):
        _commit_staging_output(
            output,
            staging,
            expected_state=expected_state,
        )

    assert output.is_dir()
    assert list(output.iterdir()) == []
    assert staging.is_dir()
    backups = sorted(tmp_path.glob(".publication-race.backup-*"))
    assert len(backups) == 1
    assert verify_bundle(backups[0])["ok"]


def test_windows_publication_normalizes_existing_error(
    tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    error = OSError(5, "access denied")
    error.winerror = 183  # type: ignore[attr-defined]
    monkeypatch.setattr(pipeline.os, "rename", lambda *_args: (_ for _ in ()).throw(error))
    with pytest.raises(FileExistsError):
        pipeline._rename_directory_noreplace_windows(
            tmp_path / "staging", tmp_path / "target"
        )


def test_linux_publication_uses_renameat2_and_normalizes_errno(
    tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    class FakeRenameat2:
        argtypes = None
        restype = None

        def __call__(self, *_args) -> int:
            ctypes.set_errno(errno.EEXIST)
            return -1

    class FakeLibc:
        renameat2 = FakeRenameat2()

    monkeypatch.setattr(pipeline.ctypes, "CDLL", lambda *_args, **_kwargs: FakeLibc())
    with pytest.raises(FileExistsError):
        pipeline._rename_directory_noreplace_linux(
            tmp_path / "staging", tmp_path / "target"
        )
    assert FakeLibc.renameat2.argtypes[-1] is pipeline.ctypes.c_uint
    assert FakeLibc.renameat2.restype is pipeline.ctypes.c_int


def test_linux_publication_fails_closed_without_libc_primitive(
    tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    class MissingLibc:
        def __getattr__(self, _name: str) -> object:
            raise AttributeError("missing primitive")

    monkeypatch.setattr(pipeline.ctypes, "CDLL", lambda *_args, **_kwargs: MissingLibc())
    with pytest.raises(OSError, match="refusing a racy publication"):
        pipeline._rename_directory_noreplace_linux(
            tmp_path / "staging", tmp_path / "target"
        )


def test_macos_publication_uses_renamex_np_exclusive_flag(
    tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    class FakeRenamex:
        argtypes = None
        restype = None
        flags: int | None = None

        def __call__(self, _source: bytes, _target: bytes, flags: int) -> int:
            self.flags = flags
            ctypes.set_errno(errno.EEXIST)
            return -1

    class FakeLibc:
        renamex_np = FakeRenamex()

    monkeypatch.setattr(pipeline.ctypes, "CDLL", lambda *_args, **_kwargs: FakeLibc())
    with pytest.raises(FileExistsError):
        pipeline._rename_directory_noreplace_macos(
            tmp_path / "staging", tmp_path / "target"
        )
    assert FakeLibc.renamex_np.flags == 0x00000004


def test_unsupported_posix_publication_fails_closed(
    tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    monkeypatch.setattr(pipeline.os, "name", "posix")
    monkeypatch.setattr(pipeline.sys, "platform", "freebsd")
    with pytest.raises(OSError, match="refusing a racy rename"):
        pipeline._rename_directory_noreplace(
            tmp_path / "staging", tmp_path / "target"
        )


def test_provider_copy_falls_back_without_hardlink(
    tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    source = tmp_path / "source.cif"
    source.write_bytes(b"complete provider artifact")
    destination = tmp_path / "inputs"
    original_link = pipeline.os.link

    def no_hardlink(*_args, **_kwargs) -> None:
        raise OSError(errno.EOPNOTSUPP, "hardlinks unavailable")

    monkeypatch.setattr(pipeline.os, "link", no_hardlink)
    published = pipeline._copy_provider_artifact(
        source,
        destination,
        preferred_name="source.cif",
        suffix=".cif",
    )

    assert published.read_bytes() == source.read_bytes()
    assert not list(destination.glob("*.tmp"))
    monkeypatch.setattr(pipeline.os, "link", original_link)


def test_dangling_probe_is_not_captured_as_vacant_and_commit_rejects_reparse(
    tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    class DanglingProbe:
        def stat(self, *, follow_symlinks: bool = True) -> object:
            raise FileNotFoundError

        def lstat(self) -> object:
            return SimpleNamespace(
                st_mode=0o120777,
                st_dev=1,
                st_ino=2,
                st_ctime_ns=3,
                st_mtime_ns=4,
            )

        @staticmethod
        def is_symlink() -> bool:
            return True

    captured = pipeline._capture_target_state(DanglingProbe())  # type: ignore[arg-type]
    assert captured.exists is True
    assert captured.reparse is True

    target = tmp_path / "reparse-race"
    staging = tmp_path / "reparse-staging"
    staging.mkdir()
    monkeypatch.setattr(
        pipeline,
        "_capture_target_state",
        lambda _path: _TargetState(exists=True, reparse=True),
    )
    with pytest.raises(FileExistsError, match="symbolic link or reparse"):
        _commit_staging_output(target, staging, expected_state=_TargetState(False))


def test_local_cif_copy_fails_closed_when_source_mutates(
    demo_inputs: Path, tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    source = demo_inputs / "synthetic_fcc_al.cif"
    destination = tmp_path / "staged-inputs"
    original_copy = pipeline.shutil.copy2

    def copy_then_mutate(source_path: str | Path, target: str | Path, **kwargs) -> Path:
        result = original_copy(source_path, target, **kwargs)
        if Path(source_path) == source:
            source.write_bytes(source.read_bytes() + b"\n# source changed\n")
        return result

    monkeypatch.setattr(pipeline.shutil, "copy2", copy_then_mutate)
    with pytest.raises(RuntimeError, match="source changed"):
        _copy_local_input(source, destination, include_elasticity=False)
    assert not list(destination.glob("*.cif"))
    assert not list(destination.glob("*.tmp"))


def test_local_elasticity_tensor_comes_from_verified_sidecar_snapshot(
    demo_inputs: Path, tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    source = demo_inputs / "synthetic_fcc_al.cif"
    sidecar = demo_inputs / "synthetic_fcc_al_elasticity.json"
    destination = tmp_path / "staged-inputs"
    original_copy = pipeline._copy_local_snapshot
    original_payload = sidecar.read_text(encoding="utf-8")
    expected_value = json.loads(original_payload)["stiffness_GPa"][0][0]
    changed_payload = original_payload.replace(str(expected_value), "999.0", 1)

    def copy_then_mutate(source_path: Path, target: Path, **kwargs: object) -> Path:
        result = original_copy(source_path, target, **kwargs)
        if source_path == sidecar:
            sidecar.write_text(changed_payload, encoding="utf-8")
        return result

    monkeypatch.setattr(pipeline, "_copy_local_snapshot", copy_then_mutate)
    staged_cif, tensor = _copy_local_input(
        source,
        destination,
        include_elasticity=True,
    )

    assert tensor is not None
    assert tensor.stiffness_GPa[0, 0] == pytest.approx(expected_value)
    assert tensor.raw_payload_path == destination / "synthetic_fcc_al_elasticity.json"
    assert json.loads(tensor.raw_payload_path.read_text(encoding="utf-8"))["stiffness_GPa"][0][0] == pytest.approx(expected_value)
    assert staged_cif.read_bytes() == source.read_bytes()


def test_local_elasticity_sidecar_mutation_fails_closed(
    demo_inputs: Path, tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    source = demo_inputs / "synthetic_fcc_al.cif"
    sidecar = demo_inputs / "synthetic_fcc_al_elasticity.json"
    destination = tmp_path / "staged-inputs"
    original_copy2 = pipeline.shutil.copy2

    def copy_then_mutate(source_path: str | Path, target: str | Path, **kwargs: object) -> Path:
        result = original_copy2(source_path, target, **kwargs)
        if Path(source_path) == sidecar:
            sidecar.write_text(sidecar.read_text(encoding="utf-8") + "\n", encoding="utf-8")
        return result

    monkeypatch.setattr(pipeline.shutil, "copy2", copy_then_mutate)
    with pytest.raises(RuntimeError, match="source changed"):
        _copy_local_input(source, destination, include_elasticity=True)
    assert not list(destination.glob("*.cif"))
    assert not list(destination.glob("*_elasticity.*"))
    assert not list(destination.glob("*.tmp"))


def test_provider_normalization_ignores_post_snapshot_source_mutation(
    tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    source_dir = tmp_path / "provider-source"
    source_dir.mkdir()
    source_cif = source_dir / "phase.cif"
    source_cif.write_text("data_phase\n_cell_length_a 1\n", encoding="utf-8")
    source_sidecar = source_dir / "phase_elasticity.json"
    payload = {
        "cif_filename": "phase.cif",
        "material_id": "mp-1",
        "coordinate_frame": "crystal_cartesian_from_cif_lattice",
        "stiffness_GPa": [[100.0 if i == j else 0.0 for j in range(6)] for i in range(6)],
    }
    source_sidecar.write_text(json.dumps(payload), encoding="utf-8")
    item = DownloadArtifact(
        candidate=CandidateRecord(material_id="mp-1", formula="Al"),
        cif_path=source_cif,
        elasticity_path=source_sidecar,
    )
    destination = tmp_path / "inputs"
    original_copy = pipeline._copy_provider_artifact

    def copy_then_mutate(source: str | Path, *args: object, **kwargs: object) -> Path:
        result = original_copy(source, *args, **kwargs)
        if Path(source) == source_sidecar:
            changed = dict(payload)
            changed["stiffness_GPa"] = [
                [999.0 if i == j else 0.0 for j in range(6)] for i in range(6)
            ]
            source_sidecar.write_text(json.dumps(changed), encoding="utf-8")
        return result

    monkeypatch.setattr(pipeline, "_copy_provider_artifact", copy_then_mutate)
    normalized, diagnostics = _normalize_download_artifacts(
        [item],
        destination,
        include_elasticity=True,
    )

    assert not [item for item in diagnostics if item.level == "error"]
    assert normalized[0].elasticity_path is not None
    assert normalized[0].elasticity_status == "ok"
    staged_payload = json.loads(normalized[0].elasticity_path.read_text(encoding="utf-8"))
    assert staged_payload["stiffness_GPa"][0][0] == pytest.approx(100.0)


def _write_override_sidecar(path: Path, cif_name: str, tensor: object) -> None:
    matrix = tensor.stiffness_GPa.tolist()  # type: ignore[union-attr]
    coordinate_frame = tensor.coordinate_frame  # type: ignore[union-attr]
    path.write_text(
        json.dumps(
            {
                "status": "ok",
                "cif_filename": cif_name,
                "coordinate_frame": coordinate_frame,
                "stiffness_GPa": matrix,
                "provenance": {
                    "provider": "sidecar-provider",
                    "coordinate_frame": coordinate_frame,
                    "paired_cif": cif_name,
                },
            }
        ),
        encoding="utf-8",
    )


def test_local_elastic_override_sidecar_matrix_mismatch_fails_closed(
    demo_inputs: Path, tmp_path: Path
) -> None:
    source = demo_inputs / "synthetic_fcc_al.cif"
    sidecar = tmp_path / "explicit_elasticity.json"
    override = parse_cubic_cij(111.0, 22.0, 33.0, source="explicit-provider")
    staged_sidecar_tensor = parse_cubic_cij(222.0, 44.0, 66.0, source="sidecar-provider")
    _write_override_sidecar(sidecar, source.name, staged_sidecar_tensor)
    override.raw_payload_path = sidecar
    destination = tmp_path / "override-mismatch-inputs"

    with pytest.raises(ValueError, match="stiffness matrix mismatch"):
        _copy_local_input(
            source,
            destination,
            include_elasticity=True,
            elastic_override=override,
        )
    assert not list(destination.glob("*.cif"))
    assert not list(destination.glob("*_elasticity.*"))
    assert not list(destination.glob("*.tmp"))


def test_local_elastic_override_keeps_provenance_and_rebinds_staged_sidecar(
    demo_inputs: Path, tmp_path: Path
) -> None:
    source = demo_inputs / "synthetic_fcc_al.cif"
    sidecar = tmp_path / "explicit_elasticity.json"
    override = parse_cubic_cij(111.0, 22.0, 33.0, source="explicit-provider")
    _write_override_sidecar(sidecar, source.name, override)
    override.raw_payload_path = sidecar
    destination = tmp_path / "override-matching-inputs"

    staged_cif, tensor = _copy_local_input(
        source,
        destination,
        include_elasticity=True,
        elastic_override=override,
    )

    assert tensor is not None
    assert tensor.source_record_id == "explicit-provider"
    assert tensor.nature_of_data == "explicit-provider"
    assert tensor.stiffness_GPa[0, 0] == pytest.approx(111.0)
    assert tensor.raw_payload_path == destination / "synthetic_fcc_al_elasticity.json"
    assert tensor.raw_payload_path.is_file()
    assert staged_cif.is_file()


def test_json_cleanup_failure_warns_without_masking_success(
    tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.utils as utils

    output = tmp_path / "cleanup-warning.json"
    original_unlink = utils.Path.unlink

    def fail_temp_unlink(self: Path, *args, **kwargs) -> None:
        if self.parent == tmp_path and self.suffix == ".tmp":
            raise PermissionError("simulated temporary cleanup failure")
        original_unlink(self, *args, **kwargs)

    monkeypatch.setattr(utils.Path, "unlink", fail_temp_unlink)
    with pytest.warns(RuntimeWarning, match="JSON temporary file"):
        result = write_json(output, {"ok": True})
    assert result == output
    assert json.loads(output.read_text(encoding="utf-8")) == {"ok": True}


def test_concurrent_overwrite_writers_leave_at_most_one_committed_bundle(
    demo_inputs: Path, tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    output = tmp_path / "concurrent-writers"
    analyze_cifs([demo_inputs], output, include_excel=False)

    left = tmp_path / "left"
    right = tmp_path / "right"
    left.mkdir()
    right.mkdir()
    shutil.copy2(demo_inputs / "synthetic_fcc_al.cif", left / "left.cif")
    shutil.copy2(demo_inputs / "synthetic_fcc_al.cif", right / "right.cif")

    original_export = pipeline.export_result_bundle
    ready = threading.Barrier(2)

    def export_then_race(*args, **kwargs):
        result = original_export(*args, **kwargs)
        ready.wait(timeout=30)
        return result

    monkeypatch.setattr(pipeline, "export_result_bundle", export_then_race)

    def run(source: Path) -> object:
        try:
            return analyze_cifs(
                [source], output, include_excel=False, overwrite=True
            )
        except Exception as exc:  # the losing transaction is expected to fail closed
            return exc

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(run, (left, right)))

    assert sum(not isinstance(result, Exception) for result in results) == 1
    assert sum(isinstance(result, Exception) for result in results) == 1
    assert verify_bundle(output)["ok"]
    phase_summary = (output / "phase_summary.csv").read_text(encoding="utf-8-sig")
    assert "left.cif" in phase_summary or "right.cif" in phase_summary


def test_write_json_concurrent_writers_use_distinct_temps(
    tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.utils as utils

    output = tmp_path / "concurrent.json"
    original_replace = utils.Path.replace
    replace_barrier = threading.Barrier(2)
    replace_lock = threading.Lock()

    def synchronize_replacements(self: Path, target: str | Path) -> Path:
        if self.parent == tmp_path and self.suffix == ".tmp":
            replace_barrier.wait(timeout=30)
            with replace_lock:
                return original_replace(self, target)
        return original_replace(self, target)

    monkeypatch.setattr(utils.Path, "replace", synchronize_replacements)

    def write(index: int) -> Path:
        return write_json(output, {"index": index})

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(write, (1, 2)))

    assert all(path == output for path in results)
    assert json.loads(output.read_text(encoding="utf-8"))["index"] in {1, 2}
    assert not list(tmp_path.glob("*.tmp"))


def test_write_json_encoding_failure_cleans_temporary_file(tmp_path: Path) -> None:
    output = tmp_path / "encoding-failure.json"
    with pytest.raises(UnicodeEncodeError):
        write_json(output, {"invalid": "\udcff"})
    assert not output.exists()
    assert not list(tmp_path.glob("*.tmp"))


def test_staged_cif_modified_after_analysis_fails_closed(
    demo_inputs: Path, tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    output = tmp_path / "changed-staging-cif"
    original_export = pipeline.export_result_bundle

    def modify_before_export(staging: Path, *args, **kwargs):
        cif = next((staging / "inputs").glob("*.cif"))
        cif.write_bytes(cif.read_bytes() + b"\n# changed after analysis\n")
        return original_export(staging, *args, **kwargs)

    monkeypatch.setattr(pipeline, "export_result_bundle", modify_before_export)
    with pytest.raises(RuntimeError, match="CIF.*changed"):
        analyze_cifs([demo_inputs], output, include_excel=False)

    assert not output.exists()
    assert not list(tmp_path.glob(".changed-staging-cif.diffractscout-*"))


def test_remote_provider_is_not_contacted_when_output_is_unsafe(tmp_path: Path) -> None:
    class Provider:
        name = "not-called"

        def search_subsystem(self, *_args, **_kwargs):
            raise AssertionError("provider should not be contacted")

        def download_candidates(self, *_args, **_kwargs):
            raise AssertionError("provider should not be contacted")

        def metadata(self):
            raise AssertionError("provider should not be contacted")

    output = tmp_path / "occupied"
    output.mkdir()
    (output / "user.txt").write_text("keep", encoding="utf-8")

    with pytest.raises(FileExistsError, match="not empty"):
        run_pipeline("Ti-Al", Provider(), output)
    assert (output / "user.txt").read_text(encoding="utf-8") == "keep"


def _real_directory_reparse_point(tmp_path: Path) -> tuple[Path, callable]:
    target = tmp_path / "junction-target"
    target.mkdir()
    link = tmp_path / "junction-like"
    if os.name == "nt":
        powershell = shutil.which("powershell") or shutil.which("pwsh")
        if powershell is None:
            pytest.skip("PowerShell is unavailable; cannot create a real Windows junction")
        environment = os.environ.copy()
        environment["DIFRACTSCOUT_JUNCTION_TARGET"] = str(target)
        environment["DIFRACTSCOUT_JUNCTION_LINK"] = str(link)
        command = (
            "$ErrorActionPreference='Stop'; "
            "New-Item -ItemType Junction "
            "-Path $env:DIFRACTSCOUT_JUNCTION_LINK "
            "-Target $env:DIFRACTSCOUT_JUNCTION_TARGET | Out-Null"
        )
        completed = subprocess.run(
            [powershell, "-NoProfile", "-NonInteractive", "-Command", command],
            env=environment,
            capture_output=True,
            text=True,
            check=False,
        )
        if completed.returncode != 0:
            pytest.skip(
                "Windows junction creation unavailable: "
                + (completed.stderr or completed.stdout).strip()
            )

        def cleanup() -> None:
            subprocess.run(
                [
                    powershell,
                    "-NoProfile",
                    "-NonInteractive",
                    "-Command",
                    "Remove-Item -LiteralPath $env:DIFRACTSCOUT_JUNCTION_LINK -Force",
                ],
                env=environment,
                capture_output=True,
                text=True,
                check=False,
            )

        return link, cleanup
    try:
        link.symlink_to(target, target_is_directory=True)
    except (OSError, NotImplementedError) as exc:
        pytest.skip(f"POSIX directory symlink creation unavailable: {exc}")
    return link, link.unlink


def test_pipeline_rejects_real_reparse_point_final_and_parent(tmp_path: Path) -> None:
    import diffractscout.pipeline as pipeline

    reparse_component, cleanup = _real_directory_reparse_point(tmp_path)
    try:
        assert pipeline._is_reparse_point(reparse_component)
        for output in (reparse_component, reparse_component / "nested" / "bundle"):
            with pytest.raises(FileExistsError, match="reparse-point"):
                pipeline._validate_output_target(output, overwrite=False)
    finally:
        cleanup()


def test_pipeline_reparse_detector_reads_raw_windows_attribute() -> None:
    import diffractscout.pipeline as pipeline

    class RawAttributePath:
        def stat(self, *, follow_symlinks: bool = True) -> object:
            assert follow_symlinks is False
            return SimpleNamespace(st_file_attributes=0x0400)

        @staticmethod
        def is_symlink() -> bool:
            return False

    assert pipeline._is_reparse_point(RawAttributePath())  # type: ignore[arg-type]


def test_elastic_override_prefers_canonical_path_and_rejects_legacy_ambiguity(
    demo_inputs: Path,
) -> None:
    cif = (demo_inputs / "synthetic_fcc_al.cif").resolve()
    canonical = object()
    legacy_filename = object()
    legacy_stem = object()
    assert _lookup_elastic_override(cif, {str(cif): canonical}) is canonical  # type: ignore[arg-type]
    assert _lookup_elastic_override(cif, {cif.name: legacy_filename}) is legacy_filename  # type: ignore[arg-type]
    assert _lookup_elastic_override(cif, {cif.stem: legacy_stem}) is legacy_stem  # type: ignore[arg-type]
    assert (
        _lookup_elastic_override(
            cif,
            {cif.name: legacy_filename, cif.stem: legacy_stem},
        )
        is None
    )
