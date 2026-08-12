"""Lab-view Excel sheets and expanded peak/pattern export columns."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook

from diffractscout.demo import write_demo_inputs
from diffractscout.export_views import (
    ANALYSIS_PEAK_COLUMNS,
    BEGINNER_PEAK_HEADERS_ZH,
    beginner_peak_rows_zh,
    user_guide_rows,
)
from diffractscout.exporters import PEAK_HEADERS, peak_rows, write_excel_workbook
from diffractscout.models import AnalysisSettings
from diffractscout.pipeline import analyze_cifs


def test_demo_analyze_writes_lab_view_sheets(tmp_path: Path) -> None:
    inputs = write_demo_inputs(tmp_path / "inputs")
    result = analyze_cifs(
        [inputs],
        tmp_path / "bundle",
        settings=AnalysisSettings(export_lab_views=True),
        include_excel=True,
    )
    workbook_path = result.output_dir / "results.xlsx"
    assert workbook_path.is_file()
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    assert "推荐峰表" in workbook.sheetnames
    assert "使用说明" in workbook.sheetnames
    assert "Peaks" in workbook.sheetnames
    # Lab-first workbook opens on the Chinese analysis long table.
    assert workbook.active.title == "推荐峰表"
    recommend = workbook["推荐峰表"]
    headers = [cell.value for cell in next(recommend.iter_rows(min_row=1, max_row=1))]
    assert headers[0] == "物相名称"
    assert "相对强度_相内max100" in headers
    assert "2θ_当前_deg" in headers
    assert "2θ_CuKa_deg" in headers
    assert "体积归一强度J_含LP_R_hkl别名" in headers
    guide = workbook["使用说明"]
    guide_rows = list(guide.iter_rows(min_row=1, max_row=3, values_only=True))
    assert guide_rows[0][0] == "项目"
    joined = "\n".join(
        str(cell or "") for row in guide.iter_rows(values_only=True) for cell in row
    )
    assert "残差" in joined
    assert "不是" in joined


def test_peak_and_recommend_sheets_have_freeze_and_autofilter(tmp_path: Path) -> None:
    inputs = write_demo_inputs(tmp_path / "inputs")
    result = analyze_cifs(
        [inputs],
        tmp_path / "bundle",
        settings=AnalysisSettings(export_lab_views=True, include_elasticity=True),
        include_excel=True,
    )
    workbook = load_workbook(result.output_dir / "results.xlsx", data_only=False)
    for sheet_name in ("Peaks", "推荐峰表"):
        sheet = workbook[sheet_name]
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref
        assert sheet.auto_filter.ref.startswith("A1:")
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        assert headers[0] in {"phase_name", "物相名称"}
        assert sheet.max_row >= 2


def test_analysis_peak_columns_present_in_peaks_and_csv(tmp_path: Path) -> None:
    inputs = write_demo_inputs(tmp_path / "inputs")
    result = analyze_cifs([inputs], tmp_path / "bundle", include_excel=True)
    for col in ANALYSIS_PEAK_COLUMNS:
        assert col in PEAK_HEADERS
    workbook = load_workbook(result.output_dir / "results.xlsx", data_only=False)
    peak_headers = [cell.value for cell in next(workbook["Peaks"].iter_rows(min_row=1, max_row=1))]
    for col in ANALYSIS_PEAK_COLUMNS:
        assert col in peak_headers
    with (result.output_dir / "peak_reference.csv").open(encoding="utf-8-sig", newline="") as handle:
        fieldnames = list(csv.DictReader(handle).fieldnames or [])
    for col in ANALYSIS_PEAK_COLUMNS:
        assert col in fieldnames
    # Analysis-first ordering: identity and geometry before deep SF extras.
    assert peak_headers.index("phase_name") < peak_headers.index("d_spacing_A")
    assert peak_headers.index("normalized_intensity") < peak_headers.index("structure_factor_sq")


def test_peak_reference_has_two_theta_cu_ka_column(tmp_path: Path) -> None:
    inputs = write_demo_inputs(tmp_path / "inputs")
    result = analyze_cifs([inputs], tmp_path / "bundle", include_excel=False)
    peak_path = result.output_dir / "peak_reference.csv"
    with peak_path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fieldnames = list(reader.fieldnames or [])
        rows = list(reader)
    assert "two_theta_cu_ka_deg" in fieldnames
    assert "phase_relative_R_hkl_pct" in fieldnames
    assert "inverse_R_hkl" in fieldnames
    assert "r_hkl_model_note" in fieldnames
    assert rows
    first = rows[0]
    assert first["two_theta_cu_ka_deg"]
    assert float(first["two_theta_cu_ka_deg"]) > 0


def test_pattern_profiles_include_d_axis_columns(tmp_path: Path) -> None:
    inputs = write_demo_inputs(tmp_path / "inputs")
    result = analyze_cifs([inputs], tmp_path / "bundle", include_excel=False)
    pattern_path = result.output_dir / "pattern_profiles.csv"
    with pattern_path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fieldnames = list(reader.fieldnames or [])
        rows = list(reader)
    assert "d_A" in fieldnames or "d_spacing" in fieldnames or "d_spacing_A" in fieldnames
    assert "two_theta_deg" in fieldnames
    assert "q_invA" in fieldnames
    assert "g_invA" in fieldnames
    assert "x_axis_mode" in fieldnames
    assert "x" in fieldnames
    assert "relative_intensity" in fieldnames
    assert rows
    sample = rows[len(rows) // 2]
    assert sample["d_A"]
    assert float(sample["d_A"]) > 0
    assert sample["x_axis_mode"] == "two_theta"


def test_include_patterns_false_skips_profile_export(tmp_path: Path) -> None:
    inputs = write_demo_inputs(tmp_path / "inputs")
    result = analyze_cifs(
        [inputs],
        tmp_path / "bundle",
        settings=AnalysisSettings(include_patterns=False, include_elasticity=False),
        include_excel=True,
    )
    assert not (result.output_dir / "pattern_profiles.csv").exists()
    workbook = load_workbook(result.output_dir / "results.xlsx", read_only=True)
    assert "Patterns" not in workbook.sheetnames
    assert "Peaks" in workbook.sheetnames


def test_export_lab_views_false_omits_chinese_sheets(tmp_path: Path) -> None:
    inputs = write_demo_inputs(tmp_path / "inputs")
    result = analyze_cifs(
        [inputs],
        tmp_path / "bundle",
        settings=AnalysisSettings(export_lab_views=False, include_elasticity=False),
        include_excel=True,
    )
    workbook = load_workbook(result.output_dir / "results.xlsx", read_only=True)
    assert "推荐峰表" not in workbook.sheetnames
    assert "使用说明" not in workbook.sheetnames
    assert "Peaks" in workbook.sheetnames


def test_multi_phase_peak_rows_keep_phase_name_for_filter(tmp_path: Path) -> None:
    """Shipped peak_rows + Excel Peaks support multi-phase filter via phase_name."""

    demo_dir = write_demo_inputs(tmp_path / "inputs")
    # Second synthetic phase: copy the demo CIF under another stem for multi-phase.
    src = next(demo_dir.glob("*.cif"))
    second = demo_dir / "phase_b.cif"
    second.write_bytes(src.read_bytes())
    # Pair elasticity only for first if present; analysis still works.
    result = analyze_cifs(
        [demo_dir],
        tmp_path / "bundle",
        settings=AnalysisSettings(export_lab_views=True, include_elasticity=False),
        include_excel=True,
    )
    assert len(result.analyses) >= 2
    rows = peak_rows(result.analyses)
    phases = {row["phase_name"] for row in rows}
    assert len(phases) >= 2
    workbook = load_workbook(result.output_dir / "results.xlsx", data_only=False)
    peak_headers = [cell.value for cell in next(workbook["Peaks"].iter_rows(min_row=1, max_row=1))]
    assert peak_headers[0] == "phase_name"
    rec_headers = [cell.value for cell in next(workbook["推荐峰表"].iter_rows(min_row=1, max_row=1))]
    assert rec_headers[0] == "物相名称"
    # All peak data rows carry a non-empty phase identity for Excel filter.
    for sheet_name, key in (("Peaks", "phase_name"), ("推荐峰表", "物相名称")):
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        col = headers.index(key) + 1
        values = [sheet.cell(row=r, column=col).value for r in range(2, sheet.max_row + 1)]
        assert values
        assert all(v not in (None, "") for v in values)


def test_beginner_and_guide_helpers_are_stable() -> None:
    assert BEGINNER_PEAK_HEADERS_ZH["物相名称"] == "phase_name"
    assert BEGINNER_PEAK_HEADERS_ZH["2θ_CuKa_deg"] == "two_theta_cu_ka_deg"
    mapped = beginner_peak_rows_zh(
        [{"phase_name": "Al", "normalized_intensity": 100.0, "two_theta_cu_ka_deg": 38.0}]
    )
    assert mapped[0]["物相名称"] == "Al"
    assert mapped[0]["相对强度_相内max100"] == 100.0
    guide = user_guide_rows()
    assert guide[0] == ["项目", "说明"]
    joined = "\n".join(cell for row in guide for cell in row)
    assert "不是" in joined and "残差" in joined
    assert "QPA" in joined
    assert "筛选物相" in joined


def test_user_guide_rank_and_theta_wording_is_consistent() -> None:
    """Guide must match rank_by_*=1 strongest and 2θ=2×θ (shipped user_guide_rows)."""

    guide = user_guide_rows()
    # θ explanation must state double-angle relation, not tautology.
    theta_rows = [row for row in guide if row and "theta_deg" in str(row[0])]
    assert theta_rows
    assert "2×θ" in theta_rows[0][1] or "2*θ" in theta_rows[0][1]
    assert "2θ = 2θ" not in theta_rows[0][1]
    # Rank: 1 is strongest → smaller rank is stronger.
    rank_rows = [row for row in guide if row and "rank_by" in str(row[0])]
    assert rank_rows
    assert "1 为最强" in rank_rows[0][1] or "1为最强" in rank_rows[0][1]
    assert "数值越大排名越靠前" not in rank_rows[0][1]
    # How-to: strong peaks via relative intensity DESC or rank ASC / filter 1..N
    tip_rows = [row for row in guide if row and str(row[0]) == "找强峰"]
    assert tip_rows
    tip = tip_rows[0][1]
    assert "相对强度" in tip and "降序" in tip
    assert "rank_by_intensity 降序" not in tip
    assert "1..5" in tip or "1–5" in tip


def test_write_excel_workbook_freeze_filter_on_empty_peaks(tmp_path: Path) -> None:
    path = tmp_path / "empty.xlsx"
    write_excel_workbook(
        path,
        summary=[{"key": "k", "value": "v"}],
        phases=[],
        peaks=[],
        elasticity=[],
        candidates=[],
        downloads=[],
        diagnostics=[],
        patterns=[],
        export_lab_views=True,
        include_patterns=False,
    )
    workbook = load_workbook(path)
    peaks = workbook["Peaks"]
    assert peaks.freeze_panes == "A2"
    assert peaks.auto_filter.ref
