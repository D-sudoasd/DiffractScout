from pathlib import Path
import json
import shutil

import pytest
from openpyxl import load_workbook

from diffractscout.elasticity_input import parse_cubic_cij
from diffractscout.models import AnalysisSettings
from diffractscout.pipeline import analyze_cifs
from diffractscout.quick_export import build_parser as build_quick_export_parser
from diffractscout.quick_export import main as quick_export_main
from diffractscout.quick_export import _copy_excel_atomic, quick_export
from diffractscout.validation import verify_bundle


def test_quick_export_xlsx_writes_excel_and_bundle(demo_inputs: Path, tmp_path: Path) -> None:
    excel = tmp_path / "report.xlsx"
    result = quick_export([demo_inputs], excel)
    assert excel.is_file()
    bundle = tmp_path / "report_bundle"
    assert result.output_dir == bundle.resolve()
    assert (bundle / "manifest.json").is_file()
    assert (bundle / "results.xlsx").is_file()
    assert verify_bundle(bundle)["ok"]
    workbook = load_workbook(excel, read_only=True)
    assert "Peaks" in workbook.sheetnames


def test_quick_export_does_not_replace_existing_excel_without_authorization(
    demo_inputs: Path, tmp_path: Path
) -> None:
    excel = tmp_path / "existing.xlsx"
    excel.write_bytes(b"user workbook")

    with pytest.raises(FileExistsError, match="already exists"):
        quick_export([demo_inputs], excel)

    assert excel.read_bytes() == b"user workbook"
    assert not (tmp_path / "existing_bundle").exists()


def test_atomic_excel_copy_preserves_target_created_during_export(
    tmp_path: Path, monkeypatch
) -> None:
    source = tmp_path / "source.xlsx"
    target = tmp_path / "result.xlsx"
    source.write_bytes(b"completed-workbook")

    real_copy2 = shutil.copy2

    def racing_copy(source_path: Path, temporary_path: Path) -> Path:
        copied = real_copy2(source_path, temporary_path)
        target.write_bytes(b"external-file")
        return copied

    monkeypatch.setattr("diffractscout.quick_export.shutil.copy2", racing_copy)

    with pytest.raises(FileExistsError, match="created while exporting"):
        _copy_excel_atomic(source, target, overwrite=False)

    assert target.read_bytes() == b"external-file"
    assert not list(tmp_path.glob(".result.xlsx.*.tmp"))


def test_quick_export_can_atomically_replace_existing_excel_when_authorized(
    demo_inputs: Path, tmp_path: Path
) -> None:
    excel = tmp_path / "existing.xlsx"
    excel.write_bytes(b"old workbook")

    result = quick_export([demo_inputs], excel, overwrite=True)

    assert result.output_dir == (tmp_path / "existing_bundle").resolve()
    assert excel.read_bytes() != b"old workbook"
    assert load_workbook(excel, read_only=True).sheetnames


def test_quick_export_directory_mode(demo_inputs: Path, tmp_path: Path) -> None:
    output = tmp_path / "lab_bundle"
    result = quick_export([demo_inputs], output)
    assert result.output_dir == output.resolve()
    assert (output / "results.xlsx").is_file()
    assert result.analyses[0].metadata.get("export_lab_views") is True


def test_quick_export_keyword_overrides_apply_to_explicit_settings(
    demo_inputs: Path, tmp_path: Path
) -> None:
    result = quick_export(
        [demo_inputs],
        tmp_path / "overrides",
        settings=AnalysisSettings(step_deg=0.02),
        step_deg=0.05,
        include_excel=False,
    )
    assert result.analyses[0].metadata["step_deg"] == pytest.approx(0.05)


def test_quick_export_radiation_keywords_infer_physical_mode(
    demo_inputs: Path, tmp_path: Path
) -> None:
    energy_result = quick_export(
        [demo_inputs],
        tmp_path / "energy-inferred",
        energy_keV=20.0,
        include_excel=False,
    )
    energy_analysis = energy_result.analyses[0]
    energy_provenance = json.loads(
        (energy_result.output_dir / "provenance.json").read_text(encoding="utf-8")
    )
    assert energy_provenance["analysis_settings"]["input_mode"] == "energy"
    assert energy_analysis.energy_keV == pytest.approx(20.0)
    assert energy_analysis.wavelength_A == pytest.approx(12.398419843320026 / 20.0)
    assert energy_analysis.wavelength_source == "energy_keV"

    wavelength_result = quick_export(
        [demo_inputs],
        tmp_path / "wavelength-inferred",
        wavelength_A=1.2,
        include_excel=False,
    )
    wavelength_analysis = wavelength_result.analyses[0]
    wavelength_provenance = json.loads(
        (wavelength_result.output_dir / "provenance.json").read_text(encoding="utf-8")
    )
    assert wavelength_provenance["analysis_settings"]["input_mode"] == "wavelength"
    assert wavelength_analysis.wavelength_A == pytest.approx(1.2)
    assert wavelength_analysis.energy_keV == pytest.approx(12.398419843320026 / 1.2)
    assert wavelength_analysis.wavelength_source == "wavelength_A"


@pytest.mark.parametrize(
    ("case", "settings", "expected_wavelength", "expected_energy"),
    [
        (
            "energy-clears-stale-wavelength",
            AnalysisSettings(input_mode="energy", energy_keV=20.0, wavelength_A=1.2),
            None,
            pytest.approx(20.0),
        ),
        (
            "wavelength-clears-stale-energy",
            AnalysisSettings(input_mode="wavelength", wavelength_A=1.2, energy_keV=20.0),
            pytest.approx(1.2),
            None,
        ),
        (
            "builtin-source-clears-stale-radiation",
            AnalysisSettings(
                input_mode="source",
                source_preset="Cu Ka",
                wavelength_A=1.2,
                energy_keV=20.0,
            ),
            None,
            None,
        ),
        (
            "custom-source-keeps-wavelength",
            AnalysisSettings(
                input_mode="source",
                source_preset="Custom",
                wavelength_A=1.2,
                energy_keV=20.0,
            ),
            pytest.approx(1.2),
            None,
        ),
    ],
)
def test_quick_export_normalizes_direct_settings_before_validation(
    demo_inputs: Path,
    tmp_path: Path,
    case: str,
    settings: AnalysisSettings,
    expected_wavelength: object,
    expected_energy: object,
) -> None:
    """Inactive direct-settings radiation fields do not reach pipeline validation."""

    result = quick_export(
        [demo_inputs],
        tmp_path / case,
        settings=settings,
        include_excel=False,
    )

    persisted = json.loads(
        (result.output_dir / "provenance.json").read_text(encoding="utf-8")
    )["analysis_settings"]
    assert persisted["input_mode"] == settings.input_mode
    assert persisted["source_preset"] == settings.source_preset
    assert persisted["wavelength_A"] == expected_wavelength
    assert persisted["energy_keV"] == expected_energy


def test_quick_export_settings_radiation_override_is_applied(
    demo_inputs: Path, tmp_path: Path
) -> None:
    result = quick_export(
        [demo_inputs],
        tmp_path / "settings-energy-override",
        settings=AnalysisSettings(input_mode="energy", energy_keV=30.0),
        energy_keV=20.0,
        include_excel=False,
    )
    analysis = result.analyses[0]
    assert analysis.energy_keV == pytest.approx(20.0)
    assert analysis.wavelength_A == pytest.approx(12.398419843320026 / 20.0)


def test_quick_export_radiation_kw_overrides_default_settings_mode(
    demo_inputs: Path, tmp_path: Path
) -> None:
    result = quick_export(
        [demo_inputs],
        tmp_path / "default-settings-energy",
        settings=AnalysisSettings(step_deg=0.02),
        energy_keV=20.0,
        include_excel=False,
    )
    provenance = json.loads(
        (result.output_dir / "provenance.json").read_text(encoding="utf-8")
    )
    assert provenance["analysis_settings"]["input_mode"] == "energy"
    assert provenance["analysis_settings"]["wavelength_A"] is None
    assert provenance["analysis_settings"]["energy_keV"] == pytest.approx(20.0)


def test_quick_export_wavelength_kw_overrides_cross_baseline_mode(
    demo_inputs: Path, tmp_path: Path
) -> None:
    result = quick_export(
        [demo_inputs],
        tmp_path / "cross-baseline-wavelength",
        settings=AnalysisSettings(input_mode="energy", energy_keV=83.0),
        wavelength_A=1.2,
        include_excel=False,
    )
    provenance = json.loads(
        (result.output_dir / "provenance.json").read_text(encoding="utf-8")
    )
    assert provenance["analysis_settings"]["input_mode"] == "wavelength"
    assert provenance["analysis_settings"]["wavelength_A"] == pytest.approx(1.2)
    assert provenance["analysis_settings"]["energy_keV"] is None


def test_quick_export_stale_custom_source_does_not_override_energy_baseline(
    demo_inputs: Path, tmp_path: Path
) -> None:
    result = quick_export(
        [demo_inputs],
        tmp_path / "stale-custom-baseline",
        settings=AnalysisSettings(
            input_mode="energy",
            source_preset="Custom",
            energy_keV=20.0,
        ),
        wavelength_A=1.2,
        include_excel=False,
    )
    analysis = result.analyses[0]
    assert analysis.wavelength_source == "wavelength_A"
    assert analysis.wavelength_A == pytest.approx(1.2)
    assert json.loads(
        (result.output_dir / "provenance.json").read_text(encoding="utf-8")
    )["analysis_settings"]["input_mode"] == "wavelength"


def test_quick_export_operative_custom_baseline_keeps_source_mode(
    demo_inputs: Path, tmp_path: Path
) -> None:
    result = quick_export(
        [demo_inputs],
        tmp_path / "operative-custom-baseline",
        settings=AnalysisSettings(
            input_mode="source",
            source_preset="Custom",
            wavelength_A=1.0,
        ),
        wavelength_A=1.2,
        include_excel=False,
    )
    analysis = result.analyses[0]
    assert analysis.wavelength_source == "custom_source_wavelength"
    assert analysis.wavelength_A == pytest.approx(1.2)


def test_quick_export_explicit_custom_source_from_other_baseline(
    demo_inputs: Path, tmp_path: Path
) -> None:
    result = quick_export(
        [demo_inputs],
        tmp_path / "explicit-custom-source",
        settings=AnalysisSettings(input_mode="energy", energy_keV=20.0),
        source_preset="Custom",
        wavelength_A=1.2,
        include_excel=False,
    )
    analysis = result.analyses[0]
    assert analysis.wavelength_source == "custom_source_wavelength"
    assert analysis.wavelength_A == pytest.approx(1.2)


def test_quick_export_radiation_conflicts_fail_before_output(
    demo_inputs: Path, tmp_path: Path
) -> None:
    both_output = tmp_path / "both"
    with pytest.raises(ValueError, match="energy_keV.*wavelength_A"):
        quick_export(
            [demo_inputs],
            both_output,
            energy_keV=20.0,
            wavelength_A=1.2,
            include_excel=False,
        )
    assert not both_output.exists()

    cross_output = tmp_path / "explicit-mode-conflict"
    with pytest.raises(ValueError, match="input_mode.*wavelength_A"):
        quick_export(
            [demo_inputs],
            cross_output,
            settings=AnalysisSettings(input_mode="energy", energy_keV=20.0),
            wavelength_A=1.2,
            input_mode="energy",
            include_excel=False,
        )
    assert not cross_output.exists()


def test_quick_export_custom_to_builtin_source_clears_stale_fields(
    demo_inputs: Path, tmp_path: Path
) -> None:
    result = quick_export(
        [demo_inputs],
        tmp_path / "custom-to-cu",
        settings=AnalysisSettings(
            input_mode="source",
            source_preset="Custom",
            wavelength_A=1.2,
        ),
        source_preset="Cu Ka",
        include_excel=False,
    )
    provenance = json.loads(
        (result.output_dir / "provenance.json").read_text(encoding="utf-8")
    )
    merged = provenance["analysis_settings"]
    assert merged["input_mode"] == "source"
    assert merged["source_preset"] == "Cu Ka"
    assert merged["wavelength_A"] is None
    assert merged["energy_keV"] is None


def test_quick_export_builtin_source_override_keeps_explicit_wavelength(
    demo_inputs: Path, tmp_path: Path
) -> None:
    result = quick_export(
        [demo_inputs],
        tmp_path / "custom-to-cu-with-wavelength",
        settings=AnalysisSettings(
            input_mode="source",
            source_preset="Custom",
            wavelength_A=1.0,
        ),
        source_preset="Cu Ka",
        wavelength_A=1.2,
        include_excel=False,
    )
    analysis = result.analyses[0]
    provenance = json.loads(
        (result.output_dir / "provenance.json").read_text(encoding="utf-8")
    )
    assert provenance["analysis_settings"]["input_mode"] == "wavelength"
    assert provenance["analysis_settings"]["wavelength_A"] == pytest.approx(1.2)
    assert provenance["analysis_settings"]["energy_keV"] is None
    assert analysis.wavelength_source == "wavelength_A"
    assert analysis.wavelength_A == pytest.approx(1.2)


def test_quick_export_custom_source_wavelength_contract(
    demo_inputs: Path, tmp_path: Path
) -> None:
    result = quick_export(
        [demo_inputs],
        tmp_path / "custom-source",
        source_preset="Custom",
        wavelength_A=1.2,
        include_excel=False,
    )
    analysis = result.analyses[0]
    assert analysis.wavelength_source == "custom_source_wavelength"
    assert analysis.wavelength_A == pytest.approx(1.2)


def test_quick_export_cli_entry(demo_inputs: Path, tmp_path: Path) -> None:
    excel = tmp_path / "cli_out.xlsx"
    code = quick_export_main([str(demo_inputs), "-o", str(excel)])
    assert code == 0
    assert excel.is_file()
    assert (tmp_path / "cli_out_bundle" / "manifest.json").is_file()


@pytest.mark.parametrize(
    ("radiation_option", "radiation_value", "expected_mode", "expected_wavelength", "expected_energy"),
    [
        ("--wavelength-A", "1.2", "wavelength", pytest.approx(1.2), None),
        ("--energy-keV", "20", "energy", None, pytest.approx(20.0)),
    ],
)
def test_standalone_quick_export_custom_radiation_modes(
    demo_inputs: Path,
    tmp_path: Path,
    radiation_option: str,
    radiation_value: str,
    expected_mode: str,
    expected_wavelength: object,
    expected_energy: object,
) -> None:
    output = tmp_path / f"custom-{expected_mode}"
    code = quick_export_main(
        [
            str(demo_inputs),
            "-o",
            str(output),
            "--source",
            "Custom",
            radiation_option,
            radiation_value,
            "--no-elasticity",
            "--no-excel",
        ]
    )

    assert code == 0
    provenance = json.loads(
        (output / "provenance.json").read_text(encoding="utf-8")
    )
    settings = provenance["analysis_settings"]
    assert settings["input_mode"] == expected_mode
    assert settings["wavelength_A"] == expected_wavelength
    assert settings["energy_keV"] == expected_energy


def test_standalone_quick_export_custom_source_requires_explicit_radiation(
    demo_inputs: Path, tmp_path: Path, capsys
) -> None:
    code = quick_export_main(
        [
            str(demo_inputs),
            "-o",
            str(tmp_path / "missing-custom"),
            "--source",
            "Custom",
            "--no-elasticity",
            "--no-excel",
        ]
    )

    assert code == 2
    assert "Custom source requires --wavelength-A or --energy-keV" in capsys.readouterr().err


def test_standalone_quick_export_radiation_options_are_mutually_exclusive() -> None:
    with pytest.raises(SystemExit) as exc_info:
        build_quick_export_parser().parse_args(
            [
                "sample.cif",
                "-o",
                "out",
                "--wavelength-A",
                "1.2",
                "--energy-keV",
                "20",
            ]
        )
    assert exc_info.value.code == 2


def test_standalone_quick_export_help_is_legacy_windows_console_safe() -> None:
    help_text = build_quick_export_parser().format_help()
    assert "K-alpha" in help_text
    help_text.encode("cp936")


def test_elastic_override_replaces_sidecar(demo_inputs: Path, tmp_path: Path) -> None:
    """Override keyed by stem must be used instead of the demo sidecar."""

    override = parse_cubic_cij(250.0, 100.0, 50.0, source="override_test")
    output = tmp_path / "override_bundle"
    result = analyze_cifs(
        [demo_inputs],
        output,
        include_excel=False,
        elastic_overrides={"synthetic_fcc_al": override},
    )
    tensor = result.analyses[0].elastic_tensor
    assert tensor is not None
    assert tensor.stiffness_GPa[0, 0] == pytest.approx(250.0)
    assert tensor.source_provider == "user_input"
