from __future__ import annotations

import json
import csv
import shutil
from pathlib import Path

from openpyxl import load_workbook
import pytest

from diffractscout.demo import write_demo_inputs
from diffractscout.elasticity import discover_elastic_tensor
from diffractscout.models import CandidateRecord, DiscoverySettings, DownloadArtifact
from diffractscout.pipeline import (
    _copy_provider_artifact,
    _reconcile_download_coverage,
    run_pipeline,
)
from diffractscout.validation import verify_bundle


class OfflineProvider:
    name = "offline-fixture"

    def search_subsystem(self, chemsys: str, **_kwargs: object) -> list[CandidateRecord]:
        return [
            CandidateRecord(
                material_id="fixture-1",
                formula="Al",
                energy_above_hull_eV_atom=0.0,
                is_stable=True,
                space_group="F m -3 m",
                space_group_number=225,
                queried_chemsys=chemsys,
                source_provider=self.name,
                source_url="https://example.invalid/fixture-1",
            )
        ]

    def metadata(self) -> dict[str, object]:
        return {"provider": self.name, "database_version": "offline"}

    def download_candidates(
        self,
        candidates: object,
        output_dir: Path,
        **_kwargs: object,
    ) -> list[DownloadArtifact]:
        fixture_dir = write_demo_inputs(output_dir / "_fixture")
        cif = output_dir / "fixture-1.cif"
        elastic = output_dir / "fixture-1_elasticity.json"
        shutil.copy2(fixture_dir / "synthetic_fcc_al.cif", cif)
        shutil.copy2(fixture_dir / "synthetic_fcc_al_elasticity.json", elastic)
        shutil.rmtree(fixture_dir)
        candidate = list(candidates)[0]
        payload = json.loads(elastic.read_text(encoding="utf-8"))
        payload["provenance"]["material_id"] = candidate.material_id
        elastic.write_text(json.dumps(payload), encoding="utf-8")
        return [
            DownloadArtifact(
                candidate=candidate,
                cif_path=cif,
                elasticity_path=elastic,
                status="ok",
                provider_metadata=self.metadata(),
            )
        ]


class ExternalArtifactProvider(OfflineProvider):
    """Provider fixture that deliberately returns files outside staging/inputs."""

    def download_candidates(
        self,
        candidates: object,
        output_dir: Path,
        **_kwargs: object,
    ) -> list[DownloadArtifact]:
        cache = output_dir.parent / "provider-cache"
        self.staging_token = output_dir.parent.name
        fixture_dir = write_demo_inputs(cache)
        candidate = list(candidates)[0]
        output_dir.mkdir(parents=True, exist_ok=True)
        nested_metadata_path = output_dir / "provider-metadata.txt"
        nested_metadata_path.write_text("metadata", encoding="utf-8")
        elastic = fixture_dir / "synthetic_fcc_al_elasticity.json"
        payload = json.loads(elastic.read_text(encoding="utf-8"))
        payload["provenance"]["material_id"] = candidate.material_id
        elastic.write_text(json.dumps(payload), encoding="utf-8")
        return [
            DownloadArtifact(
                candidate=candidate,
                cif_path=fixture_dir / "synthetic_fcc_al.cif",
                elasticity_path=fixture_dir / "synthetic_fcc_al_elasticity.json",
                status="ok",
                provider_metadata={
                    "provider": self.name,
                    "nested": [
                        nested_metadata_path,
                        {"tuple": (nested_metadata_path, "staging/path")},
                    ],
                    "set": {nested_metadata_path},
                },
            )
        ]


class CollisionPairProvider(OfflineProvider):
    """Return distinct same-named CIF/sidecar pairs to force collision renaming."""

    def search_subsystem(self, chemsys: str, **_kwargs: object) -> list[CandidateRecord]:
        return [
            CandidateRecord(
                material_id=f"fixture-{index}",
                formula="Al",
                energy_above_hull_eV_atom=0.0,
                is_stable=True,
                space_group="F m -3 m",
                space_group_number=225,
                queried_chemsys=chemsys,
                source_provider=self.name,
                source_url=f"https://example.invalid/fixture-{index}",
            )
            for index in (1, 2)
        ]

    def download_candidates(
        self,
        candidates: object,
        output_dir: Path,
        **_kwargs: object,
    ) -> list[DownloadArtifact]:
        fixture_dir = write_demo_inputs(output_dir.parent / "collision-fixture")
        results: list[DownloadArtifact] = []
        for index, candidate in enumerate(list(candidates), start=1):
            source_dir = output_dir.parent / f"collision-source-{index}"
            source_dir.mkdir(parents=True, exist_ok=True)
            cif = source_dir / "shared.cif"
            shutil.copy2(fixture_dir / "synthetic_fcc_al.cif", cif)
            if index == 2:
                cif.write_text("# collision variant\n" + cif.read_text(encoding="utf-8"), encoding="utf-8")
            sidecar = source_dir / "shared_elasticity.json"
            payload = json.loads(
                (fixture_dir / "synthetic_fcc_al_elasticity.json").read_text(
                    encoding="utf-8"
                )
            )
            payload["cif_filename"] = cif.name
            payload["provenance"]["paired_cif"] = cif.name
            payload["provenance"]["material_id"] = candidate.material_id
            payload["diffractscout"]["paired_cif"] = cif.name
            sidecar.write_text(json.dumps(payload), encoding="utf-8")
            results.append(
                DownloadArtifact(
                    candidate=candidate,
                    cif_path=cif,
                    elasticity_path=sidecar,
                    status="ok",
                )
            )
        return results


class MismatchedSidecarProvider(OfflineProvider):
    def download_candidates(
        self,
        candidates: object,
        output_dir: Path,
        **_kwargs: object,
    ) -> list[DownloadArtifact]:
        fixture_dir = write_demo_inputs(output_dir.parent / "mismatch-fixture")
        cif = output_dir.parent / "mismatch.cif"
        sidecar = output_dir.parent / "mismatch_elasticity.json"
        shutil.copy2(fixture_dir / "synthetic_fcc_al.cif", cif)
        payload = json.loads(
            (fixture_dir / "synthetic_fcc_al_elasticity.json").read_text(encoding="utf-8")
        )
        payload["cif_filename"] = "different.cif"
        payload["provenance"]["paired_cif"] = "different.cif"
        sidecar.write_text(json.dumps(payload), encoding="utf-8")
        return [
            DownloadArtifact(
                candidate=list(candidates)[0],
                cif_path=cif,
                elasticity_path=sidecar,
                status="ok",
            )
        ]


def test_discovery_to_diffraction_bundle_without_network(tmp_path: Path) -> None:
    output = tmp_path / "complete"
    result = run_pipeline(
        "Al",
        OfflineProvider(),
        output,
        discovery_settings=DiscoverySettings(mode="single_chemsys", max_total=1),
        confirm_above=5,
    )
    assert len(result.discovery.candidates) == 1
    assert len(result.downloads) == 1
    assert len(result.analyses) == 1
    assert result.analyses[0].reflections[0].hkl == (1, 1, 1)
    assert verify_bundle(output)["ok"]


def test_provider_artifacts_are_copied_and_rebound_into_committed_bundle(
    tmp_path: Path,
) -> None:
    output = tmp_path / "external-artifacts"
    provider = ExternalArtifactProvider()
    result = run_pipeline(
        "Al",
        provider,
        output,
        discovery_settings=DiscoverySettings(mode="single_chemsys", max_total=1),
        confirm_above=5,
    )
    assert result.downloads[0].cif_path == output / "inputs" / "synthetic_fcc_al.cif"
    assert result.downloads[0].elasticity_path == (
        output / "inputs" / "synthetic_fcc_al_elasticity.json"
    )
    assert result.downloads[0].cif_path.is_file()
    assert result.downloads[0].elasticity_path.is_file()
    assert result.analyses[0].elastic_tensor is not None
    assert result.analyses[0].elastic_tensor.raw_payload_path == (
        output / "inputs" / "synthetic_fcc_al_elasticity.json"
    )
    metadata = result.downloads[0].provider_metadata
    assert metadata["nested"][0] == output / "inputs" / "provider-metadata.txt"
    assert metadata["nested"][1]["tuple"][0] == (
        output / "inputs" / "provider-metadata.txt"
    )
    assert metadata["nested"][1]["tuple"][1] == "staging/path"
    assert metadata["set"] == {output / "inputs" / "provider-metadata.txt"}
    assert metadata["nested"][0].is_file()
    # The provider's transaction-directory token must not survive in exports.
    assert provider.staging_token not in (
        result.output_dir / "provenance.json"
    ).read_text(encoding="utf-8")
    for name in ("download_index.csv", "phase_summary.csv", "elasticity.csv", "results.xlsx"):
        assert provider.staging_token not in (result.output_dir / name).read_bytes().decode(
            "utf-8", errors="ignore"
        )
    provenance = json.loads(
        (result.output_dir / "provenance.json").read_text(encoding="utf-8")
    )
    nested = provenance["downloads"][0]["provider_metadata"]["nested"]
    assert nested[0] == "inputs/provider-metadata.txt"
    assert nested[1]["tuple"][0] == "inputs/provider-metadata.txt"
    assert verify_bundle(output)["ok"]


def test_provider_artifact_copy_cleans_partial_file_on_failure(
    tmp_path: Path, monkeypatch
) -> None:
    import diffractscout.pipeline as pipeline

    source = tmp_path / "source.cif"
    source.write_bytes(b"complete CIF payload")
    destination = tmp_path / "inputs"

    def partial_copy(_source: str | Path, target: str | Path, **_kwargs: object) -> None:
        Path(target).write_bytes(b"partial")
        raise OSError("simulated copy failure")

    monkeypatch.setattr(pipeline.shutil, "copy2", partial_copy)
    with pytest.raises(OSError, match="simulated copy failure"):
        _copy_provider_artifact(
            source,
            destination,
            preferred_name="source.cif",
            suffix=".cif",
        )

    assert not list(destination.iterdir())


def test_provider_download_reconciliation_matches_material_ids_case_insensitively() -> None:
    requested = CandidateRecord(material_id="MP-1", formula="Al")
    returned_candidate = CandidateRecord(material_id="mp-1", formula="Al")
    returned = DownloadArtifact(
        candidate=returned_candidate,
        cif_path=None,
        status="failed",
        error="provider fixture",
    )

    reconciled, diagnostics = _reconcile_download_coverage([requested], [returned])

    assert diagnostics == []
    assert len(reconciled) == 1
    assert reconciled[0].candidate is returned_candidate
    assert reconciled[0].candidate.material_id == "mp-1"


def test_successful_candidate_without_cif_or_elasticity_is_diagnosed(tmp_path: Path) -> None:
    class MissingArtifactsProvider(OfflineProvider):
        def download_candidates(
            self,
            candidates: object,
            output_dir: Path,
            **_kwargs: object,
        ) -> list[DownloadArtifact]:
            return [
                DownloadArtifact(
                    candidate=list(candidates)[0],
                    cif_path=None,
                    elasticity_path=None,
                    status="ok",
                    elasticity_status="ok",
                )
            ]

    result = run_pipeline(
        "Al",
        MissingArtifactsProvider(),
        tmp_path / "missing-artifacts",
        discovery_settings=DiscoverySettings(mode="single_chemsys", max_total=1),
        confirm_above=5,
    )
    messages = [item.message for item in result.diagnostics if item.level == "error"]
    assert any("without a CIF artifact" in message for message in messages)
    assert any("without an elasticity artifact" in message for message in messages)
    download = result.downloads[0]
    assert download.elasticity_status == "no_elasticity_data"
    assert download.elasticity_error
    with (result.output_dir / "download_index.csv").open(
        encoding="utf-8-sig", newline=""
    ) as handle:
        row = next(csv.DictReader(handle))
    assert row["elasticity_status"] == download.elasticity_status
    assert row["elasticity_error"] == download.elasticity_error
    workbook = load_workbook(result.output_dir / "results.xlsx", data_only=True, read_only=True)
    sheet = workbook["Downloads"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    values = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
    exported = dict(zip(headers, values, strict=True))
    assert exported["elasticity_status"] == download.elasticity_status
    assert exported["elasticity_error"] == download.elasticity_error


def test_provider_download_coverage_rejects_missing_duplicate_and_unrequested(
    tmp_path: Path,
) -> None:
    class CoverageProvider(OfflineProvider):
        def search_subsystem(self, chemsys: str, **_kwargs: object) -> list[CandidateRecord]:
            return [
                CandidateRecord(
                    material_id=f"coverage-{index}",
                    formula="Al",
                    energy_above_hull_eV_atom=0.0,
                    is_stable=True,
                    queried_chemsys=chemsys,
                    source_provider=self.name,
                )
                for index in (1, 2)
            ]

        def download_candidates(
            self,
            candidates: object,
            _output_dir: Path,
            **_kwargs: object,
        ) -> list[DownloadArtifact]:
            requested = list(candidates)
            duplicate = DownloadArtifact(
                candidate=requested[0],
                cif_path=None,
                status="ok",
            )
            return [
                duplicate,
                duplicate,
                DownloadArtifact(
                    candidate=CandidateRecord(material_id="coverage-extra", formula="Al"),
                    cif_path=None,
                    status="ok",
                ),
            ]

    result = run_pipeline(
        "Al",
        CoverageProvider(),
        tmp_path / "coverage-failures",
        discovery_settings=DiscoverySettings(mode="single_chemsys", max_total=2),
        include_elasticity=False,
        confirm_above=5,
    )

    assert [item.candidate.material_id for item in result.downloads] == [
        "coverage-1",
        "coverage-1",
        "coverage-2",
        "coverage-extra",
    ]
    assert all(item.status == "failed" for item in result.downloads)
    messages = [item.message for item in result.diagnostics if item.level == "error"]
    assert any("duplicate provider download" in message for message in messages)
    assert any("missing provider download" in message for message in messages)
    assert any("unrequested provider download" in message for message in messages)
    assert result.analyses == []


def test_copy_failure_outputs_do_not_leak_staging_transaction_path(tmp_path: Path) -> None:
    class CopyFailureProvider(OfflineProvider):
        def download_candidates(
            self,
            candidates: object,
            output_dir: Path,
            **_kwargs: object,
        ) -> list[DownloadArtifact]:
            fixture_dir = write_demo_inputs(output_dir.parent / "copy-failure-fixture")
            output_dir.mkdir(parents=True, exist_ok=True)
            cif = output_dir / "copy-failure.cif"
            shutil.copy2(fixture_dir / "synthetic_fcc_al.cif", cif)
            self.staging_token = output_dir.parent.name
            return [
                DownloadArtifact(
                    candidate=list(candidates)[0],
                    cif_path=cif,
                    elasticity_path=output_dir / "missing_elasticity.json",
                    status="ok",
                )
            ]

    output = tmp_path / "copy-failure"
    provider = CopyFailureProvider()
    result = run_pipeline(
        "Al",
        provider,
        output,
        discovery_settings=DiscoverySettings(mode="single_chemsys", max_total=1),
        confirm_above=5,
    )

    assert result.downloads[0].elasticity_status == "elasticity_query_failed"
    assert provider.staging_token not in result.downloads[0].elasticity_error
    assert provider.staging_token not in "\n".join(
        item.message for item in result.diagnostics
    )
    for name in ("diagnostics.csv", "download_index.csv", "provenance.json"):
        text = (output / name).read_text(encoding="utf-8-sig")
        assert provider.staging_token not in text


def test_collision_renamed_sidecar_remains_paired_and_rediscoverable(
    tmp_path: Path,
) -> None:
    output = tmp_path / "collision-pairs"
    result = run_pipeline(
        "Al",
        CollisionPairProvider(),
        output,
        discovery_settings=DiscoverySettings(mode="single_chemsys", max_total=2),
        confirm_above=5,
    )
    assert len(result.downloads) == 2
    assert result.downloads[0].elasticity_status == "ok"
    assert result.downloads[1].elasticity_status == "ok"
    names = [item.cif_path.name for item in result.downloads if item.cif_path]
    assert names[0] == "shared.cif"
    assert names[1].startswith("shared_")
    for index, item in enumerate(result.downloads):
        assert item.cif_path is not None
        assert item.elasticity_path is not None
        assert item.cif_path.is_file()
        assert item.elasticity_path.is_file()
        payload = json.loads(item.elasticity_path.read_text(encoding="utf-8"))
        assert payload["cif_filename"] == item.cif_path.name
        assert payload["provenance"]["paired_cif"] == item.cif_path.name
        assert discover_elastic_tensor(item.cif_path) is not None
        if index == 1:
            assert payload["provenance"]["original_pairing"]["cif_filename"] == "shared.cif"
    assert verify_bundle(output)["ok"]


def test_mismatched_sidecar_is_failed_closed_in_memory_csv_and_excel(
    tmp_path: Path,
) -> None:
    output = tmp_path / "mismatched-sidecar"
    result = run_pipeline(
        "Al",
        MismatchedSidecarProvider(),
        output,
        discovery_settings=DiscoverySettings(mode="single_chemsys", max_total=1),
        confirm_above=5,
    )
    download = result.downloads[0]
    assert download.elasticity_status == "invalid"
    assert "do not match" in download.elasticity_error
    assert result.analyses[0].elastic_tensor is None
    assert any(
        item.level == "error" and download.elasticity_error in item.message
        for item in result.diagnostics
    )
    with (output / "download_index.csv").open(encoding="utf-8-sig", newline="") as handle:
        row = next(csv.DictReader(handle))
    assert row["elasticity_status"] == download.elasticity_status
    assert row["elasticity_error"] == download.elasticity_error
    workbook = load_workbook(output / "results.xlsx", data_only=True, read_only=True)
    sheet = workbook["Downloads"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    values = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
    exported = dict(zip(headers, values, strict=True))
    assert exported["elasticity_status"] == download.elasticity_status
    assert exported["elasticity_error"] == download.elasticity_error
    assert verify_bundle(output)["ok"]


def test_unreadable_sidecar_updates_download_status_and_error(tmp_path: Path) -> None:
    class MissingSidecarProvider(OfflineProvider):
        def download_candidates(
            self,
            candidates: object,
            output_dir: Path,
            **_kwargs: object,
        ) -> list[DownloadArtifact]:
            fixture_dir = write_demo_inputs(output_dir.parent / "missing-sidecar-fixture")
            cif = output_dir.parent / "missing-sidecar.cif"
            shutil.copy2(fixture_dir / "synthetic_fcc_al.cif", cif)
            return [
                DownloadArtifact(
                    candidate=list(candidates)[0],
                    cif_path=cif,
                    elasticity_path=output_dir.parent / "not-present.json",
                    status="ok",
                )
            ]

    output = tmp_path / "unreadable-sidecar"
    result = run_pipeline(
        "Al",
        MissingSidecarProvider(),
        output,
        discovery_settings=DiscoverySettings(mode="single_chemsys", max_total=1),
        confirm_above=5,
    )
    download = result.downloads[0]
    assert download.elasticity_status == "elasticity_query_failed"
    assert download.elasticity_error
    assert result.analyses[0].elastic_tensor is None
    with (output / "download_index.csv").open(encoding="utf-8-sig", newline="") as handle:
        row = next(csv.DictReader(handle))
    assert row["elasticity_status"] == "elasticity_query_failed"
    assert row["elasticity_error"] == download.elasticity_error


def test_primitive_download_rejects_automatic_elastic_pairing(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="Primitive-cell downloads"):
        run_pipeline(
            "Al",
            OfflineProvider(),
            tmp_path / "primitive",
            conventional_unit_cell=False,
            include_elasticity=True,
        )
    assert not (tmp_path / "primitive").exists()
