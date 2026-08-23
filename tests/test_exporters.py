import csv
import json
from pathlib import Path

import pytest

from openpyxl import load_workbook

import diffractscout.exporters as exporters
from diffractscout.demo import write_demo_inputs
from diffractscout.exporters import (
    _bundle_readme,
    _pattern_axis_coordinates,
    _summary_rows,
    _write_csv,
    export_result_bundle,
    write_excel_workbook,
)
from diffractscout.models import AnalysisSettings
from diffractscout.pipeline import analyze_cifs


def test_csv_and_excel_escape_formula_like_external_text(tmp_path: Path) -> None:
    csv_path = tmp_path / "safe.csv"
    _write_csv(csv_path, [{"value": "=HYPERLINK(\"https://example.invalid\")"}], ["value"])
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        row = next(csv.DictReader(handle))
    assert row["value"].startswith("'=")

    workbook_path = tmp_path / "safe.xlsx"
    write_excel_workbook(
        workbook_path,
        summary=[{"key": "external", "value": "+SUM(1,1)"}],
        phases=[],
        peaks=[],
        elasticity=[],
        candidates=[],
        downloads=[],
        diagnostics=[],
        patterns=[],
    )
    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    assert workbook["Summary"]["B2"].value == "'+SUM(1,1)"


def test_empty_csv_keeps_a_stable_header(tmp_path: Path) -> None:
    path = tmp_path / "empty.csv"
    _write_csv(path, [], ["a", "b"])
    assert path.read_text(encoding="utf-8-sig").splitlines() == ["a,b"]


def test_pattern_axis_coordinates_keep_zero_q_and_g_finite() -> None:
    assert _pattern_axis_coordinates(0.0, 1.5406, "q") == (None, 0.0, 0.0, 0.0)
    assert _pattern_axis_coordinates(0.0, 1.5406, "g") == (None, 0.0, 0.0, 0.0)
    d_A, q_invA, g_invA, selected = _pattern_axis_coordinates(30.0, 1.5406, "q")
    assert d_A == pytest.approx(2.976210656)
    assert q_invA == pytest.approx(2.0 * 3.141592653589793 / d_A)
    assert g_invA == pytest.approx(1.0 / d_A)
    assert selected == q_invA


def test_summary_rows_distinguish_requested_and_effective_settings() -> None:
    settings = AnalysisSettings(
        input_mode="energy",
        energy_keV=20.0,
        two_theta_min_deg=5.0,
        two_theta_max_deg=120.0,
        d_min_A=0.8,
        d_max_A=4.0,
    )
    summary = _summary_rows([], None, settings, [], include_excel=False)
    values = {row["key"]: row["value"] for row in summary}
    assert values["two_theta_range_deg"] == [5.0, 120.0]
    assert values["requested_two_theta_range_deg"] == [5.0, 120.0]
    assert values["effective_two_theta_range_deg"] == [5.0, 120.0]
    assert values["effective_wavelength_A"] is None
    assert values["effective_energy_keV"] is None
    assert values["source_preset_applied"] is None
    assert values["export_lab_views"] is True
    assert values["effective_export_lab_views"] is False
    assert values["include_excel"] is False


def test_summary_rows_read_back_effective_d_filtered_energy_analysis(
    tmp_path: Path,
) -> None:
    inputs = write_demo_inputs(tmp_path / "inputs")
    settings = AnalysisSettings(
        input_mode="energy",
        energy_keV=20.0,
        two_theta_min_deg=5.0,
        two_theta_max_deg=120.0,
        d_min_A=0.8,
        d_max_A=4.0,
    )
    result = analyze_cifs(
        [inputs],
        tmp_path / "energy-d-filtered",
        settings=settings,
        include_excel=True,
    )
    workbook = load_workbook(
        result.output_dir / "results.xlsx", read_only=True, data_only=True
    )
    values = {
        row[0]: row[1]
        for row in workbook["Summary"].iter_rows(min_row=2, values_only=True)
        if row[0]
    }
    workbook.close()
    metadata = result.analyses[0].metadata
    assert json.loads(values["two_theta_range_deg"]) == [5.0, 120.0]
    assert json.loads(values["requested_two_theta_range_deg"]) == [5.0, 120.0]
    assert json.loads(values["analysis_two_theta_range_deg"]) == pytest.approx(
        metadata["two_theta_range_deg"]
    )
    assert json.loads(
        values["profile_sampled_two_theta_range_deg"]
    ) == pytest.approx(metadata["profile_sampled_two_theta_range_deg"])
    assert json.loads(values["effective_two_theta_range_deg"]) == pytest.approx(
        metadata["two_theta_range_deg"]
    )
    assert values["effective_wavelength_A"] == pytest.approx(12.398419843320026 / 20.0)
    assert values["effective_energy_keV"] == pytest.approx(20.0)
    assert values["effective_radiation_source"] == "energy_keV"
    assert values["source_preset_applied"] in {None, ""}


def test_provenance_records_effective_conditional_output_flags(tmp_path: Path) -> None:
    inputs = write_demo_inputs(tmp_path / "inputs")
    settings = AnalysisSettings(
        include_patterns=False,
        include_figures=False,
        export_lab_views=True,
    )
    result = analyze_cifs(
        [inputs],
        tmp_path / "no-optional-outputs",
        settings=settings,
        include_excel=False,
    )
    payload = json.loads((result.output_dir / "provenance.json").read_text(encoding="utf-8"))
    flags = payload["output_flags"]

    assert flags == {
        "include_excel": False,
        "include_patterns": False,
        "include_figures": False,
        "export_lab_views": False,
        "effective_export_lab_views": False,
    }
    assert not (result.output_dir / "results.xlsx").exists()
    assert not (result.output_dir / "pattern_profiles.csv").exists()
    assert "results.xlsx" not in (result.output_dir / "README.md").read_text(encoding="utf-8")
    assert "pattern_profiles.csv" not in (result.output_dir / "README.md").read_text(encoding="utf-8")
    assert "always" not in payload["definitions"]["pattern_axis_columns"].lower()


def test_bundle_readme_is_conditional_and_names_profile_model() -> None:
    settings = AnalysisSettings(include_patterns=False, profile_model="gaussian")
    readme = _bundle_readme([], None, [], settings, include_excel=False)
    assert "pattern_profiles.csv" not in readme
    assert "results.xlsx" not in readme
    assert "gaussian" in readme.lower()

    enabled = AnalysisSettings(include_patterns=True, profile_model="lorentzian", export_lab_views=True)
    enabled_readme = _bundle_readme([], None, [], enabled, include_excel=True)
    assert "pattern_profiles.csv" in enabled_readme
    assert "results.xlsx" in enabled_readme
    assert "lorentzian" in enabled_readme
    assert "lab views" in enabled_readme.lower()


def test_excel_omission_warning_reaches_caller_diagnostics(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(exporters, "EXCEL_DATA_ROW_LIMIT", 2)
    monkeypatch.setattr(
        exporters,
        "phase_rows",
        lambda _analyses: [{"phase_name": str(index)} for index in range(3)],
    )
    diagnostics = []
    bundle = tmp_path / "diagnostic-retention"
    export_result_bundle(
        bundle,
        analyses=[],
        settings=AnalysisSettings(include_patterns=False),
        diagnostics=diagnostics,
        include_excel=True,
    )
    assert diagnostics
    warning = diagnostics[0]
    assert warning.level == "warning"
    assert "Excel data-row limit" in warning.message
    csv_text = (bundle / "diagnostics.csv").read_text(encoding="utf-8-sig")
    assert warning.message in csv_text
    workbook = load_workbook(bundle / "results.xlsx", read_only=True, data_only=False)
    assert warning.message in str(workbook["Diagnostics"]["D2"].value)


def test_write_text_atomic_cleans_unique_temp_after_replace_failure(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "README.md"
    path.write_text("old", encoding="utf-8")

    def fail_replace(self: Path, _target: Path) -> Path:
        raise OSError("simulated replace failure")

    monkeypatch.setattr(Path, "replace", fail_replace)
    with pytest.raises(OSError, match="simulated replace failure"):
        exporters._write_text_atomic(path, "new")

    assert path.read_text(encoding="utf-8") == "old"
    assert not list(tmp_path.glob("*.tmp"))
    assert not list(tmp_path.glob(".README.md.*.tmp"))


def test_write_text_atomic_preserves_replace_error_when_cleanup_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "README.md"

    def fail_replace(self: Path, _target: Path) -> Path:
        raise OSError("primary replace failure")

    def fail_unlink(self: Path, *, missing_ok: bool = False) -> None:
        raise PermissionError("cleanup failure")

    monkeypatch.setattr(Path, "replace", fail_replace)
    monkeypatch.setattr(Path, "unlink", fail_unlink)

    with pytest.raises(OSError, match="primary replace failure"):
        exporters._write_text_atomic(path, "new")


def test_write_text_atomic_returns_published_path_when_cleanup_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "README.md"

    def fail_unlink(self: Path, *, missing_ok: bool = False) -> None:
        raise PermissionError("cleanup failure")

    monkeypatch.setattr(Path, "unlink", fail_unlink)

    assert exporters._write_text_atomic(path, "new") == path
    assert path.read_text(encoding="utf-8") == "new"


def test_excel_uses_an_omission_note_above_row_limit(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setattr(exporters, "EXCEL_DATA_ROW_LIMIT", 2)
    workbook_path = tmp_path / "limited.xlsx"
    write_excel_workbook(
        workbook_path,
        summary=[],
        phases=[],
        peaks=[{"phase_name": str(index)} for index in range(3)],
        elasticity=[],
        candidates=[],
        downloads=[],
        diagnostics=[],
        patterns=[],
    )
    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    assert workbook["Peaks"]["A2"].value == "omitted_from_workbook"
    assert "corresponding CSV" in workbook["Peaks"]["B2"].value


def test_provenance_records_portable_runtime_environment(tmp_path: Path) -> None:
    inputs = write_demo_inputs(tmp_path / "inputs")
    result = analyze_cifs([inputs], tmp_path / "bundle", include_excel=False)
    payload = json.loads((result.output_dir / "provenance.json").read_text(encoding="utf-8"))
    environment = payload["runtime_environment"]
    assert environment["python"]
    assert environment["operating_system"]
    assert "executable" not in environment
    assert "hostname" not in environment
